"""Build the bundled built-in report templates (docx + preparation specs).

Run from the repo root:
    .venv/bin/python scripts/build_builtin_template.py

Thirteen templates mirror the measured Workbuddy/tencent-docx capability set:
8 genre layouts, each with its canonical theme, plus the 5 theme variants of
the business-report layout. Every template is a single-section DOCX (the
import constraint) whose cover paragraphs carry {{title}}/{{organization}}/
{{period}}/{{report_date}} placeholders; chapter headings exist only as style
anchors and are removed by prepare(); each sample table exists so
table_defaults() captures its borders and shading. The spec records
body-block indices, so any structural change must be regenerated with the
spec — both files are written in one run.
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ASSETS = Path(__file__).resolve().parent.parent / 'src' / 'briefloop' / 'template_assets'
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


# ---------------------------------------------------------------- primitives

def style_fonts(style, *, size, bold=False, color='1E2320', east='宋体', western='Arial'):
    style.font.name = western
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn('w:eastAsia'), east)


def run_props(run, *, size=None, bold=False, color='1E2320', east='宋体', western='Arial'):
    run.font.name = western
    if size: run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), east)


def shade_element(ppr, fill):
    element = OxmlElement('w:shd')
    element.set(qn('w:val'), 'clear')
    element.set(qn('w:fill'), fill)
    ppr.append(element)


def shade(paragraph, fill):
    shade_element(paragraph._p.get_or_add_pPr(), fill)


def bottom_border_ppr(ppr, *, color, sz='6', space='4'):
    border = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for key, value in (('val', 'single'), ('sz', sz), ('color', color), ('space', space)):
        bottom.set(qn('w:' + key), value)
    border.append(bottom)
    ppr.append(border)


def bottom_border(paragraph, *, color, sz='6', space='4'):
    bottom_border_ppr(paragraph._p.get_or_add_pPr(), color=color, sz=sz, space=space)


def page_break(doc):
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    return paragraph


def table_borders(table, spec):
    borders = OxmlElement('w:tblBorders')
    for edge, kind, size in spec:
        element = OxmlElement('w:' + edge)
        element.set(qn('w:val'), kind)
        element.set(qn('w:sz'), size)
        element.set(qn('w:color'), '1E2320')
        borders.append(element)
    table._tbl.tblPr.append(borders)


THREE_LINE = (('top', 'single', '12'), ('bottom', 'single', '12'),
              ('left', 'none', '0'), ('right', 'none', '0'),
              ('insideH', 'none', '0'), ('insideV', 'none', '0'))
FULL_BORDER = (('top', 'single', '4'), ('bottom', 'single', '4'),
               ('left', 'single', '4'), ('right', 'single', '4'),
               ('insideH', 'single', '4'), ('insideV', 'single', '4'))
HORIZONTAL = (('top', 'single', '8'), ('bottom', 'single', '8'),
              ('left', 'none', '0'), ('right', 'none', '0'),
              ('insideH', 'single', '4'), ('insideV', 'none', '0'))


def add_sample_table(doc, *, kind, header_fill=None, rows=3, cols=3):
    table = doc.add_table(rows=rows, cols=cols)
    table_borders(table, {'three_line': THREE_LINE, 'full': FULL_BORDER, 'horizontal': HORIZONTAL}[kind])
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.paragraphs[0].add_run(f'示例{r+1}-{c+1}')
            if r == 0:
                cell.paragraphs[0].runs[0].font.bold = True
                if header_fill:
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:val'), 'clear')
                    shading.set(qn('w:fill'), header_fill)
                    cell._tc.get_or_add_tcPr().append(shading)
                cell_borders = OxmlElement('w:tcBorders')
                bottom = OxmlElement('w:bottom')
                for key, value in (('val', 'single'), ('sz', '8'), ('color', '1E2320')):
                    bottom.set(qn('w:' + key), value)
                cell_borders.append(bottom)
                cell._tc.get_or_add_tcPr().append(cell_borders)
    return table


def add_page_number(section, mode):
    section.different_first_page_header_footer = True
    section.first_page_footer.paragraphs[0].text = ''
    section.first_page_header.paragraphs[0].text = ''
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), ' PAGE ')
    inner = OxmlElement('w:r')
    text = OxmlElement('w:t'); text.text = '1'
    inner.append(text); field.append(inner)
    if mode == 'gov':
        footer.add_run('— '); footer._p.append(field); footer.add_run(' —')
    elif mode == 'total':
        footer.add_run('第 '); footer._p.append(field)
        total = OxmlElement('w:fldSimple')
        total.set(qn('w:instr'), ' NUMPAGES ')
        inner2 = OxmlElement('w:r'); text2 = OxmlElement('w:t'); text2.text = '1'
        inner2.append(text2); total.append(inner2)
        footer._p.append(total); footer.add_run(' 页')
    else:
        footer.add_run('第 '); footer._p.append(field); footer.add_run(' 页')
    for run in footer.runs: run.font.size = Pt(9)


def add_title_header(section, text, color):
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(text)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(color)
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    bottom_border(header, color='DEDFD8', sz='4')


def add_toc_block(doc, color):
    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(6)
    label.paragraph_format.space_after = Pt(10)
    run = label.add_run('目录')
    run_props(run, size=15, bold=True, color=color, east='黑体')
    holder = doc.add_paragraph()
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), ' TOC \\o "1-3" \\h \\z \\u ')
    placeholder = OxmlElement('w:r')
    placeholder_text = OxmlElement('w:t')
    placeholder_text.text = '目录将在打开文档时自动生成'
    placeholder.append(placeholder_text)
    field.append(placeholder)
    holder._p.append(field)


def customize(doc, cfg):
    style_fonts(doc.styles['Normal'], size=cfg['body']['size'], color=cfg['body']['color'], east=cfg['body']['east'])
    doc.styles['Normal'].paragraph_format.line_spacing = cfg.get('line_spacing', 1.4)
    for level, spec in cfg['headings'].items():
        style = doc.styles[f'Heading {level}']
        style_fonts(style, size=spec['size'], bold=True, color=spec['color'], east=spec['east'])
        if spec.get('fill'):
            shade_element(style.element.get_or_add_pPr(), spec['fill'])
        if spec.get('border'):
            bottom_border_ppr(style.element.get_or_add_pPr(), color=spec['border'])


def chapter(doc, text, spec):
    paragraph = doc.add_paragraph(style=f"Heading {spec.get('level', 1)}")
    run = paragraph.add_run(text)
    run_props(run, size=spec['size'], bold=True, color=spec['color'], east=spec['east'])
    if spec.get('fill'):
        shade(paragraph, spec['fill'])
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(8)
    if spec.get('border'):
        bottom_border(paragraph, color=spec['border'])
    return paragraph


def centered(doc, text, *, size, bold=False, color='1E2320', east='宋体', space_before=0):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(space_before)
    if text:
        run = paragraph.add_run(text)
        run_props(run, size=size, bold=bold, color=color, east=east)
    return paragraph


def body_sample(doc, text='这是一段正文样式示例：BriefLoop 会把该段的字体、字号、行距与格式作为报告正文样式。Numbers like 1,234.56 keep Arial while 汉字落到中文字体。'):
    return doc.add_paragraph(text)


def key_value_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table_borders(table, HORIZONTAL)
    for r, (key, value) in enumerate(rows):
        left = table.rows[r].cells[0].paragraphs[0].add_run(key)
        run_props(left, size=10.5, bold=True)
        right = table.rows[r].cells[1].paragraphs[0].add_run(value)
        run_props(right, size=10.5)
    return table


# --------------------------------------------------------------- genre covers

def cover_standard(doc, cfg):
    centered(doc, '{{organization}}', size=11, color=cfg.get('muted', '6A706B'), space_before=90)
    title = centered(doc, '{{title}}', size=cfg.get('title_size', 24), bold=True, color=cfg['headings'][1]['color'], east=cfg['headings'][1]['east'])
    title.style = doc.styles['Heading 1']
    centered(doc, '{{period}} · {{report_date}}', size=11, color=cfg.get('muted', '6A706B'), space_before=16)


def cover_gov(doc, cfg):
    centered(doc, '{{organization}}', size=36, bold=True, color='C8102E', east='宋体', space_before=60)
    centered(doc, '〔2026〕第　　号', size=14, color='1E2320', east='仿宋_GB2312', space_before=8)
    rule = doc.add_paragraph()
    bottom_border(rule, color='C8102E', sz='18', space='1')
    title = centered(doc, '{{title}}', size=22, bold=True, color='1E2320', east='方正小标宋简体', space_before=24)
    centered(doc, '{{report_date}}', size=14, color='1E2320', east='仿宋_GB2312', space_before=10)
    receiver = doc.add_paragraph('各有关单位：')
    run_props(receiver.runs[0], size=16, color='1E2320', east='仿宋_GB2312')


def cover_contract(doc, cfg):
    centered(doc, '{{title}}', size=18, bold=True, color=cfg['headings'][1]['color'], east=cfg['headings'][1]['east'], space_before=24)
    centered(doc, '合同编号：________　签订日期：{{report_date}}', size=10.5, color='6A706B', space_before=8)
    doc.add_paragraph()
    key_value_table(doc, [('甲方（委托方）：', '＿＿＿＿＿＿＿＿＿＿'), ('乙方（受托方）：', '＿＿＿＿＿＿＿＿＿＿'),
                          ('签订地点：', '＿＿＿＿＿＿＿＿＿＿'), ('{{organization}}', '＿')])


def cover_minutes(doc, cfg):
    centered(doc, '{{title}}', size=18, bold=True, color=cfg['headings'][1]['color'], east=cfg['headings'][1]['east'], space_before=16)
    doc.add_paragraph()
    key_value_table(doc, [('会议时间：', '{{report_date}}'), ('会议地点：', '＿＿＿＿＿＿'), ('主持人：', '＿＿＿＿＿＿'),
                          ('记录人：', '＿＿＿＿＿＿'), ('出席人员：', '＿＿＿＿＿＿'), ('{{organization}}', '＿')])


def cover_research(doc, cfg):
    centered(doc, '{{organization}}', size=11, color=cfg.get('muted', '6A706B'), space_before=48)
    title = centered(doc, '{{title}}', size=20, bold=True, color='FFFFFF', east=cfg['headings'][1]['east'])
    shade(title, cfg['headings'][1]['fill'])
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    centered(doc, '{{period}} · {{report_date}}', size=11, color=cfg.get('muted', '6A706B'), space_before=12)
    doc.add_paragraph()
    key_value_table(doc, [('投资评级：', '待定（首次覆盖时更新）'), ('目标价：', '＿＿＿＿'), ('现价：', '＿＿＿＿'), ('报告日期：', '{{report_date}}')])


# ------------------------------------------------------------ genre registry

INK = '1E2320'
BODY_STANDARD = {'size': 10.5, 'color': INK, 'east': '宋体'}
BUSINESS_BODY = {'size': 9, 'color': '262626', 'east': '微软雅黑'}
GOV_BODY = {'size': 16, 'color': INK, 'east': '仿宋_GB2312'}
ACADEMIC_BODY = {'size': 12, 'color': INK, 'east': '宋体'}

NAVY = '1A237E'
BUSINESS_HEADINGS = {1: {'size': 13, 'color': 'FFFFFF', 'east': '微软雅黑', 'fill': NAVY},
                     2: {'size': 11, 'color': NAVY, 'east': '微软雅黑'},
                     3: {'size': 10, 'color': '37474F', 'east': '微软雅黑'}}

GENRES = (
    {'stem': 'general-report-zh', 'label': '通用报告',
     'sections': (('summary', '摘要', '本期核心结论'), ('background', '背景', '事件与背景'),
                  ('analysis', '分析', '影响推演'), ('conclusion', '结论', '判断与建议'),
                  ('risks', '风险提示', '不确定性')),
     'body': BODY_STANDARD, 'muted': '6A706B', 'accent': '006838', 'line_spacing': 1.4,
     'headings': {1: {'size': 15, 'color': '006838', 'east': '黑体', 'border': '006838'},
                  2: {'size': 12.5, 'color': INK, 'east': '黑体'},
                  3: {'size': 11, 'color': '37474F', 'east': '黑体'}},
     'cover': cover_standard, 'table': {'kind': 'three_line'}, 'footer': 'page', 'header': True, 'toc': True},

    {'stem': 'business-report-zh', 'label': '商务报告',
     'sections': (('views', '核心观点', '开门见山的判断'), ('events', '事件回顾', '本期大事'),
                  ('impact', '影响分析', '短中长期影响'), ('data', '数据附录', '关键指标'),
                  ('risks', '风险提示', '不确定性')),
     'body': BUSINESS_BODY, 'muted': '78909C', 'accent': NAVY, 'line_spacing': 1.3,
     'headings': BUSINESS_HEADINGS,
     'cover': cover_standard, 'table': {'kind': 'horizontal', 'header_fill': 'EEF2F9'}, 'footer': 'page', 'header': True,
     'toc': True, 'title_size': 22},

    {'stem': 'academic-paper-zh', 'label': '学术论文',
     'sections': (('abstract', '摘要', '研究概述'), ('intro', '引言', '问题与贡献'),
                  ('method', '研究方法', '设计与数据'), ('results', '结果', '发现'),
                  ('discussion', '讨论', '解释与局限'), ('conclusion', '结论', '结论与展望'),
                  ('references', '参考文献', '引用列表')),
     'body': ACADEMIC_BODY, 'muted': '555555', 'accent': INK, 'line_spacing': 1.5,
     'headings': {1: {'size': 15, 'color': INK, 'east': '黑体'},
                  2: {'size': 13, 'color': INK, 'east': '黑体'},
                  3: {'size': 12, 'color': INK, 'east': '宋体', 'bold': True}},
     'cover': cover_standard, 'table': {'kind': 'three_line'}, 'footer': 'page', 'header': False, 'toc': False,
     'title_size': 20},

    {'stem': 'government-doc-zh', 'label': '政府公文',
     'sections': (('first', '一、总体要求', '总体要求'), ('second', '二、重点任务', '重点任务'),
                  ('third', '三、保障措施', '保障措施')),
     'body': GOV_BODY, 'muted': '1E2320', 'accent': 'C8102E', 'line_spacing': 1.5,
     'headings': {1: {'size': 16, 'color': INK, 'east': '黑体'},
                  2: {'size': 16, 'color': INK, 'east': '楷体'},
                  3: {'size': 16, 'color': INK, 'east': '仿宋_GB2312', 'bold': True}},
     'cover': cover_gov, 'table': {'kind': 'full'}, 'footer': 'gov', 'header': False, 'toc': False},

    {'stem': 'annual-report-zh', 'label': '上市公司年报',
     'sections': (('profile', '公司概况', '基本情况'), ('financials', '主要财务数据', '关键数字'),
                  ('discussion', '经营层讨论与分析', '经营回顾'), ('risks', '风险因素', '风险与应对'),
                  ('appendix', '附录', '治理与备查')),
     'body': BODY_STANDARD, 'muted': '6A706B', 'accent': 'C00000', 'line_spacing': 1.35,
     'headings': {1: {'size': 15, 'color': 'C00000', 'east': '黑体', 'border': 'C00000'},
                  2: {'size': 12.5, 'color': 'C00000', 'east': '黑体'},
                  3: {'size': 11, 'color': INK, 'east': '黑体'}},
     'cover': cover_standard, 'table': {'kind': 'full'}, 'footer': 'page', 'header': True, 'toc': True},

    {'stem': 'legal-contract-zh', 'label': '合同',
     'sections': (('clause1', '第一条 合作宗旨', '宗旨'), ('clause2', '第二条 权利与义务', '权利义务'),
                  ('clause3', '第三条 费用与支付', '费用支付'), ('clause4', '第四条 保密条款', '保密'),
                  ('clause5', '第五条 违约责任', '违约责任')),
     'body': ACADEMIC_BODY, 'muted': '555555', 'accent': INK, 'line_spacing': 1.5,
     'headings': {1: {'size': 13, 'color': INK, 'east': '黑体'},
                  2: {'size': 12, 'color': INK, 'east': '宋体', 'bold': True}},
     'cover': cover_contract, 'table': {'kind': 'full'}, 'footer': 'total', 'header': False, 'toc': False},

    {'stem': 'meeting-minutes-zh', 'label': '会议纪要',
     'sections': (('topics', '会议议题', '讨论议题'), ('decisions', '决议事项', '形成的决议'),
                  ('actions', '待办事项', '责任与时限'), ('others', '其他事项', '补充记录')),
     'body': ACADEMIC_BODY, 'muted': '555555', 'accent': INK, 'line_spacing': 1.5,
     'headings': {1: {'size': 14, 'color': INK, 'east': '黑体'},
                  2: {'size': 12, 'color': '37474F', 'east': '黑体'}},
     'cover': cover_minutes, 'table': {'kind': 'full'}, 'footer': 'page', 'header': False, 'toc': False},

    {'stem': 'stock-research-zh', 'label': '券商研报',
     'sections': (('views', '核心观点', '核心判断'), ('events', '事件回顾', '触发事件'),
                  ('forecast', '盈利预测与估值', '预测与估值'), ('risks', '风险提示', '风险因素'),
                  ('disclaimer', '免责声明', '五项免责')),
     'body': BUSINESS_BODY, 'muted': '78909C', 'accent': NAVY, 'line_spacing': 1.3,
     'headings': BUSINESS_HEADINGS,
     'cover': cover_research, 'table': {'kind': 'horizontal', 'header_fill': 'EEF2F9'}, 'footer': 'page', 'header': True,
     'toc': False, 'title_size': 20},
)

THEME_VARIANTS = (
    ('t1', '极简蓝', '2563EB', {'size': 11, 'color': INK, 'east': 'PingFang SC'}),
    ('t2', '商务蓝', '1565C0', {'size': 9.5, 'color': '262626', 'east': '微软雅黑'}),
    ('t3', '学术黑', '000000', {'size': 12, 'color': INK, 'east': '宋体'}),
    ('t4', '政务蓝红', '003087', {'size': 14, 'color': INK, 'east': '仿宋_GB2312'}),
    ('t5', '创意橙', 'FF6B35', {'size': 12, 'color': '262626', 'east': '阿里巴巴普惠体'}),
)
for code, variant, primary, body in THEME_VARIANTS:
    GENRES = GENRES + (
        {'stem': f'business-report-{code}', 'label': f'商务报告·{variant}',
         'sections': (('views', '核心观点', '开门见山的判断'), ('events', '事件回顾', '本期大事'),
                      ('impact', '影响分析', '短中长期影响'), ('data', '数据附录', '关键指标'),
                      ('risks', '风险提示', '不确定性')),
         'body': body, 'muted': '78909C', 'accent': primary, 'line_spacing': 1.35,
         'headings': {1: {'size': 13, 'color': 'FFFFFF', 'east': body['east'], 'fill': primary},
                      2: {'size': 11, 'color': primary, 'east': body['east']},
                      3: {'size': 10, 'color': '37474F', 'east': body['east']}},
         'cover': cover_standard, 'table': {'kind': 'horizontal', 'header_fill': 'EEF2F9'}, 'footer': 'page',
         'header': True, 'toc': True, 'title_size': 22},)

TEMPLATE_INDEX = {item['stem']: item for item in GENRES}


# --------------------------------------------------------------------- build

def build_one(cfg):
    doc = Document()
    customize(doc, cfg)
    add_page_number(doc.sections[0], cfg['footer'])
    if cfg.get('header'):
        add_title_header(doc.sections[0], '{{title}}', cfg.get('muted', '6A706B'))
    cfg['cover'](doc, cfg)
    if cfg.get('toc'):
        page_break(doc)
        add_toc_block(doc, cfg['accent'])
    page_break(doc)
    for _, title, _ in cfg['sections']:
        chapter(doc, title, {**cfg['headings'][1], 'level': 1})
    body_sample(doc)
    add_sample_table(doc, **cfg['table'])

    blocks = list(doc.element.body)
    def index_of(predicate, label):
        for index, element in enumerate(blocks):
            if element.tag == qn('w:p') and predicate(''.join(t.text or '' for t in element.iter(qn('w:t')))):
                return index
        raise SystemExit(f"{cfg['stem']} 模板缺少锚点：{label}")

    sections = [
        {'section_id': sid, 'title': title, 'index': index_of(lambda t, l=title: t == l, title), 'purpose': purpose}
        for sid, title, purpose in cfg['sections']
    ]
    first_section = min(s['index'] for s in sections)
    keep_blocks = [index for index, element in enumerate(blocks)
                   if index < first_section and element.tag in (qn('w:p'), qn('w:tbl'))]
    spec = {
        'sections': sections,
        'keep_blocks': keep_blocks,
        'paragraph_index': index_of(lambda t: t.startswith('这是一段正文样式示例'), '正文样例'),
        'fields': [
            {'old': '{{title}}', 'field': 'title'},
            {'old': '{{organization}}', 'field': 'organization'},
            {'old': '{{period}}', 'field': 'period'},
            {'old': '{{report_date}}', 'field': 'report_date'},
        ],
    }
    if sorted(set(keep_blocks)) != keep_blocks or max(keep_blocks) >= first_section:
        raise SystemExit(f"{cfg['stem']} 封面块索引必须小于全部章节索引")
    doc.save(ASSETS / f"{cfg['stem']}.docx")
    (ASSETS / f"{cfg['stem']}.spec.json").write_text(json.dumps(spec, ensure_ascii=False, indent=1), encoding='utf-8')
    return {'stem': cfg['stem'], 'label': cfg['label'], 'sections': len(sections), 'blocks': len(blocks)}


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    for stale in ASSETS.glob('*.docx'):
        stale.unlink()
    for stale in ASSETS.glob('*.spec.json'):
        stale.unlink()
    summary = [build_one(cfg) for cfg in GENRES]
    print(json.dumps(summary, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    sys.exit(main())
