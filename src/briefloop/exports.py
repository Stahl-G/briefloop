"""Readable downloads from a saved version; never rewrite the original."""
from io import BytesIO
import json
import re
from markdown_it import MarkdownIt


def reader_markdown(store,brief):
    text=brief['markdown'];refs=json.loads(brief['detail']).get('citations',[]);used=[]
    def replace(match):
        sid=match.group(1).replace('\\','')
        if sid not in used:used.append(sid)
        return '['+str(used.index(sid)+1)+']'
    text=re.sub(r'\\?\[@(src\\?_[a-zA-Z0-9]+)\\?\]',replace,text)
    if used:
        text+='\n\n## 来源\n\n'
        for i,sid in enumerate(used,1):
            try:source=store.one('sources',sid)
            except ValueError:
                text+=f'{i}. 引用未关联到来源，请补充核对。\n';continue
            locators=list(dict.fromkeys(r.get('locator','') for r in refs if r['source_id']==sid and r.get('locator')))
            title=source['name']
            if source['url']:title=f'[{title}]({source["url"]})'
            text+=f'{i}. {title}'+(' · '+'；'.join(locators) if locators else '')+'\n'
    return text


def docx_bytes(markdown, *, report_profile="brief", title="", report_date="", organization="", period="", report_data=None, industry="", figures=None):
    from docx import Document
    from .industry_export import append_inline, configure_document, style_heading, style_table, append_data_chart
    industry_report = report_profile == 'industry_periodic'
    tokens=MarkdownIt('commonmark').enable('table').parse(markdown)
    levels=[int(t.tag[1]) for t in tokens if t.type=='heading_open']
    if industry_report and len(tokens)>=3 and tokens[0].type=='heading_open' and tokens[0].tag=='h1' and levels.count(1)==1 and 2 in levels:
        title=title or tokens[1].content
        tokens=tokens[3:]
    doc=Document()
    if industry_report: configure_document(doc, title=title, report_date=report_date, organization=organization, period=period, industry=industry)
    from docx.shared import Mm
    section=doc.sections[0]
    figure_width=section.page_width-section.left_margin-section.right_margin
    figure_height=min(Mm(180),section.page_height-section.top_margin-section.bottom_margin-Mm(25))
    explicit_figures=any(child.type=='image' and (child.attrGet('src') or '').startswith('briefloop-figure:')
                         for token in tokens for child in (token.children or []))
    paragraph=None;table=None;row=None;cell=None;list_type=None
    heading_count=0
    heading_levels=[int(t.tag[1]) for t in tokens if t.type=='heading_open']
    main_level=1 if heading_levels.count(1)>1 else 2
    for i,t in enumerate(tokens):
        if t.type=='heading_open':
            level=int(t.tag[1]); paragraph=doc.add_heading('',level=level)
            if industry_report and level==main_level:
                style_heading(paragraph, 1, first=heading_count==0); heading_count+=1
        elif t.type=='bullet_list_open':list_type='List Bullet'
        elif t.type=='ordered_list_open':list_type='List Number'
        elif t.type in ('bullet_list_close','ordered_list_close'):list_type=None
        elif t.type=='paragraph_open' and table is None:paragraph=doc.add_paragraph(style=list_type)
        elif t.type=='table_open':
            first_row=[]
            for n in tokens[i+1:]:
                if n.type=='tr_close':break
                if n.type in ('th_open','td_open'):first_row.append(n)
            table=doc.add_table(rows=0,cols=max(1,len(first_row)));table.style='Table Grid'
        elif t.type=='tr_open' and table is not None:row=table.add_row();cell=-1
        elif t.type in ('th_open','td_open') and row is not None:
            cell+=1;paragraph=row.cells[cell].paragraphs[0]
        elif t.type=='table_close':
            if industry_report: style_table(table)
            table=None;row=None;cell=None;paragraph=None
        elif t.type=='inline':
            if paragraph is None:paragraph=doc.add_paragraph()
            available_width=figure_width
            if row is not None and cell is not None and cell>=0:
                cell_width=row.cells[cell].width
                if cell_width:available_width=min(available_width,max(Mm(5),cell_width-Mm(4)))
            paragraph=append_inline(paragraph,t.children,figures=figures,max_figure_width=available_width,max_figure_height=figure_height)
        elif t.type in ('fence','code_block'):doc.add_paragraph(t.content)
    # Old callers retain the legacy chart; new callers pass a mapping, even {}.
    if industry_report and report_data and figures is None and not explicit_figures: append_data_chart(doc, report_data)
    buf=BytesIO();doc.save(buf);return buf.getvalue()
