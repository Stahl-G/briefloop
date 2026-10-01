"""Per-report market colors, applied only to explicit change data.

Direction is semantic, not a positive/negative-money rule. These helpers never
rewrite report text or recolor registered raster figures.
"""
import json
import re
from copy import deepcopy

_COLORS = {'up': 'D9363E', 'down': '1E8E4F', 'flat': '5F6675'}
_NUMBER = r'[+\-−]?\d+(?:,\d{3})*(?:\.\d+)?(?:\s*(?:%|％|个百分点|percentage points?|pp|bps?|基点))?'
_PREFIX = r'上涨|增长|增幅|上升|提高|增加|下跌|下降|下滑|减少|降低|跌幅|涨跌|变动|变化|同比|环比|\b(?:up|down|rose|rise|grew|growth|increase[ds]?|decrease[ds]?|decline[ds]?|fell|fall|change|delta|yoy|mom|qoq)\b'
_DOWN = re.compile(r'下跌|下降|下滑|减少|降低|跌幅|\b(?:down|decrease[ds]?|decline[ds]?|fell|fall)\b', re.I)
_CONTEXT = re.compile(_PREFIX, re.I)
_CHANGES = re.compile(r'(?P<prefix>' + _PREFIX + r')\s*(?:(?:by|of)\s+)?[:：]?\s*(?P<number>' + _NUMBER + r')', re.I)
_ARROWS = re.compile(r'(?P<arrow>[↑↓↗↘▲▼])\s*(?P<number>' + _NUMBER + r')', re.I)
_WHOLE_NUMBER = re.compile(r'\s*(' + _NUMBER + r')\s*', re.I)


def resolve_market(requirements=None):
    """Resolve legacy reports without changing their saved requirements."""
    from .models import report_language
    requirements = requirements or {}
    chosen = requirements.get('market_convention')
    return chosen if chosen in ('cn', 'intl') else ('intl' if report_language(requirements.get('language')) == 'en' else 'cn')


def market_colors(requirements=None):
    colors = dict(_COLORS)
    if resolve_market(requirements) == 'intl':
        colors['up'], colors['down'] = colors['down'], colors['up']
    return colors


def _direction(number, *, down=False):
    value = float(re.match(r'[+\-]?\d+(?:,\d{3})*(?:\.\d+)?', number.replace('−', '-')).group().replace(',', ''))
    return 'flat' if value == 0 else 'down' if down or value < 0 else 'up'


def semantic_delta_spans(text, *, context=''):
    """(start, end, direction) spans, with offsets into the unchanged text.

    Arrows + numbers and changes with an immediate directional/change label are
    eligible. A whole numeric cell is eligible when its header denotes a change.
    Bare signed amounts, margins, rates and percentages are deliberately ignored.
    """
    text = str(text)
    spans = []
    for match in _ARROWS.finditer(text):
        spans.append((match.start(), match.end(), _direction(match['number'], down=match['arrow'] in '↓↘▼')))
    for match in _CHANGES.finditer(text):
        start, end = match.span('number')
        if not any(start < other_end and end > other_start for other_start, other_end, _ in spans):
            spans.append((start, end, _direction(match['number'], down=bool(_DOWN.search(match['prefix'])))))
    numeric = _WHOLE_NUMBER.fullmatch(text)
    if not spans and numeric and _CONTEXT.search(str(context)):
        spans.append((*numeric.span(1), _direction(numeric[1], down=bool(_DOWN.search(str(context))))))
    return sorted(spans)


def save_report_settings(store, body):
    """Atomically change presentation only, without revalidating frozen research."""
    allowed = {'run_id', 'workspace_id', 'market_convention'}
    if not isinstance(body, dict) or set(body) - allowed:
        raise ValueError('报告设置只支持涨跌配色')
    if body.get('workspace_id') != store.meta('workspace_id'):
        raise ValueError('工作区已切换，请回到原工作区重新选择报告')
    if body.get('market_convention') not in ('cn', 'intl'):
        raise ValueError('涨跌配色只支持 cn 或 intl')
    from .store import dump
    with store.tx() as connection:
        row = connection.execute('SELECT requirements FROM runs WHERE id=?', (body.get('run_id'),)).fetchone()
        if row is None:
            raise ValueError('报告不存在')
        requirements = json.loads(row['requirements'])
        requirements['market_convention'] = body['market_convention']
        connection.execute('UPDATE runs SET requirements=? WHERE id=?', (dump(requirements), body['run_id']))
    return {'run_id': body['run_id'], 'market_convention': body['market_convention'], 'requirements': requirements}


def apply_docx_market_colors(document, requirements=None, *, protected_runs=None):
    """Color numeric/arrow runs in a Word document, preserving authored styles.

    Splitting copies run properties, so bold/italic/link relationships survive.
    Explicit user text colors take precedence. No cell shading is introduced.
    render_document may collect protected_runs to distinguish user colors/code
    from inherited template foregrounds. Without it, all direct colors are kept.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.run import Run
    from .ooxml_order import insert_ordered
    colors = market_colors(requirements)
    visited = set()

    def paragraph(value, context=''):
        if value._p in visited:
            return
        visited.add(value._p)
        runs = [Run(run, value) for run in value._p.iter(qn('w:r'))]
        text = ''.join(run.text for run in runs)
        spans = semantic_delta_spans(text, context=context)
        if context and spans and (_WHOLE_NUMBER.fullmatch(text) or _ARROWS.fullmatch(text.strip())) and value.alignment is None:
            # Table values only. A saved/user/template alignment is authoritative.
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            value.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        offset = 0
        for run in runs:
            text = run.text
            start, offset = offset, offset + len(text)
            hits = [(max(a, start)-start, min(b, offset)-start, direction)
                    for a, b, direction in spans if a < offset and b > start]
            # Do not disturb fields, images, breaks, tabs, or explicitly colored text.
            protected = (run._r in protected_runs if protected_runs is not None else
                         run._r.find('./' + qn('w:rPr') + '/' + qn('w:color')) is not None)
            if not hits or protected:
                continue
            if any(child.tag not in (qn('w:rPr'), qn('w:t')) for child in run._r):
                continue
            cuts = sorted({0, len(text), *(p for a, b, _ in hits for p in (a, b))})
            for a, b in zip(cuts, cuts[1:]):
                part = deepcopy(run._r)
                for child in list(part):
                    if child.tag != qn('w:rPr'):
                        part.remove(child)
                element = OxmlElement('w:t'); element.text = text[a:b]
                element.set(qn('xml:space'), 'preserve'); part.append(element)
                direction = next((d for left, right, d in hits if left <= a and b <= right), None)
                if direction:
                    properties = part.get_or_add_rPr()
                    color = OxmlElement('w:color'); color.set(qn('w:val'), colors[direction])
                    insert_ordered(properties, color)
                run._r.addprevious(part)
            run._r.getparent().remove(run._r)

    def table(value):
        headers = [cell.text for cell in value.rows[0].cells] if value.rows else []
        merged = bool(value._tbl.xpath('.//w:vMerge | .//w:gridSpan'))
        for index, row in enumerate(value.rows):
            for column, cell in enumerate(row.cells):
                context = headers[column] if not merged and index and column < len(headers) else ''
                for item in cell.paragraphs:
                    paragraph(item, context)
                for nested in cell.tables:
                    table(nested)

    for item in document.paragraphs:
        paragraph(item)
    for item in document.tables:
        table(item)
    return document


def chart_presentation(requirements=None):
    """Authoring metadata, not permission to recolor existing raster assets."""
    return {'market_convention': resolve_market(requirements),
            'direction_colors': {key: '#' + value for key, value in market_colors(requirements).items()},
            'series_colors': ['#2448B8', '#0B6A78', '#9A6B1F', '#6741D9', '#8A909C'],
            'instructions': '新图的涨跌色只用于变化数字、箭头和表示变化的系列，不用于底色，不按正负值给营收、余额、利润率等任意财务数值着色。普通图表系列按 series_colors 顺序，超过五组改分面或表格。已登记图像保持原样，换配色须依据原数据生成并登记新图，不把重绘称为原图复制。'}
