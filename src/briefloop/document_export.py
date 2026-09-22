"""Render saved rich content directly to Word, without a Markdown round trip."""
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from .document_model import normalize_document, table_layout, citation_label_text

ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
         'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}


def reader_source_blocks(document, sources):
    """Replace only a recognized internal source index, not authored data tables."""
    def text(node):
        return node.get('text', '') + ''.join(text(c) for c in node.get('content', []))
    blocks = document.get('content', [])
    kept, indexed = [], []
    for node in blocks:
        rows = node.get('content', [])
        headers = [text(c).strip() for c in rows[0].get('content', [])] if rows else []
        ids = []
        if (node.get('type') == 'table' and kept and kept[-1].get('type') == 'heading'
                and text(kept[-1]).strip() in ('来源', '参考来源', '参考资料')
                and headers in (['编号', '主体/来源', 'source_id'], ['编号', '主体来源', 'source_id'])
                and len(rows) > 1):
            for row in rows[1:]:
                cells = row.get('content', [])
                if len(cells) != 3 or not text(cells[0]).strip().isdigit():
                    break
                if any(len(c.get('content', [])) != 1 or c['content'][0].get('type') != 'paragraph'
                       or any(n.get('type') != 'text' for n in c['content'][0].get('content', []))
                       for c in cells):
                    break
                sid = text(cells[2]).strip()
                if sid not in sources or not sources[sid].get('name'):
                    break
                # Do not discard merged cells or additional structured content.
                if any(c.get('attrs', {}).get(k, 1) != 1 for c in cells for k in ('colspan', 'rowspan')):
                    break
                ids.append(sid)
            if len(ids) == len(rows) - 1:
                kept.pop()
                indexed.extend(ids)
                continue
        kept.append(node)
    return kept, indexed


def without_duplicate_cover_heading(document,title):
    """Drop one repeated cover H1, preserving every surrounding block and table."""
    from copy import deepcopy
    result=deepcopy(document)
    label=' '.join(title.split())
    if not label:return result
    for index,node in enumerate(result.get('content',[])):
        if node.get('type')!='heading':continue
        children=node.get('content',[])
        plain=all(child.get('type')=='text' and not any(mark.get('type')=='link' for mark in child.get('marks',[])) for child in children)
        text=' '.join(''.join(child.get('text','') for child in children).split())
        if node.get('attrs',{}).get('level')==1 and plain and text==label:
            del result['content'][index]
        break
    return result


def render_document(doc, document, *, figures=None, sources=None, append_sources=True, styles=None, citations=None):
    from .industry_export import append_figure
    document = normalize_document(document); figures = figures or {}; sources = sources or {}; styles = styles or {}
    used = []; section = doc.sections[0]
    max_width = section.page_width - section.left_margin - section.right_margin
    max_height = min(Mm(180), section.page_height - section.top_margin - section.bottom_margin - Mm(25))

    def inherit_properties(target,serialized):
        if not serialized:return
        from docx.oxml import parse_xml
        from copy import deepcopy
        for child in parse_xml(serialized):
            if target.find(child.tag) is None:target.append(deepcopy(child))

    existing_bookmarks = doc.element.xpath('.//w:bookmarkStart')
    bookmark_count = [max((int(node.get(qn('w:id'))) for node in existing_bookmarks
                           if (node.get(qn('w:id')) or '').isdigit()), default=0)]

    def anchor_heading(paragraph, name):
        # Word bookmark names must start with a letter and avoid spaces;
        # document_model's blockId ("block_<hex>") already qualifies.
        identifier = str(bookmark_count[0] + 1)
        start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), identifier); start.set(qn('w:name'), name)
        end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), identifier)
        paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
        paragraph._p.append(end)
        bookmark_count[0] += 1

    def cite(sid):
        if sid not in used: used.append(sid)
        return '[' + str(used.index(sid) + 1) + ']'

    def citation_run(paragraph, sid):
        label = cite(sid)
        run = paragraph.add_run(label)
        url = sources.get(sid, {}).get('url') or ''
        link = OxmlElement('w:hyperlink')
        if url.startswith(('https://', 'http://')):
            link.set(qn('r:id'), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        elif append_sources:
            link.set(qn('w:anchor'), 'briefloop_source_' + str(used.index(sid) + 1))
        else:
            return
        run.font.color.rgb = RGBColor.from_string('0563C1')
        run.font.underline = True
        link.append(run._r);paragraph._p.append(link)

    def inline(paragraph, nodes):
        for node in nodes:
            kind = node['type']
            if kind == 'hardBreak': paragraph.add_run().add_break(); continue
            if kind == 'citation': citation_run(paragraph, node['attrs']['sourceId']); continue
            if kind != 'text': raise ValueError('段落中包含不支持的内容')
            marks = {m['type']: m.get('attrs', {}) for m in node.get('marks', [])}
            href = marks.get('link', {}).get('href')
            reference = None
            if href and href.startswith('#source-'):
                reference = cite(href[8:])
                if not citation_label_text(node['text']):
                    citation_run(paragraph, href[8:]); continue
                # Text typed into a citation link is still authored content.
                # Keep it in full, including any numeric prefix we cannot infer.
                href = None
            run = paragraph.add_run(node['text'])
            if 'bold' in marks: run.bold = True
            if 'italic' in marks: run.italic = True
            if 'underline' in marks: run.underline = True
            if 'strike' in marks: run.font.strike = True
            if 'code' in marks: run.font.name = 'Consolas'
            if marks.get('textStyle', {}).get('color'):
                run.font.color.rgb = RGBColor.from_string(marks['textStyle']['color'][1:])
            if href:
                link = OxmlElement('w:hyperlink')
                link.set(qn('r:id'), paragraph.part.relate_to(href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                link.append(run._r); paragraph._p.append(link)
                run.font.underline = True
            if reference: citation_run(paragraph, marks['link']['href'][8:])

    def paragraph(container, style=None):
        # Word creates an initial empty paragraph inside a cell.
        if hasattr(container, '_tc') and len(container.paragraphs) == 1 and not container.paragraphs[0].text and not container.paragraphs[0]._p.xpath('.//w:drawing'):
            p = container.paragraphs[0]
            if style: p.style = style
            return p
        return container.add_paragraph(style=style)

    def block(node, container=doc, list_style=None, depth=0):
        kind = node['type']; attrs = node.get('attrs', {}); children = node.get('content', [])
        if kind in ('paragraph', 'heading', 'codeBlock'):
            style = styles.get('heading:'+attrs.get('blockId',''),styles.get('heading' + str(attrs.get('level', 2)), 'Heading ' + str(attrs.get('level', 2)))) if kind == 'heading' else list_style or styles.get('paragraph')
            p = paragraph(container, style)
            if attrs.get('textAlign'): p.alignment = ALIGN[attrs['textAlign']]
            if depth: p.paragraph_format.left_indent = Mm(depth * 5)
            inline(p, children)
            # Short standalone bold labels are commonly authored as paragraphs
            # before lists. Keep the label with its first item in Word.
            if (kind == 'paragraph' and not list_style and children
                    and all(c.get('type') == 'text' and any(m.get('type') == 'bold' for m in c.get('marks', [])) for c in children)
                    and 0 < len(p.text.strip()) <= 40):
                p.paragraph_format.keep_with_next = True
            if kind == 'heading' and attrs.get('blockId'): anchor_heading(p, attrs['blockId'])
            if kind == 'codeBlock':
                for run in p.runs: run.font.name = 'Consolas'; run.font.size = Pt(9)
        elif kind in ('bulletList', 'orderedList'):
            for child in children: block(child, container, 'List Number' if kind == 'orderedList' else 'List Bullet', depth)
        elif kind == 'listItem':
            for index, child in enumerate(children): block(child, container, list_style if index == 0 else None, depth + (1 if child['type'] in ('bulletList', 'orderedList') else 0))
        elif kind == 'blockquote':
            for child in children: block(child, container, None, depth + 1)
        elif kind == 'horizontalRule':
            p = paragraph(container); borders = OxmlElement('w:pBdr'); border = OxmlElement('w:bottom')
            for key, value in [('val', 'single'), ('sz', '4'), ('color', 'CCCCCC')]: border.set(qn('w:' + key), value)
            borders.append(border); p._p.get_or_add_pPr().append(borders)
        elif kind == 'image':
            fid = attrs['src'].split(':', 1)[1]
            if fid not in figures: raise ValueError('图表资源未登记到此版本：' + fid)
            figure = dict(figures[fid])
            if attrs.get('caption') is not None: figure['caption'] = attrs['caption']
            available = min(max_width, max(Mm(5), container.width - Mm(4))) if hasattr(container, '_tc') and container.width else max_width
            width = min(available, attrs['width'] * 9525) if attrs.get('width') else available
            height = min(max_height, attrs['height'] * 9525) if attrs.get('height') else max_height
            append_figure(paragraph(container), fid, figure, attrs.get('alt', ''), max_width=width, max_height=height)
        elif kind == 'table':
            nr, nc, cells = table_layout(node)
            table = container.add_table(rows=nr, cols=nc)
            table.style = styles.get('table', 'Table Grid')
            inherit_properties(table._tbl.tblPr,styles.get('table_properties'))
            default_widths=styles.get('table_widths',[])
            if len(default_widths)==nc and all(default_widths):
                ratios=default_widths
            else:
                # A template's sample may have fewer columns than the report.
                # Give descriptive cells room while keeping numeric columns compact.
                def text_length(value):
                    return len(value.get('text','')) + sum(text_length(child) for child in value.get('content',[]))
                ratios=[6]*nc
                for _,c,_,cs,value in cells:
                    if cs==1:ratios[c]=max(ratios[c],min(32,text_length(value)))
            for index,column in enumerate(table.columns):column.width=int(max_width*ratios[index]/sum(ratios))
            for r, c, rs, cs, value in cells:
                cell = table.cell(r, c)
                if rs > 1 or cs > 1: cell = cell.merge(table.cell(r + rs - 1, c + cs - 1))
                cell.text = ''
                ca = value.get('attrs', {})
                profile_name='table_header' if value['type']=='tableHeader' else 'table_body' if r%2 else 'table_alternate'
                profiles=styles.get(profile_name) or styles.get('table_body') or []
                profile=profiles[min(c,len(profiles)-1)] if profiles else {}
                if ca.get('backgroundColor'):
                    shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), ca['backgroundColor'][1:]); cell._tc.get_or_add_tcPr().append(shading)
                widths = ca.get('colwidth')
                cell.width = min(max_width, sum(widths) * 9525) if widths and all(widths) else sum(table.columns[i].width for i in range(c,c+cs))
                inherit_properties(cell._tc.get_or_add_tcPr(),profile.get('cell'))
                for child in value.get('content', []): block(child, cell)
                for p in cell.paragraphs:
                    if profile.get('paragraph_style') and p.style.name==styles.get('paragraph'):
                        p.style=profile['paragraph_style']
                    if ca.get('textAlign'): p.alignment = ALIGN[ca['textAlign']]
                    inherit_properties(p._p.get_or_add_pPr(),profile.get('paragraph'))
                    for run in p.runs:inherit_properties(run._r.get_or_add_rPr(),profile.get('run'))
                    if value['type'] == 'tableHeader':
                        for run in p.runs: run.bold = True
            if all(cell['type'] == 'tableHeader' for cell in children[0]['content']):
                table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                # Repeating a header does not stop WPS leaving its first copy
                # alone at a page bottom. Keep it with the first data row.
                if len(table.rows) > 1:
                    for cell in table.rows[0].cells:
                        for p in cell.paragraphs:
                            p.paragraph_format.keep_with_next = True
            for row in table.rows:
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        else: raise ValueError('不支持的导出内容：' + kind)

    blocks, indexed = reader_source_blocks(document, sources) if append_sources else (document.get('content', []), [])
    for node in blocks: block(node)
    for sid in indexed:
        if sid not in used: used.append(sid)
    if append_sources and used:
        doc.add_heading('来源', level=2)
        for i, sid in enumerate(used, 1):
            source = sources.get(sid, {})
            p = doc.add_paragraph(f'{i}. ')
            anchor_heading(p, 'briefloop_source_' + str(i))
            label = source.get('name') or '引用来源未关联，请补充核对'
            url = source.get('url') or ''
            if url.startswith(('https://', 'http://')):
                inline(p, [{'type': 'text', 'text': label, 'marks': [{'type': 'link', 'attrs': {'href': url}},
                        {'type': 'textStyle', 'attrs': {'color': '#0563C1'}}]}])
            else:
                p.add_run(label)
            # Locations belong to this saved version, not the global source.
            # Keep the title's hyperlink and the source paragraph's bookmark.
            locators=list(dict.fromkeys(ref.get('locator','') for ref in (citations or [])
                                        if ref.get('source_id')==sid and ref.get('locator')))
            if locators:p.add_run(' · '+'；'.join(locators))
    from .industry_export import populate_table_of_contents
    populate_table_of_contents(doc)
    return doc
