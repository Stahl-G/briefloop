"""Button-triggered, version-bound Excel workbook jobs. No model invocation.

The base artifact is pure openpyxl and needs nothing but the document itself;
OfficeCLI enhancement (SUM formulas, number formats, data bars) is applied on a
staging copy afterwards and any failure only degrades back to the base file.
"""
import hashlib
import json
import math
import os
import re
import shutil
from io import BytesIO

from .store import dump, now, uid
from .document_model import brief_document, table_layout

LAYOUTS = ('sheets', 'single')
XLSX_RENDERER_VERSION = 1
NO_TABLES_MESSAGE = '报告没有可导出的表格，无需生成 Excel'
INDEX_SHEET_TITLE = '目录'
# openpyxl rejects ':\\/?*[]' and control characters in sheet titles, silently
# renames duplicates and only warns past 31 characters while other applications
# then fail to open the file — normalize and dedupe before creating any sheet.
_SHEET_ILLEGAL = re.compile(r'[:：\\/?*\[\]\x00-\x1f]')
_SHEET_MAX = 31
_SHEET_BASE_MAX = 28  # room for a '-12' dedupe suffix while staying at 31
# Excel keeps 15 significant digits; anything longer would silently change the
# value, and leading zeros carry meaning of their own ('007' is a code).
_NUMERIC_TEXT = re.compile(r'^-?(?:\d{1,3}(?:,\d{3})*|\d+)(?:\.\d+)?$')
_PERCENT_TEXT = re.compile(r'^(-?(?:\d{1,3}(?:,\d{3})*|\d+)(?:\.\d+)?)%$')
_CURRENCY_PREFIXES = ('¥', '￥', '$', '€', '£')
_TOTAL_LABELS = {'合计', '总计', '小计', 'total', 'sum'}
_MAX_SIGNIFICANT_DIGITS = 15
_HEADER_FILL = 'E8EEF7'
_BORDER_COLOR = 'BFBFBF'
_TITLE_INDEX_MAX = 60


def export_xlsx_input(store, brief, layout_id, *, enhanced=None):
    """Frozen export identity. Unlike the Word identity it deliberately has no
    sources/figures projection: cell text flattening drops citation nodes, so
    neither can change the rendered bytes."""
    from . import office_cli
    requirements = json.loads(store.one('runs', brief['run_id'])['requirements'])
    document = brief_document(brief)
    if enhanced is None:
        enhanced = office_cli.enhancement_ready(store)
    return {'renderer': XLSX_RENDERER_VERSION, 'layout': layout_id, 'enhanced': bool(enhanced),
            'version_id': brief['id'], 'brief_hash': brief['hash'],
            'document': document, 'detail': json.loads(brief['detail']),
            'requirements': requirements}


def enqueue_export_xlsx(store, version_id, layout_id='sheets'):
    if layout_id not in LAYOUTS:
        raise ValueError('未知 Excel 版式')
    brief = store.one('briefs', version_id)
    identity = export_xlsx_input(store, brief, layout_id)
    if not table_titles(identity['document']):
        raise ValueError(NO_TABLES_MESSAGE)
    digest = hashlib.sha256(dump(identity).encode()).hexdigest()
    payload = {'version_id': version_id, 'run_id': brief['run_id'], 'fingerprint': digest,
               'layout': layout_id, 'enhanced': bool(identity['enhanced'])}
    # Admission is serialized by SQLite across request handlers and processes,
    # mirroring enqueue_export: lookup and insert share one transaction and the
    # jobs row is written directly, never through Store.enqueue's main-lane gate.
    with store.tx() as c:
        rows = c.execute("SELECT * FROM jobs WHERE kind='export_xlsx' "
                         "AND json_extract(payload,'$.fingerprint')=? "
                         "AND status IN ('queued','running','complete') ORDER BY rowid DESC", (digest,))
        for row in rows:
            job = dict(row)
            if job['status'] != 'complete':
                return job
            try:
                path = output_path_xlsx(store, job)
                if path.is_file() and hashlib.sha256(path.read_bytes()).hexdigest() == json.loads(job['result'] or '{}').get('sha256'):
                    return job
            except (OSError, ValueError):
                # A missing, unreadable or invalid cached artifact can be rebuilt.
                continue
        jid = uid('job')
        c.execute("INSERT INTO jobs VALUES(?,?,?,?,?,?,?,?)",
                  (jid, 'export_xlsx', 'queued', dump(payload), None, None, now(), now()))
    job = store.one('jobs', jid)
    from .task_notify import notify
    notify(store, job, 'queued')
    store.wake_jobs()
    return job


def output_path_xlsx(store, job):
    if job['kind'] != 'export_xlsx': raise ValueError('不是 Excel 文件任务')
    path = store.root / 'exports' / job['id'] / 'report.xlsx'
    if not path.resolve().is_relative_to((store.root / 'exports').resolve()): raise ValueError('无效导出路径')
    result = json.loads(job.get('result') or '{}')
    if result.get('path'):
        saved = (store.root / result['path']).resolve()
        if saved.parent != path.parent.resolve() or saved.suffix.lower() != '.xlsx':
            raise ValueError('无效导出结果路径')
        path = saved
    return path


def _inline_text(node):
    parts = []
    for child in node.get('content', []):
        kind = child.get('type')
        if kind == 'text': parts.append(child['text'])
        elif kind == 'hardBreak': parts.append('\n')
        # Citation nodes carry no printable cell text and are dropped here.
    return ''.join(parts)


def _block_text(node):
    kind = node.get('type')
    if kind in ('paragraph', 'heading'):
        return _inline_text(node)
    if kind in ('bulletList', 'orderedList', 'blockquote', 'listItem'):
        return '\n'.join(text for text in (_block_text(child) for child in node.get('content', [])) if text)
    if kind == 'codeBlock':
        return ''.join(child.get('text', '') for child in node.get('content', []))
    return ''  # images, rules and nested tables are not cell text


def _cell_text(cell):
    if cell is None: return ''  # shadow position inside a merged region
    return '\n'.join(text for text in (_block_text(child) for child in cell.get('content', [])) if text)


def _numeric_value(text):
    """Number for a fully compliant numeric cell text, else None (stays text).

    Thousands separators are accepted and dropped — the base artifact shows the
    plain General form; the enhanced file restores '#,##0'. Guard rails: no
    leading zeros ('0' and '0.x' pass, '007' does not) and at most 15
    significant digits, matching what Excel itself can store.
    """
    if not _NUMERIC_TEXT.fullmatch(text):
        return None
    body = text[1:] if text[:1] == '-' else text
    integer, _, fraction = body.partition('.')
    plain_integer = integer.replace(',', '')
    if len(plain_integer) > 1 and plain_integer[0] == '0':
        return None
    if len((plain_integer + fraction).lstrip('0')) > _MAX_SIGNIFICANT_DIGITS:
        return None
    plain = text.replace(',', '')
    return float(plain) if fraction else int(plain)


def _format_decimals(fraction):
    return ('.' + '0' * len(fraction)) if fraction else ''


def table_titles(document):
    """(表格标题, 表节点) in document order; the title is the nearest preceding
    heading's text, or '表N' (1-based table number) when there is none."""
    heading = None
    found = []
    for node in document.get('content', []):
        if node.get('type') == 'heading':
            heading = _inline_text(node).strip()
        elif node.get('type') == 'table':
            found.append((heading or f'表{len(found) + 1}', node))
    return found


def _sheet_base(raw):
    name = _SHEET_ILLEGAL.sub('-', raw.strip())[:_SHEET_BASE_MAX].strip()
    return name or '表'


def _dedupe_sheet_names(raw_names):
    issued = set()
    result = []
    for raw in raw_names:
        name, attempt = raw, 1
        while name in issued:
            suffix = f'-{attempt + 1}'
            name = raw[:_SHEET_MAX - len(suffix)] + suffix
            attempt += 1
        issued.add(name)
        result.append(name)
    return result


def _report_title(detail, requirements):
    return (detail.get('title') or requirements.get('title') or '').strip()


def _layout(document, layout_id, *, report_title=''):
    """Positioned table models shared by the renderer and the enhancement plan."""
    models = []
    for index, (title, node) in enumerate(table_titles(document)):
        nr, nc, cells = table_layout(node)
        rows = [[None] * nc for _ in range(nr)]
        spans = []
        for r, c, rs, cs, cell in cells:
            rows[r][c] = cell
            if rs > 1 or cs > 1:
                spans.append((r + 1, c + 1, r + rs, c + cs))
        header = bool(rows[0]) and all(cell.get('type') == 'tableHeader' for cell in rows[0])
        texts = [[_cell_text(cell) for cell in row] for row in rows]
        models.append({'title': title, 'rows': rows, 'texts': texts, 'width': nc, 'height': nr,
                       'spans': spans, 'header': header, 'index': index})
    if layout_id == 'single':
        sheet = _sheet_base(report_title)[: _SHEET_BASE_MAX] or '报表'
        cursor = 1
        for model in models:
            model['sheet'] = sheet
            model['title_row'] = cursor
            cursor += 1 + model['height'] + 1  # title + grid + one blank spacer row
        return models
    names = _dedupe_sheet_names([_sheet_base(INDEX_SHEET_TITLE)] +
                                [_sheet_base(model['title']) for model in models])
    for model, name in zip(models, names[1:]):
        model['sheet'] = name
        model['title_row'] = 1
    return models


def _grid_row(model, index):
    """Absolute 1-based sheet row of grid row `index` (0-based)."""
    return model['title_row'] + 1 + index


def _column_letter(index):
    from openpyxl.utils import get_column_letter
    return get_column_letter(index + 1)


def _cell_number(text):
    """Numeric value behind a cell text: plain, thousands, percent or currency.
    This is what an enhanced workbook can compute on; the base artifact keeps
    everything but the plain/thousands forms as text."""
    number = _numeric_value(text)
    if number is not None:
        return number
    found = _PERCENT_TEXT.fullmatch(text)
    if found:
        value = _numeric_value(found.group(1))
        return value / 100 if value is not None else None
    for prefix in _CURRENCY_PREFIXES:
        if text.startswith(prefix):
            return _numeric_value(text[len(prefix):])
    return None


def _cell_number_and_format(text):
    """Props only the enhanced artifact applies, keyed for officecli batch."""
    number = _numeric_value(text)
    if number is not None:
        fraction = text.replace(',', '').partition('.')[2]
        if ',' in text:
            return {'numberformat': '#,##0' + _format_decimals(fraction)}
        return None  # plain number: General already shows it exactly
    found = _PERCENT_TEXT.fullmatch(text)
    if found:
        value = _numeric_value(found.group(1))
        if value is not None:
            fraction = found.group(1).replace(',', '').partition('.')[2]
            return {'value': value / 100, 'numberformat': '0' + _format_decimals(fraction) + '%'}
    for prefix in _CURRENCY_PREFIXES:
        if text.startswith(prefix):
            value = _numeric_value(text[len(prefix):])
            if value is not None:
                rest = text[len(prefix):]
                fraction = rest.replace(',', '').partition('.')[2]
                grouping = '#,##0' if ',' in rest else '0'
                return {'value': value, 'numberformat': f'"{prefix}"' + grouping + _format_decimals(fraction)}
    return None


def plan_enhancements(tables):
    """Pure officecli batch plan for positioned table models (see _layout).

    Returns {'commands': batch items, 'formulas': expected-value records}. The
    executor turns commands into one `officecli batch` invocation and verifies
    every formula's cached value against the number the report actually shows.
    """
    commands, formulas = [], []
    for model in tables:
        props = {}
        height, header = model['height'], model['header']
        first_data = 1 if header else 0
        # Numeric-cell detection includes percent/currency totals: those become
        # real numbers in the enhanced workbook, so their column gets a formula.
        values = [[_cell_number(text) for text in row] for row in model['texts']]

        def is_total(index):
            label = model['texts'][index][0].strip().lower() if model['width'] else ''
            return (label in _TOTAL_LABELS
                    and sum(value is not None for value in values[index]) >= 2)

        totals = {index for index in range(height) if is_total(index)}
        # 合计行：每个数值格换成该列数据区的 SUM 公式，期望值=报告原数值。
        for index in sorted(totals):
            for column in range(model['width']):
                expected = values[index][column]
                if expected is None:
                    continue
                segments, start = [], None
                for row in range(first_data, index):
                    if row in totals:
                        if start is not None: segments.append((start, row - 1)); start = None
                    else:
                        start = start if start is not None else row
                if start is not None: segments.append((start, index - 1))
                if not segments:
                    continue
                letter = _column_letter(column)
                formula = 'SUM(' + ','.join(
                    (f'{letter}{_grid_row(model, top)}' if top == bottom
                     else f'{letter}{_grid_row(model, top)}:{letter}{_grid_row(model, bottom)}')
                    for top, bottom in segments) + ')'
                coordinate = f'{letter}{_grid_row(model, index)}'
                props[coordinate] = {**props.get(coordinate, {}), 'formula': formula}
                formulas.append({'sheet': model['sheet'], 'cell': coordinate,
                                 'formula': formula, 'expected': expected})
        # 显示格式：千分位/百分数/货币前缀在基础件里是纯文本或 General 数值。
        for index in range(height):
            for column, text in enumerate(model['texts'][index]):
                found = _cell_number_and_format(text)
                if found:
                    coordinate = f'{_column_letter(column)}{_grid_row(model, index)}'
                    existing = props.get(coordinate, {})
                    if 'formula' in existing:
                        # The formula computes this cell; setting a literal value
                        # too would overwrite it inside one officecli set command.
                        found = {key: value for key, value in found.items() if key != 'value'}
                    props[coordinate] = {**existing, **found}
        # 数据条：数据列（含合计行）至少 3 个数值格才值得可视化。
        for column in range(model['width']):
            letter = _column_letter(column)
            rows = [index for index in range(first_data, height)
                    if values[index][column] is not None]
            if len(rows) >= 3:
                commands.append({'command': 'add', 'parent': '/' + model['sheet'], 'type': 'databar',
                                 'props': {'ref': f'{letter}{_grid_row(model, rows[0])}:'
                                                  f'{letter}{_grid_row(model, rows[-1])}'}})
        # Value/format conversions first, formulas last: a SUM computed while
        # its inputs are still text caches 0, so percent/currency cells must be
        # numeric before any formula that spans them (verified against 1.0.152).
        for coordinate, cell_props in props.items():
            if 'formula' not in cell_props:
                commands.append({'command': 'set', 'path': f"/{model['sheet']}/{coordinate}",
                                 'props': cell_props})
        for coordinate, cell_props in props.items():
            if 'formula' in cell_props:
                commands.append({'command': 'set', 'path': f"/{model['sheet']}/{coordinate}",
                                 'props': cell_props})
    return {'commands': commands, 'formulas': formulas}


def _border():
    from openpyxl.styles import Border, Side
    thin = Side(border_style='thin', color=_BORDER_COLOR)
    return Border(left=thin, right=thin, top=thin, bottom=thin)


_BORDER = None


def _paint_cell(target, text, cell, *, header_row):
    from openpyxl.styles import Alignment, Font, PatternFill
    attrs = cell.get('attrs', {}) if cell else {}
    number = _numeric_value(text) if text else None
    if number is not None:
        target.value = number
    elif text:
        target.value = text
    target.border = _BORDER
    fill = attrs.get('backgroundColor')
    if fill:
        target.fill = PatternFill(start_color=fill[1:], end_color=fill[1:], fill_type='solid')
    elif header_row and cell is not None and cell.get('type') == 'tableHeader':
        target.fill = PatternFill(start_color=_HEADER_FILL, end_color=_HEADER_FILL, fill_type='solid')
    alignment = {'vertical': 'top'}
    if attrs.get('textAlign'): alignment['horizontal'] = attrs['textAlign']
    if '\n' in text: alignment['wrap_text'] = True
    target.alignment = Alignment(**alignment)
    if cell is not None and cell.get('type') == 'tableHeader':
        target.font = Font(bold=True)


def _paint_table(ws, model, top):
    """Paint the grid at sheet row `top`; the title row is painted by the caller."""
    from openpyxl.utils import get_column_letter
    first, header = top, model['header']
    for index, row in enumerate(model['rows']):
        for column in range(model['width']):
            coordinate = ws.cell(row=first + index, column=column + 1)
            _paint_cell(coordinate, model['texts'][index][column], row[column],
                        header_row=index == 0 and header)
    for r, c, r2, c2 in model['spans']:
        ws.merge_cells(start_row=first + r - 1, start_column=c, end_row=first + r2 - 1, end_column=c2)
    for row in model['rows']:
        for column, cell in enumerate(row):
            widths = (cell.get('attrs', {}) or {}).get('colwidth') if cell else None
            if widths and len(widths) == 1 and widths[0]:
                letter = get_column_letter(column + 1)
                ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0,
                                                         round(widths[0] / 7, 2))


def xlsx_bytes(document, detail, requirements, layout_id):
    """Deterministic openpyxl rendering; the base artifact needs no OfficeCLI."""
    from openpyxl import Workbook
    global _BORDER
    if _BORDER is None: _BORDER = _border()
    title = _report_title(detail, requirements)
    report_date = requirements.get('report_date') or ''
    models = _layout(document, layout_id, report_title=title)
    workbook = Workbook()
    workbook.remove(workbook.active)
    from openpyxl.styles import Font
    if layout_id == 'single':
        sheet = workbook.create_sheet(models[0]['sheet'] if models else _sheet_base(title) or '报表')
        for model in models:
            cell = sheet.cell(row=model['title_row'], column=1, value=model['title'])
            cell.font = Font(bold=True)
            _paint_table(sheet, model, model['title_row'] + 1)
    else:
        index = workbook.create_sheet(_dedupe_sheet_names([_sheet_base(INDEX_SHEET_TITLE)])[0])
        index.cell(row=1, column=1, value=title).font = Font(bold=True)
        if report_date:
            index.cell(row=2, column=1, value=report_date)
        for column, label in enumerate(('序号', '表格标题', '工作表'), 1):
            cell = index.cell(row=3, column=column, value=label)
            cell.font = Font(bold=True)
        for row, model in enumerate(models, 4):
            index.cell(row=row, column=1, value=row - 3)
            index.cell(row=row, column=2, value=model['title'][:_TITLE_INDEX_MAX])
            index.cell(row=row, column=3, value=model['sheet'])
        for model in models:
            sheet = workbook.create_sheet(model['sheet'])
            sheet.cell(row=1, column=1, value=model['title']).font = Font(bold=True)
            _paint_table(sheet, model, 2)
            # Freeze below the header row (title row 1 + header row 2 → A3); a
            # worksheet has one frozen pane, so 'single' freezes nothing at all.
            sheet.freeze_panes = f'A{model["title_row"] + (1 if model["header"] else 0) + 1}'
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _enhanced_mismatches(staging, formulas):
    """Read-only formula-cache check; never a reason to fail the job."""
    if not formulas:
        return []
    import warnings
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        # officecli writes data bars as an x14 extension; openpyxl warns on read.
        warnings.simplefilter('ignore')
        workbook = load_workbook(staging, data_only=True, read_only=True)
        try:
            mismatched = []
            for item in formulas:
                found = workbook[item['sheet']][item['cell']].value
                if not isinstance(found, (int, float)) or not math.isclose(
                        found, item['expected'], rel_tol=1e-9, abs_tol=1e-9):
                    mismatched.append({'sheet': item['sheet'], 'cell': item['cell'],
                                       'expected': item['expected'], 'found': found})
            return mismatched
        finally:
            workbook.close()


def generate_xlsx(store, job, cancelled):
    from . import office_cli
    payload = json.loads(job['payload']); brief = store.one('briefs', payload['version_id'])
    def stage(step, message):
        if cancelled.is_set(): raise InterruptedError('Excel 制作已停止')
        store.event(job['id'], 'export_progress', {'step': step, 'total': 4, 'message': message})
    stage(1, '读取已保存的报告版本')
    # Recompute with the enhanced flag frozen at enqueue time: a settings or
    # PATH drift after queueing is no reason to call the input "changed".
    identity = export_xlsx_input(store, brief, payload['layout'], enhanced=payload.get('enhanced'))
    if not table_titles(identity['document']):
        raise ValueError(NO_TABLES_MESSAGE)
    if hashlib.sha256(dump(identity).encode()).hexdigest() != payload['fingerprint']:
        raise ValueError('导出输入已变化，请对当前报告重新生成 Excel')
    detail, requirements, layout_id = identity['detail'], identity['requirements'], payload['layout']
    stage(2, '展开表格并生成工作簿')
    blob = xlsx_bytes(identity['document'], detail, requirements, layout_id)
    stage(3, '写入 Excel 文件' + ('并应用增强' if payload.get('enhanced') else ''))
    destination = output_path_xlsx(store, job); destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix('.tmp'); temporary.write_bytes(blob)
    if cancelled.is_set(): temporary.unlink(); raise InterruptedError('Excel 制作已停止')
    try:
        os.replace(temporary, destination)
    except PermissionError:
        if os.name != 'nt' or not destination.exists(): raise
        # Office/WPS may hold a deny-delete handle; keep this rendered version.
        from uuid import uuid4
        destination = destination.with_name('report-' + uuid4().hex + '.xlsx')
        os.replace(temporary, destination)
        store.event(job['id'], 'export_saved_as', {'path': str(destination.relative_to(store.root)),
                                                   'reason': '原文件被占用或不可替换，已另存本次 Excel'})
    enhanced, reason = bool(payload.get('enhanced')), None
    if enhanced:
        plan = plan_enhancements(_layout(identity['document'], layout_id,
                                         report_title=_report_title(detail, requirements)))
        staging = destination.with_name('report.staging.xlsx')
        try:
            if plan['commands']:
                shutil.copyfile(destination, staging)
                outcome = office_cli.enhance_workbook(store, staging, plan)
                if not outcome['applied']:
                    enhanced, reason = False, outcome['reason'] or 'OfficeCLI 增强未执行'
                else:
                    mismatched = _enhanced_mismatches(staging, plan['formulas'])
                    if mismatched:
                        enhanced, reason = False, '公式结果与报告数值不一致'
                        store.event(job['id'], 'export_xlsx_formula_mismatch', {'cells': mismatched})
                    else:
                        os.replace(staging, destination)  # enhanced bytes become the artifact
            else:
                enhanced, reason = False, None  # nothing to enhance in this report
        finally:
            if staging.exists(): staging.unlink()
        if not enhanced and reason and reason != '公式结果与报告数值不一致':
            store.event(job['id'], 'export_xlsx_enhancement_failed', {'reason': reason})
    stage(4, 'Excel 已生成，可以下载')
    data = destination.read_bytes()
    result = {'version_id': brief['id'], 'fingerprint': payload['fingerprint'],
              'path': str(destination.relative_to(store.root)), 'sha256': hashlib.sha256(data).hexdigest(),
              'download_url': '/api/export-file?job=' + job['id'],
              'layout': layout_id, 'enhanced': enhanced}
    if not enhanced and reason: result['enhance_reason'] = reason
    # Optional local quality gate after the atomic write: any failure is only a
    # recorded check result; this hook swallows everything itself as a second guard.
    try:
        office = office_cli.check_file(store, destination, job_id=job['id'],
                                       version_id=payload.get('version_id'))
    except Exception:
        office = None
    if office is not None:
        result['office'] = office
    return result
