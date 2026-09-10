"""Deterministic industry report styling; content and facts remain in the saved draft."""
from urllib.parse import urlsplit
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from .default_fonts import WESTERN_FONT

BLUE = '17466B'


def _element(parent, tag, **attributes):
    node = OxmlElement('w:' + tag)
    for key, value in attributes.items():
        node.set(qn('w:' + key), str(value))
    parent.append(node)
    return node


def _paragraph_after(paragraph, *, style=None):
    from docx.text.paragraph import Paragraph
    node=OxmlElement('w:p')
    paragraph._p.addnext(node)
    result=Paragraph(node,paragraph._parent)
    if style:result.style=style
    return result


def append_figure(paragraph, figure_id, figure, alt, *, max_width, max_height):
    """Use only already-authorized in-memory bytes, never an MD path or URL."""
    from io import BytesIO
    from PIL import Image
    payload=figure.get('image_bytes')
    if not isinstance(payload,(bytes,bytearray,memoryview)):
        raise ValueError('图表 '+figure_id+' 缺少已授权的图片数据')
    try:
        with Image.open(BytesIO(payload)) as image:
            width,height=image.size
            if image.format!='PNG' or width<=0 or height<=0:raise ValueError('expected PNG')
            image.verify()
    except Exception as exc:
        raise ValueError('图表 '+figure_id+' 的 PNG 数据无效') from exc
    if paragraph._p.xpath('.//w:t | .//w:drawing'):
        paragraph=_paragraph_after(paragraph)
    # 9525 EMU/px corresponds to 96 dpi; don't upscale a tiny image.
    scale=min(float(max_width)/width,float(max_height)/height,9525)
    paragraph.paragraph_format.line_spacing=1.0
    paragraph.paragraph_format.keep_with_next=True
    paragraph.paragraph_format.space_after=Pt(4)
    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    picture=paragraph.add_run().add_picture(BytesIO(payload),width=max(1,int(width*scale)),height=max(1,int(height*scale)))
    picture._inline.docPr.set('descr',alt or str(figure.get('title') or ''))
    labels=[]
    for value in (figure.get('title') or alt,figure.get('caption')):
        if value and str(value).strip() not in labels:labels.append(str(value).strip())
    sources=figure.get('source_labels') or []
    if isinstance(sources,str):sources=[sources]
    if sources:labels.append('来源：'+'；'.join(str(value) for value in sources if value))
    if labels:
        caption=_paragraph_after(paragraph,style='Caption')
        caption.paragraph_format.line_spacing=1.0
        caption.paragraph_format.keep_together=True
        caption.paragraph_format.space_after=Pt(8)
        caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run('\n'.join(labels))
        return caption
    paragraph.paragraph_format.keep_with_next=False
    return paragraph


def append_inline(paragraph, children, *, figures=None, max_figure_width=Mm(150), max_figure_height=Mm(180)):
    """Preserve normal inline content; registered figures occupy their MD position."""
    from docx.opc.constants import RELATIONSHIP_TYPE
    bold=italic=False;link=None;after_figure=False
    for child in children or []:
        if child.type=='strong_open':bold=True
        elif child.type=='strong_close':bold=False
        elif child.type=='em_open':italic=True
        elif child.type=='em_close':italic=False
        elif child.type=='link_open':
            if after_figure:
                paragraph=_paragraph_after(paragraph);after_figure=False
            href=child.attrGet('href') or ''
            if urlsplit(href).scheme.lower() in ('http','https','mailto'):
                link=OxmlElement('w:hyperlink')
                link.set(qn('r:id'),paragraph.part.relate_to(href,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
                paragraph._p.append(link)
        elif child.type=='link_close':link=None
        elif child.type=='image' and (child.attrGet('src') or '').startswith('briefloop-figure:'):
            figure_id=child.attrGet('src').split(':',1)[1]
            if figure_id not in (figures or {}):raise ValueError('图表 '+figure_id+' 未提供已授权的导出资源')
            paragraph=append_figure(paragraph,figure_id,figures[figure_id],child.content,
                                    max_width=max_figure_width,max_height=max_figure_height)
            after_figure=True;link=None
        elif child.type in ('text','code_inline','softbreak','hardbreak','image'):
            content='\n' if child.type in ('softbreak','hardbreak') else child.content
            if after_figure:
                if not content.strip():continue
                paragraph=_paragraph_after(paragraph);after_figure=False
            # Nonregistered local/remote images remain alt text; no fetch is done.
            run=paragraph.add_run(content);run.bold,run.italic=bold,italic
            if link is not None:
                run.font.color.rgb=RGBColor.from_string(BLUE);run.font.underline=True;link.append(run._r)
    return paragraph


def configure_document(doc, *, title='', report_date='', organization='', period='', industry=''):
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(25)
    section.header_distance = section.footer_distance = Mm(12.7)
    section.different_first_page_header_footer = True
    normal = doc.styles['Normal']
    normal.font.name = WESTERN_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string('2D2D2D')
    normal.paragraph_format.line_spacing = Pt(18)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 28), ('Subtitle', 16), ('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11)]:
        style = doc.styles[name]
        style.font.name = WESTERN_FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(16)
        style.paragraph_format.space_after = Pt(10)
    header = section.header.paragraphs[0]
    header.text = ' · '.join(filter(None, [organization, (industry + '行业定期报告') if industry else '行业定期报告', report_date]))
    header.style = doc.styles['Caption']
    header.runs[0].font.color.rgb = RGBColor.from_string(BLUE)
    borders = _element(header._p.get_or_add_pPr(), 'pBdr')
    _element(borders, 'bottom', val='single', sz=6, color=BLUE, space=6)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run((organization + '  ' if organization else '') + '第 ')
    _element(footer._p, 'fldSimple', instr='PAGE')
    footer.add_run(' 页')
    for run in footer.runs: run.font.size = Pt(9)
    cover = doc.add_paragraph(style='Title')
    cover.paragraph_format.space_before = Pt(100)
    cover.add_run(title or ((industry + '行业定期报告') if industry else '行业定期报告'))
    if organization: doc.add_paragraph(organization, style='Subtitle')
    if report_date: doc.add_paragraph(report_date)
    if period: doc.add_paragraph('覆盖期间：' + period)
    doc.add_page_break()


def style_heading(paragraph, level, *, first=False):
    if level == 1:
        paragraph.paragraph_format.page_break_before = not first
        border = _element(paragraph._p.get_or_add_pPr(), 'pBdr')
        _element(border, 'bottom', val='single', sz=5, color=BLUE, space=8)


def style_table(table):
    table.autofit = False
    count = len(table.columns)
    # First column names generally need more width than numeric period columns.
    widths = [160 / count] * count
    if count >= 4:
        widths = [42] + [118 / (count - 1)] * (count - 1)
    for index, column in enumerate(table.columns): column.width = Mm(widths[index])
    for row_index, row in enumerate(table.rows):
        props = row._tr.get_or_add_trPr()
        _element(props, 'cantSplit')
        if row_index == 0: _element(props, 'tblHeader')
        for index, cell in enumerate(row.cells):
            cell.width = Mm(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            _element(tcpr, 'shd', fill=BLUE if row_index == 0 else ('EBF2F8' if row_index % 2 else 'FFFFFF'))
            margins = _element(tcpr, 'tcMar')
            for edge in ('top', 'bottom'): _element(margins, edge, w=80, type='dxa')
            for edge in ('left', 'right'): _element(margins, edge, w=90, type='dxa')
            borders = _element(tcpr, 'tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'): _element(borders, edge, val='single', sz=4, color='D9D9D9')
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.0 if p._p.xpath('.//w:drawing') else Pt(14)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.keep_with_next = row_index == 0
                if index: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(9)
                    if row_index == 0:
                        r.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)


def append_data_chart(doc, report_data):
    """Draw one explicitly comparable forecast series; never fabricate missing periods."""
    from collections import defaultdict
    from datetime import date
    from decimal import Decimal, InvalidOperation
    from io import BytesIO
    import math
    groups = defaultdict(list)
    for record in (report_data or {}).get('records', []):
        if not isinstance(record, dict) or record.get('category') != 'forecast': continue
        try:
            as_of = date.fromisoformat(str(record['as_of']))
            timestamp = date.fromisoformat(str(record['current_date']))
            value = float(Decimal(str(record['current'])))
            if not math.isfinite(value): continue
        except (KeyError, ValueError, TypeError, InvalidOperation): continue
        key = tuple(str(record.get(k, '')) for k in ('metric', 'unit', 'region', 'product', 'tax_basis')) + (as_of.isoformat(),)
        groups[key].append((timestamp, value, record))
    series = []
    for key, values in groups.items():
        dates = [item[0] for item in values]
        if 2 <= len(values) <= 24 and len(set(dates)) == len(dates):
            series.append((key, sorted(values)))
    if not series: return
    from PIL import Image, ImageDraw
    key, points = max(series, key=lambda item: len(item[1]))
    image = Image.new('RGB', (1200, 540), 'white')
    draw = ImageDraw.Draw(image)
    values = [v for _, v, _ in points]
    bottom, top = min(0, min(values)), max(values)
    if top == bottom: top = bottom + 1
    if not math.isfinite(top - bottom): return
    padding = (top - bottom) * .12
    top += padding
    if not math.isfinite(top): return
    x0, x1, y0, y1 = 90, 1140, 40, 430
    for i in range(5):
        value = bottom + (top - bottom) * i / 4
        y = y1 - (y1 - y0) * i / 4
        draw.line([(x0, y), (x1, y)], fill='#DEE7ED', width=2)
        draw.text((12, y - 8), f'{value:g}', fill='#40505E', font_size=20)
    elapsed = (points[-1][0] - points[0][0]).days
    coordinates = [(x0 + (day - points[0][0]).days / elapsed * (x1 - x0), y1 - (value - bottom) / (top - bottom) * (y1 - y0)) for day, value, _ in points]
    draw.line(coordinates, fill='#17466B', width=4)
    stride = max(1, (len(points) + 5) // 6)
    for index, ((day, value, _), (x, y)) in enumerate(zip(points, coordinates)):
        draw.ellipse((x-5,y-5,x+5,y+5), fill='#17466B')
        if index % stride == 0 or index == len(points)-1:
            draw.text((x-45, 448), day.isoformat(), fill='#40505E', font_size=17)
            draw.text((x-20, y-28), f'{value:g}', fill='#17466B', font_size=19)
    doc.add_heading('数据图表', level=2)
    caption = doc.add_paragraph(' · '.join(filter(None, key[:-1])) + '（预测值，截至 ' + key[-1] + '）')
    caption.paragraph_format.keep_with_next = True
    output = BytesIO(); image.save(output, format='PNG'); output.seek(0)
    picture = doc.add_paragraph()
    picture.paragraph_format.line_spacing = 1.0
    picture.paragraph_format.keep_with_next = True
    picture.add_run().add_picture(output, width=Mm(160))
    source_ids, source_labels = [], []
    for _, _, record in points:
        sid = str(record.get('source_id', ''))
        if sid not in source_ids: source_ids.append(sid)
        label = str(record.get('source_label') or '').strip() or f'来源{source_ids.index(sid) + 1}'
        text = ' · '.join(filter(None, [label, str(record.get('locator') or '')]))
        if text not in source_labels: source_labels.append(text)
    doc.add_paragraph('数据依据：' + '；'.join(filter(None, source_labels)))
