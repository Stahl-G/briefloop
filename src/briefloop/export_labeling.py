"""Export-only AI provenance labels; never alter the saved report or assert review.

OOXML/PDF layouts follow TC260-PG-20258A §§6.1/6.3 (2025-08), published by
全国网络安全标准化技术委员会秘书处:
https://www.tc260.org.cn/tc260/sjzn/202508/f0a84ada0c11458e9cb6da56c1cb2426/files/760387fb30e94c058c31b8c2b2230020.pdf
The OOXML AIGC property contains the seven fields directly, without a wrapper.
"""
from io import BytesIO
import json
import math
from zipfile import ZipFile
from lxml import etree

LABELING_VERSION = 1
NOTICE_ZH = '本文件包含人工智能生成内容，请核实后使用'
NOTICE_EN = 'This file contains AI-generated content. Please verify before use.'
FIELDS = ('Label', 'ContentProducer', 'ProduceID', 'ReservedCode1',
          'ContentPropagator', 'PropagateID', 'ReservedCode2')
CUSTOM_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/custom-properties'
VT_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'
REL_NS = 'http://schemas.openxmlformats.org/package/2006/relationships'
CONTENT_NS = 'http://schemas.openxmlformats.org/package/2006/content-types'
CUSTOM_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties'
CUSTOM_TYPE = 'application/vnd.openxmlformats-officedocument.custom-properties+xml'


def brief_label(store, brief, *, language=None):
    """Positive provenance only: an agent-created version or known ancestor.

    Human edits do not erase known AI provenance. User-only/imported reports and
    unknown ancestry are not guessed to be AI. No edit-percentage classifier and
    no claim that a human reviewed the content is made here.
    """
    current, seen = brief, set()
    generated = False
    while current and current.get('id') not in seen:
        seen.add(current.get('id'))
        if current.get('author') == 'agent':
            generated = True
            break
        parent = current.get('parent_id')
        if not parent:
            break
        try:
            current = store.one('briefs', parent)
        except ValueError:
            break
        if current.get('run_id') != brief.get('run_id'):
            break
    if not generated:
        return None
    if language is None:
        language = json.loads(store.one('runs', brief['run_id'])['requirements']).get('language')
    from .models import report_language
    # BriefLoop is the actual local producing application name, not a purported
    # regulator-assigned provider code. The saved version is a real content ID.
    identifier = str(brief['id'])
    metadata = dict(zip(FIELDS, ('1', 'BriefLoop', identifier, '', 'BriefLoop', identifier, '')))
    return {'schema_version': LABELING_VERSION,
            'notice': NOTICE_EN if report_language(language) == 'en' else NOTICE_ZH,
            'metadata': metadata}


def metadata_json(label):
    values = label['metadata']
    if set(values) != set(FIELDS) or any(not isinstance(values[key], str) for key in FIELDS):
        raise ValueError('AI 标识元数据格式无效')
    if values['Label'] not in ('1', '2', '3'):
        raise ValueError('AI 标识属性无效')
    return json.dumps({key: values[key] for key in FIELDS}, ensure_ascii=False, separators=(',', ':'))


def export_properties(label=None, requirements=None):
    from .market_convention import resolve_market
    values = {'BriefLoopMarketConvention': resolve_market(requirements or {})}
    if label:
        values['AIGC'] = metadata_json(label)
    return values


def office_properties(data, values):
    """Replace named properties, preserving unrelated parts and custom fields.

    Register the custom part in both OPC relationship and content-type tables,
    use unique property IDs, and remove duplicate copies of managed properties.
    """
    if not values:
        return data
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    with ZipFile(BytesIO(data)) as original:
        parts = {entry.filename: original.read(entry) for entry in original.infolist()}
        root = etree.fromstring(parts['docProps/custom.xml'], parser) if 'docProps/custom.xml' in parts else etree.Element(
            '{'+CUSTOM_NS+'}Properties', nsmap={None: CUSTOM_NS, 'vt': VT_NS})
        for child in list(root):
            if child.get('name') in values:
                root.remove(child)
        used = {int(child.get('pid')) for child in root if str(child.get('pid', '')).isdigit()}
        pid = 2
        for name, value in values.items():
            while pid in used:
                pid += 1
            prop = etree.SubElement(root, '{'+CUSTOM_NS+'}property',
                                    fmtid='{D5CDD505-2E9C-101B-9397-08002B2CF9AE}', pid=str(pid), name=name)
            etree.SubElement(prop, '{'+VT_NS+'}lpwstr').text = str(value)
            used.add(pid)
        parts['docProps/custom.xml'] = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        rels = etree.fromstring(parts['_rels/.rels'], parser)
        existing = [item for item in rels if item.get('Type') == CUSTOM_REL]
        if existing:
            existing[0].set('Target', 'docProps/custom.xml')
            for item in existing[1:]:
                rels.remove(item)
        else:
            ids = {item.get('Id') for item in rels}
            index = 1
            while 'rId'+str(index) in ids:
                index += 1
            etree.SubElement(rels, '{'+REL_NS+'}Relationship', Id='rId'+str(index), Type=CUSTOM_REL, Target='docProps/custom.xml')
        parts['_rels/.rels'] = etree.tostring(rels, xml_declaration=True, encoding='UTF-8', standalone=True)
        types = etree.fromstring(parts['[Content_Types].xml'], parser)
        existing = [item for item in types if item.get('PartName') == '/docProps/custom.xml']
        if existing:
            existing[0].set('ContentType', CUSTOM_TYPE)
            for item in existing[1:]:
                types.remove(item)
        else:
            etree.SubElement(types, '{'+CONTENT_NS+'}Override', PartName='/docProps/custom.xml', ContentType=CUSTOM_TYPE)
        parts['[Content_Types].xml'] = etree.tostring(types, xml_declaration=True, encoding='UTF-8', standalone=True)
        output = BytesIO()
        with ZipFile(output, 'w') as target:
            written = set()
            for entry in original.infolist():
                if entry.filename in written:
                    continue
                target.writestr(entry, parts.pop(entry.filename))
                written.add(entry.filename)
            for name, body in parts.items():
                target.writestr(name, body)
    return output.getvalue()


def add_docx_notice(doc, label):
    if not label:
        return
    from docx.shared import Pt, RGBColor
    text = label['notice']
    def paint(paragraph, size):
        run = paragraph.add_run(text)
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string('5F6675')
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(0)
    if not any(p.text == text for p in doc.paragraphs):
        paint(doc.add_paragraph(), 9)
    seen = set()
    for section in doc.sections:
        # Mark every footer variant, including imported multi-section templates.
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            element = footer._element
            if element in seen:
                continue
            seen.add(element)
            if not any(p.text == text for p in footer.paragraphs):
                paint(footer.add_paragraph(), 8)


def add_xlsx_notice(workbook, label):
    if not label:
        return
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    text = label['notice']
    for sheet in workbook:
        row, last = sheet.max_row + 2, max(1, sheet.max_column)
        cell = sheet.cell(row, 1, text)
        cell.data_type = 's'
        cell.font = Font(size=9, color='5F6675')
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        if last > 1:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
        width = sum(sheet.column_dimensions[get_column_letter(col)].width or 13 for col in range(1, last+1))
        visual = sum(2 if ord(ch) > 127 else 1 for ch in text)
        sheet.row_dimensions[row].height = max(30, math.ceil(visual / max(1, width * .85)) * 13 + 6)
        for footer in (sheet.oddFooter, sheet.evenFooter, sheet.firstFooter):
            footer.center.text = text
            footer.center.size = 8
            footer.center.color = '5F6675'


def markdown_label(text, label):
    if not label:
        return text
    # TC260 §6.5: metadata in the Markdown file head, visible label at the end.
    values = json.loads(metadata_json(label))
    header = '---\nAIGC:\n' + ''.join('  '+key+': '+json.dumps(value, ensure_ascii=False)+'\n' for key, value in values.items()) + '---\n\n'
    return header + text.rstrip() + '\n\n' + label['notice'] + '\n'


def pdf_properties(data, label=None, requirements=None):
    """Preserve rendered pages and attach the format-specific PDF info metadata."""
    from pypdf import PdfReader, PdfWriter
    reader = PdfReader(BytesIO(data))
    writer = PdfWriter(clone_from=reader)
    writer.add_metadata({'/'+key: value for key, value in export_properties(label, requirements).items()})
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def read_office_properties(data):
    with ZipFile(BytesIO(data)) as archive:
        if 'docProps/custom.xml' not in archive.namelist(): return {}
        root = etree.fromstring(archive.read('docProps/custom.xml'), etree.XMLParser(resolve_entities=False, no_network=True))
        return {item.get('name'): ''.join(item.itertext()) for item in root}
