"""Immutable evidence locations and version-scoped claim bindings.

Location validation proves what was saved/read, never semantic support or truth.
"""
import hashlib
import json
from copy import deepcopy
from typing import Literal
from pydantic import Field
from .models import Model
from .store import dump, uid, now

SCHEMA = '''
CREATE TABLE IF NOT EXISTS evidence_spans(id TEXT PRIMARY KEY,source_id TEXT NOT NULL REFERENCES sources(id),
 source_hash TEXT NOT NULL,data TEXT NOT NULL,created TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS claims(id TEXT PRIMARY KEY,run_id TEXT NOT NULL REFERENCES runs(id),
 previous_id TEXT REFERENCES claims(id),data TEXT NOT NULL,created TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS claim_bindings(id TEXT PRIMARY KEY,version_id TEXT NOT NULL REFERENCES briefs(id),
 claim_id TEXT NOT NULL REFERENCES claims(id),block_id TEXT NOT NULL,quote TEXT NOT NULL,
 block_hash TEXT NOT NULL,created TEXT NOT NULL);
'''


class Locator(Model):
    kind: Literal['text','pdf','xlsx','image']
    start_line: int | None = Field(default=None,ge=1)
    end_line: int | None = Field(default=None,ge=1)
    page: int | None = Field(default=None,ge=1)
    sheet: str | None = None
    cells: str | None = None
    region: list[float] | None = None


class EvidenceInput(Model):
    source_id: str
    locator: Locator
    excerpt: str = ''
    entity: str = ''
    metric: str = ''
    value: str | None = None
    unit: str = ''
    period: str = ''
    category: Literal['actual','plan','forecast','opinion','unknown'] = 'unknown'


class Support(Model):
    span_id: str
    supports_quote: str = Field(min_length=1)
    rationale: str = ''


class ClaimInput(Model):
    statement: str = Field(min_length=1,max_length=6000)
    kind: Literal['fact','source_opinion','calculation','inference','recommendation']
    importance: Literal['core','supporting'] = 'core'
    entity: str = ''
    metric: str = ''
    period: str = ''
    scope: str = ''
    requirement_ids: list[str] = Field(default_factory=list)
    supports: list[Support] = Field(default_factory=list)
    premise_claim_ids: list[str] = Field(default_factory=list)
    reasoning: str = ''
    assumptions: list[str] = Field(default_factory=list)
    figure_ids: list[str] = Field(default_factory=list)


def digest(value):return hashlib.sha256(value.encode()).hexdigest()


def record(store,table,identity):
    if table not in ('evidence_spans','claims'):raise ValueError('无效证据记录类型')
    rows=store.rows('SELECT * FROM '+table+' WHERE id=?',(identity,))
    if not rows:raise ValueError('证据或主张记录不存在')
    return {**rows[0],'data':json.loads(rows[0]['data'])}


def _read_location(store,value):
    from .media import source_files
    loc=value.locator
    source,provenance,original=source_files(store,value.source_id)
    text=store.source_text(value.source_id)
    raw_hash=hashlib.sha256(original.read_bytes()).hexdigest() if original else None
    method='saved_text';located='';status='located';payload={}
    if loc.kind=='text':
        lines=text.splitlines()
        if loc.start_line is None or loc.end_line is None or not 1<=loc.start_line<=loc.end_line<=len(lines) or loc.end_line-loc.start_line>200:
            raise ValueError('证据行范围超出已保存来源')
        located='\n'.join(lines[loc.start_line-1:loc.end_line])
    elif loc.kind=='pdf':
        if not original or original.suffix.lower()!='.pdf' or loc.page is None:raise ValueError('PDF 证据需要原件及页码')
        from pypdf import PdfReader
        pdf=PdfReader(original)
        if loc.page>len(pdf.pages):raise ValueError('PDF 页码超出原件')
        located=pdf.pages[loc.page-1].extract_text() or '';method='pdf_page_text'
        if not located.strip():status='visual_review_needed';method='pdf_page_visual'
    elif loc.kind=='xlsx':
        if not original or original.suffix.lower() not in ('.xlsx','.xlsm') or not loc.sheet or not loc.cells:
            raise ValueError('表格证据需要 XLSX 原件、工作表和单元格范围')
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries, get_column_letter
        left,top,right,bottom=range_boundaries(loc.cells)
        if not all((left,top,right,bottom)) or left>right or top>bottom or (right-left+1)*(bottom-top+1)>200:
            raise ValueError('请将证据范围限定在 200 个单元格内')
        formula=load_workbook(original,data_only=False,read_only=True);cached=load_workbook(original,data_only=True,read_only=True)
        try:
            if loc.sheet not in formula.sheetnames:raise ValueError('工作表不存在')
            cells=[]
            formula_rows=formula[loc.sheet].iter_rows(min_row=top,max_row=bottom,min_col=left,max_col=right)
            cached_rows=cached[loc.sheet].iter_rows(min_row=top,max_row=bottom,min_col=left,max_col=right)
            for row_number,(row,cache_row) in enumerate(zip(formula_rows,cached_rows),top):
                for column_number,(cell,cached_cell) in enumerate(zip(row,cache_row),left):
                    coordinate=f'{get_column_letter(column_number)}{row_number}'
                    raw=cell.value;cache=cached_cell.value
                    cells.append({'cell':coordinate,'formula':raw if cell.data_type=='f' else None,
                                  'value':str(cache) if cache is not None else None,'number_format':cell.number_format})
            payload['cells']=cells;located='\n'.join(f"{x['cell']}: {x['value']}" for x in cells)
            method='saved_workbook_cells'
            if any(x['formula'] and x['value'] is None for x in cells):status='missing_formula_cache'
        finally:formula.close();cached.close()
    else:
        if not original:raise ValueError('图像证据需要已保留原件')
        from PIL import Image
        with Image.open(original) as image:
            width,height=image.size
        if loc.region:
            if len(loc.region)!=4:raise ValueError('图像区域需要 x/y/宽/高')
            x,y,w,h=loc.region
            if min(x,y)<0 or min(w,h)<=0 or x+w>width or y+h>height:raise ValueError('图像区域超出原件')
        payload['image_size']=[width,height];method='image_visual';status='visual_review_needed'
    if value.excerpt and status not in ('visual_review_needed',) and value.excerpt not in located:
        raise ValueError('证据摘录不在指定位置，不能使用全文其他位置替代')
    return source,{'locator':loc.model_dump(exclude_none=True),'excerpt':value.excerpt or located,
                   'located_text':located,'raw_hash':raw_hash,'extraction_method':method,
                   'location_status':status,**payload}


def create_span(store,request):
    value=EvidenceInput.model_validate(request);source,location=_read_location(store,value)
    data={**value.model_dump(exclude={'source_id','locator','excerpt'}),**location,
          'excerpt_hash':digest(location['excerpt']),'support_status':'unreviewed'}
    eid='span_'+digest(dump({'source_id':source['id'],'source_hash':source['hash'],'data':data}))[:24]
    with store.tx() as c:c.execute('INSERT OR IGNORE INTO evidence_spans VALUES(?,?,?,?,?)',(eid,source['id'],source['hash'],dump(data),now()))
    return record(store,'evidence_spans',eid)


def create_claim(store,run_id,request,previous_id=None):
    value=ClaimInput.model_validate(request);store.one('runs',run_id)
    allowed=set(store.source_ids(run_id));references=set(json.loads(store.one('runs',run_id)['requirements']).get('reference_source_ids',[]))
    from .deliverable_spec import requirement_items
    requirement_ids={x['requirement_id'] for x in requirement_items(json.loads(store.one('runs',run_id)['requirements']))}
    if not set(value.requirement_ids).issubset(requirement_ids):raise ValueError('主张关联了未登记的报告要求')
    if previous_id and record(store,'claims',previous_id)['run_id']!=run_id:raise ValueError('主张修订属于另一报告')
    for support in value.supports:
        evidence=record(store,'evidence_spans',support.span_id)
        if evidence['source_id'] not in allowed-references:raise ValueError('主张证据未登记到本轮报告')
        if support.supports_quote not in value.statement:raise ValueError('支持范围必须明确对应本条主张中的片段')
    for identity in value.premise_claim_ids:
        if record(store,'claims',identity)['run_id']!=run_id:raise ValueError('推断前提属于另一报告')
    if value.kind in ('inference','recommendation','calculation') and not value.reasoning.strip():
        raise ValueError('计算、推断与建议必须记录依据和推理说明')
    from .figures import read_figure
    figures=[]
    for fid in value.figure_ids:
        figure=read_figure(store,fid)
        if not set(figure['source_ids']).issubset(allowed-references):raise ValueError('图表依据不属于本轮报告')
        figures.append({'figure_id':fid,'manifest_hash':digest(dump(figure))})
    identity=uid('claim');data={**value.model_dump(),'figures':figures,'review_status':'unreviewed'}
    with store.tx() as c:c.execute('INSERT INTO claims VALUES(?,?,?,?,?)',(identity,run_id,previous_id,dump(data),now()))
    return record(store,'claims',identity)


def node_text(node):
    if node.get('type')=='text':return node.get('text','')
    if node.get('type')=='image':return node.get('attrs',{}).get('caption') or node.get('attrs',{}).get('alt','')
    return ''.join(node_text(child) for child in node.get('content',[]))


def node_signature(node):
    """Ignore typography but retain structure, values, source and image identity."""
    attrs={k:v for k,v in node.get('attrs',{}).items() if k in ('src','caption','sourceId','colspan','rowspan')}
    links=[m.get('attrs',{}).get('href') for m in node.get('marks',[]) if m.get('type')=='link']
    return digest(dump({'type':node['type'],'text':node.get('text'),'attrs':attrs,**({'links':links} if links else {}),
                        'content':[node_signature(x) for x in node.get('content',[])]}))


def blocks(document):
    out={}
    def walk(node):
        identity=node.get('attrs',{}).get('blockId')
        if identity:
            if identity in out:raise ValueError('文档含重复正文块 ID，需重新定位')
            out[identity]=node
        for child in node.get('content',[]):walk(child)
    walk(document);return out


def bind_claim(store,version_id,claim_id,block_id,quote):
    from .document_model import brief_document
    brief=store.one('briefs',version_id);claim=record(store,'claims',claim_id)
    if claim['run_id']!=brief['run_id']:raise ValueError('主张不属于该报告')
    node=blocks(brief_document(brief)).get(block_id)
    if node is None or not isinstance(quote,str) or not quote or node_text(node).count(quote)!=1:
        raise ValueError('正文锚点缺失或不唯一，请指定原文所在块，不能全文模糊匹配')
    identity=uid('binding')
    with store.tx() as c:c.execute('INSERT INTO claim_bindings VALUES(?,?,?,?,?,?,?)',
        (identity,version_id,claim_id,block_id,quote,node_signature(node),now()))
    return {'id':identity,'version_id':version_id,'claim_id':claim_id,'status':'unreviewed'}


def claim_closure(store,claim_id,trail=()):
    if claim_id in trail:return {'claim_id':claim_id,'status':'premise_cycle','evidence':[],'premises':[]}
    claim=record(store,'claims',claim_id);evidence=[];state='unreviewed'
    for support in claim['data']['supports']:
        span=record(store,'evidence_spans',support['span_id'])
        try:
            current,location=_read_location(store,EvidenceInput(source_id=span['source_id'],locator=span['data']['locator'],excerpt=span['data']['excerpt']))
            intact=current['hash']==span['source_hash'] and location['raw_hash']==span['data']['raw_hash']
        except (ValueError,OSError):intact=False
        if not intact:state='source_changed'
        evidence.append({**span,'source_name':store.one('sources',span['source_id'])['name'],'intact':intact,
                         'supports_quote':support['supports_quote'],'rationale':support['rationale']})
    premises=[claim_closure(store,identity,(*trail,claim_id)) for identity in claim['data'].get('premise_claim_ids',[])]
    if any(p['status']!='unreviewed' for p in premises):state='premise_changed'
    return {'claim_id':claim_id,'claim':claim,'status':state,'evidence':evidence,'premises':premises}


def inspect_bindings(store,version_id):
    from .document_model import brief_document
    brief=store.one('briefs',version_id);nodes=blocks(brief_document(brief));versions=[];cursor=brief
    while cursor:
        versions.append(cursor['id']);cursor=store.one('briefs',cursor['parent_id']) if cursor['parent_id'] else None
    candidates=[];seen=set()
    for vid in versions:
        for binding in store.rows('SELECT * FROM claim_bindings WHERE version_id=? ORDER BY rowid DESC',(vid,)):
            key=(binding['claim_id'],binding['block_id'])
            if key in seen:continue
            seen.add(key);candidates.append(binding)
    superseded=set()
    for binding in candidates:
        claim=record(store,'claims',binding['claim_id']);previous=claim['previous_id'];visited=set()
        while previous and previous not in visited:
            visited.add(previous);superseded.add((previous,binding['block_id']))
            previous=record(store,'claims',previous)['previous_id']
    result=[]
    for binding in candidates:
        if (binding['claim_id'],binding['block_id']) in superseded:continue
        closure=claim_closure(store,binding['claim_id']);node=nodes.get(binding['block_id']);state=closure['status']
        if node is None:state='anchor_missing'
        elif node_signature(node)!=binding['block_hash'] or node_text(node).count(binding['quote'])!=1:state='needs_review'
        if closure['status']!='unreviewed':state=closure['status']
        result.append({**binding,'target_version':version_id,'inherited':binding['version_id']!=version_id,
                       **closure,'status':state})
    return {'version_id':version_id,'bindings':result,'coverage_status':'not_reviewed',
            'note':'已登记绑定不代表完整覆盖；重要主张遗漏及语义支持仍待独立审阅。'}
