"""Worksheet selection and cache binding; no simulated model judgments."""
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import hashlib
import json

from openpyxl import Workbook
from PIL import Image
import pytest

from briefloop import office_cli
from briefloop.store import Store


def _workbook(path):
    workbook = Workbook()
    workbook.active.title = '目录'
    workbook.active.append(['目录中的标题'])
    hidden = workbook.create_sheet('隐藏表')
    hidden.sheet_state = 'hidden'
    detail = workbook.create_sheet('中文! "明细"')
    detail.append(['品名', '编号'])
    detail.append(['甲', '007'])
    workbook.save(path)


def _renderer(monkeypatch):
    calls = []
    images = []
    for color in ('red', 'blue'):
        buffer = BytesIO()
        Image.new('RGB', (12, 6), color).save(buffer, format='PNG')
        images.append(buffer.getvalue())

    def run(args, **kwargs):
        assert '--page' not in args
        region = args[args.index('--range') + 1]
        calls.append(region)
        Path(args[args.index('-o') + 1]).write_bytes(images[0 if region.startswith('/目录/') else 1])
        return {'ok': True, 'data': {}, 'reason': None}

    monkeypatch.setattr(office_cli, 'find', lambda: 'synthetic-renderer')
    monkeypatch.setattr(office_cli, 'version', lambda *_: 'synthetic')
    monkeypatch.setattr(office_cli, 'run_json', run)
    return calls, images


def test_xlsx_preview_selects_visible_sheets_and_replaces_legacy_cache(tmp_path, monkeypatch):
    target = tmp_path / 'two.xlsx'
    _workbook(target)
    frozen = target.read_bytes()
    store = SimpleNamespace(root=tmp_path / 'workspace')
    calls, images = _renderer(monkeypatch)
    digest = hashlib.sha256(frozen).hexdigest()
    directory = office_cli._render_directory(store, digest)
    legacy = directory / 'page-0002.png'
    legacy.write_bytes(images[0])  # The old --page 2 image showed the index.
    record = directory / 'page-0002.json'
    record.write_text(json.dumps({'source_sha256': digest, 'page': 2,
                                 'image_sha256': hashlib.sha256(images[0]).hexdigest()}))
    with pytest.raises(ValueError, match='绑定不一致'):
        office_cli.office_image(store, digest, 2)
    first = office_cli.render_page(store, target, 1)
    second = office_cli.render_page(store, target, 2)
    assert calls == ['/目录/A1:A1', '/中文! \\"明细\\"/A1:B2']
    assert first['worksheet']['name'] == '目录'
    assert second['worksheet'] == {'name': '中文! "明细"', 'cells': 'A1:B2',
                                   'range': calls[1], 'count': 2}
    assert first['image_sha256'] != second['image_sha256']
    assert office_cli.office_image(store, digest, 2) == images[1]
    assert office_cli.render_page(store, target, 2)['cached'] is True and len(calls) == 2
    # A new-format image still cannot be reused for another worksheet.
    metadata = json.loads(record.read_text())
    metadata['worksheet'] = first['worksheet']
    record.write_text(json.dumps(metadata))
    assert office_cli.render_page(store, target, 2)['cached'] is False and len(calls) == 3
    with pytest.raises(ValueError, match='2 张可见工作表'):
        office_cli.render_page(store, target, 3)
    assert len(calls) == 3 and target.read_bytes() == frozen


def test_xlsx_preview_returns_worksheet_names_and_ranges(tmp_path, monkeypatch):
    store = Store(tmp_path / 'workspace')
    store.update_settings({'officecli_enabled': True})
    target = tmp_path / 'two.xlsx'
    _workbook(target)
    from briefloop.sources import upload
    source = upload(store, 'two.xlsx', target.read_bytes())
    calls, _ = _renderer(monkeypatch)
    result = office_cli.render_preview(store, {'source_id': source['id'], 'pages': [1, 2]})
    assert [page['worksheet']['name'] for page in result['pages']] == ['目录', '中文! "明细"']
    assert [page['worksheet']['cells'] for page in result['pages']] == ['A1:A1', 'A1:B2']
    assert result['cached'] is False and len(calls) == 2
    assert office_cli.render_preview(store, {'source_id': source['id'], 'pages': [2]})['cached'] is True
