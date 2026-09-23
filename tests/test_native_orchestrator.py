"""Exercise the production main-agent tools with a scripted, unbilled engine."""
import json
from pathlib import Path
import queue
import threading
import time
from types import SimpleNamespace

import pytest

from briefloop.store import Store, dump
from briefloop.native_harness import NativeHarness
from briefloop.native_roles import run_tool
from briefloop.interactive_runtime import InteractiveRuntime
from briefloop.runtime import Worker
from briefloop.deliverable_spec import resolve, reader_contract_schema


def contract(store, rid):
    spec = resolve(json.loads(store.one('runs', rid)['requirements']))
    return {'source_fingerprint':reader_contract_schema(spec)['properties']['source_fingerprint']['const'],
        'clauses':[{'requirement_id':r['requirement_id'], 'kind':{'manual':'manual_assignment','writing':'writing_preference'}.get(r['kind'],'reader_content'),
                    'source_quote':r['text'], 'instruction':r['text']} for r in spec['requirement_items']]}


class FlowEngine:
    def __init__(self, store, run, source):
        self.process = object();self.sinks = {};self.sessions = {};self.queues = {};self.calls = []
        self.store, self.run, self.source = store, run, source
        self.scouts = threading.Barrier(2)
    def subscribe(self, eid):
        self.sinks[eid] = queue.Queue();return self.sinks[eid]
    def unsubscribe(self, eid):
        self.sinks.pop(eid, None)
    def call(self, method, p, **kw):
        self.calls.append((method,p))
        sid = p.get('session_id')
        if method == 'session_create':
            self.sessions[sid] = p
            return {'session_file':str(Path(p['session_dir'])/'transcript.jsonl'), 'model':p['model']}
        if method == 'turn_start':
            self.sessions[sid]['eid'] = p['execution_id']
            role = self.sessions[sid]['role']
            if role == 'orchestrator':
                data=json.loads((Path(self.sessions[sid]['packet_root'])/'input.json').read_text(encoding='utf-8'))
                round_tools = [
                    ('save_research_handoff', {'handoff':{'learnings':[],'follow_ups':[],'covered':['Synthetic source'],'open_questions':[]}}),
                    ('workspace_action', {'request':{'action':'finish_research_round','gaps':[],'summary':'Synthetic round complete'}})
                ] if data.get('research_plan') else []
                self.queues[sid] = [
                    ('save_plan', {'plan':{'summary':'Explain revenue', 'reader_contract':contract(self.store,self.run['id'])}}),
                    ('run_scouts', {'tasks':[{'slot_id':'scout-1','assignment':'Read revenue'}, {'slot_id':'scout-2','assignment':'Read scope'}]}),
                    *round_tools,
                    ('workspace_action', {'request':{'action':'reconciliation_save','reconciliation':{'status':'not_applicable','examined_claim_ids':[],'unexamined_claim_ids':[],'relations':[],'open_questions':[],'coverage_notes':'Synthetic fixture has no source claims'}}}),
                    ('write_report', {'instructions':'State revenue and keep the period'}), ('finish_task', {})]
            elif role == 'scout':
                # Neither Scout can submit until both child sessions are live.
                self.queues[sid] = [
                    ('source_read', {'source_id':self.source['id']}),
                    ('record_evidence', {'items':[{'id':'revenue','source_id':self.source['id'],
                        'source_hash':self.source['hash'],'locator':'line 1','quote':'2025年收入1200万元。',
                        'facts':['2025年收入1200万元。'],'coverage_status':'complete'}]}),
                    ('submit_scout_result', {'gaps':['No further synthetic material']})]
                threading.Thread(target=lambda:(self.scouts.wait(4),self.next(sid)),daemon=True).start()
                return {}
            elif role == 'analyst':
                draft={'title':'经营简报', 'editor_document':{'type':'doc','content':[{'type':'paragraph','content':[
                    {'type':'text','text':'2025年收入1200万元。'}, {'type':'citation','attrs':{'sourceId':self.source['id']}}]}]},
                    'citations':[{'source_id':self.source['id'],'locator':'line 1'}]}
                packet=json.loads((Path(self.sessions[sid]['packet_root'])/'input.json').read_text())
                self.queues[sid]=([('save_revision_metadata',{'responses':[],'bindings':[]})] if packet.get('mode')=='revision' else [])+[('save_draft',draft),('check_draft',{}),('submit_draft',{})]
            elif role == 'evaluator':
                packet=json.loads((Path(self.sessions[sid]['packet_root'])/'input.json').read_text())
                self.queues[sid]=[('submit_assessment',{'assessment':{'brief_hash':packet['brief']['hash'],'summary':'synthetic check','overall':'达到要求','evidence':3,'coverage':3,'analysis':3,'expression':3}})]
            else:raise AssertionError(role)
            self.next(sid)
        elif method == 'tool_result':
            if not p['ok']:
                self.sinks[self.sessions[sid]['eid']].put({'kind':'end','status':'failed','error':p['error']})
                return {}
            if self.sessions[sid]['role']=='analyst' and 'content' in p and 'settle' not in p:
                receipt=json.loads(p['content'][0]['text'])
                if receipt.get('revision'):
                    self.sessions[sid]['draft_revision']=receipt['revision']
            if 'settle' in p:
                self.sinks[self.sessions[sid]['eid']].put({'kind':'end','status':'completed','final_text':p['settle']})
            else:self.next(sid)
        return {}
    def next(self,sid):
        name,args=self.queues[sid].pop(0)
        if name in ('check_draft','submit_draft'):
            args={'revision':self.sessions[sid]['draft_revision']}
        self.sinks[self.sessions[sid]['eid']].put({'kind':'tool_request','tool':name,'args':args,'request_id':name})
    def close(self):pass


def setup(tmp_path, *, research_protocol='quality_v1', research_tier='standard'):
    store=Store(tmp_path)
    store.set_meta('settings',{**store.settings(),'agent_backend':'briefloop-native','model':'fixture/model',
        'model_selection_required':False,'auto_revision':False,'max_parallel':2})
    source=store.add_source('财报','2025年收入1200万元。')
    run=store.create_run({'title':'经营简报','objective':'解释经营情况','allow_web':False,
        'research_tier':research_tier},[source['id']],research_protocol=research_protocol)
    job=store.enqueue('generate',{'run_id':run['id'],'single_evaluation':False})
    store.update_job(job['id'], status='running')
    return store,source,run,store.one('jobs',job['id'])


def test_worker_native_plan_parallel_research_writer_and_saved_draft(tmp_path):
    store,source,run,job=setup(tmp_path)
    engine=FlowEngine(store,run,source);harness=NativeHarness(store,engine)
    worker=Worker(store);worker.runtime=InteractiveRuntime(store,backends={'briefloop-native':harness})
    result=worker.generate(job,score=False)
    brief=store.one('briefs',result['version_id'])
    assert brief['run_id']==run['id'] and '1200' in brief['markdown']
    assert json.loads(brief['detail'])['reconciliation_id']
    assert len(store.rows('SELECT * FROM briefs'))==1
    assert sorted(s['role'] for s in engine.sessions.values())==['analyst','orchestrator','scout','scout']
    assert {s['model'] for s in engine.sessions.values()}=={'fixture/model'}
    assert not harness._child_runtimes[next(sid for sid,s in engine.sessions.items() if s['role']=='orchestrator')]
    before=len([c for c in engine.calls if c[0]=='turn_start'])
    assert worker.generate(job,score=False)['version_id']==brief['id']
    assert len([c for c in engine.calls if c[0]=='turn_start'])==before


def test_native_quick_run_without_frozen_plan_reads_local_sources_and_saves_report(tmp_path):
    from briefloop.research_plan import frozen
    store,source,run,job=setup(tmp_path,research_protocol=None,research_tier='quick')
    engine=FlowEngine(store,run,source);harness=NativeHarness(store,engine)
    worker=Worker(store);worker.runtime=InteractiveRuntime(store,backends={'briefloop-native':harness})
    try:
        result=worker.generate(job,score=False)
        folder=worker.folder(job)
        research=json.loads((folder/'research.json').read_text(encoding='utf-8'))
        assert len(research['sources'])==1  # both Scouts read the same local source
        assert research['sources'][0]['source_id']==source['id']
        assert research['sources'][0]['excerpt']=='2025年收入1200万元。'
        assert not any('未完成' in gap for gap in research['gaps'])
        reads=[p for method,p in engine.calls if method=='tool_result' and p['request_id']=='source_read']
        assert len(reads)==2 and all(p['ok'] and 'source_hash: '+source['hash'] in p['content'][0]['text'] for p in reads)
        scouts=[s for s in engine.sessions.values() if s['role']=='scout']
        assert len(scouts)==2
        for session in scouts:
            task=json.loads((Path(session['packet_root'])/'task.json').read_text(encoding='utf-8'))
            assert task['research_handoff'] is None and task['allow_web'] is False
            assert 'web_search' not in {tool['name'] for tool in session['runner_tools']}
        brief=store.one('briefs',result['version_id'])
        assert '1200' in brief['markdown'] and json.loads(brief['detail'])['citations'][0]['source_id']==source['id']
        assert len(store.rows('SELECT * FROM briefs'))==1 and frozen(store,run['id']) is None
    finally:harness.close()


def test_chat_permissions_frozen_runtime_and_submit_replay(tmp_path):
    store=Store(tmp_path)
    h=NativeHarness(store,SimpleNamespace(process=None))
    sid=h.create_session(runtime={'model':'fixture/model'})['id']
    config={'native_role':'chat','session_id':sid,'attempt_id':'m1','_harness':h,'model':'fixture/model','permission':'read-only','allow_web':False}
    source=store.add_source('Synthetic','Data')
    req={'request':{'action':'generate','source_ids':[source['id']],'runtime':{'agent_backend':'codex','model':'other'},'requirements':{'title':'T','objective':'O','allow_web':False}}}
    assert not run_tool(store,config,'workspace_action',req)['ok']
    config['permission']='workspace-write'
    config['discuss_only']=True
    assert not run_tool(store,config,'workspace_action',req)['ok']
    config['discuss_only']=False
    req['request']['requirements']['allow_web']=True
    assert not run_tool(store,config,'workspace_action',req)['ok']
    req['request']['requirements']['allow_web']=False
    h.chat.message(sid, '原始请求：10分钟目标，12次搜索，保留四个问题', mid='m1')
    req['request']['requirements'].update(target_minutes=10, hard_timeout_minutes=0,
        research_budget={'search_requests':12,'candidate_urls':60,'source_pages':20}, key_questions=['问题一'])
    first=run_tool(store,config,'workspace_action',req)
    assert first['ok'],first
    admitted=json.loads(first['content'][0]['text'])
    assert admitted['accepted_requirements']['target_minutes']==10
    assert admitted['accepted_requirements']['research_budget']['search_requests']==12
    assert admitted['accepted_requirements']['key_questions']==['问题一']
    assert json.loads(store.one('runs',admitted['run_id'])['requirements'])['raw_input'].startswith('原始请求')
    assert json.loads(run_tool(store,config,'workspace_action',req)['content'][0]['text'])==json.loads(first['content'][0]['text'])
    assert len(store.rows('SELECT * FROM jobs'))==1
    payload=json.loads(store.rows('SELECT * FROM jobs')[0]['payload'])
    assert payload['agent_backend']=='briefloop-native' and payload['runtime']['model']=='fixture/model'
    h.cancel(sid)
    assert not run_tool(store,config,'workspace_action',{'request':{'action':'profile_update','profile':{'name':'late'}}})['ok']


def test_parent_cancel_only_cancels_its_children(tmp_path):
    h=NativeHarness(Store(tmp_path),SimpleNamespace(process=None))
    a=h.create_session(runtime={'model':'fixture/model'})['id'];b=h.create_session(runtime={'model':'fixture/model'})['id']
    class Child:
        def __init__(self):self.cancelled=False
        def cancel(self):self.cancelled=True
    ca,cb=Child(),Child()
    with h.child_runtime(a,ca),h.child_runtime(b,cb):
        h.cancel(a)
        assert ca.cancelled and not cb.cancelled
        with pytest.raises(InterruptedError):
            with h.child_runtime(a,Child()):pass


def test_native_revision_retains_original_and_rechecks_only_once(tmp_path):
    store,source,run,job=setup(tmp_path)
    engine=FlowEngine(store,run,source);harness=NativeHarness(store,engine)
    worker=Worker(store);worker.runtime=InteractiveRuntime(store,backends={'briefloop-native':harness})
    result=worker.generate(job,score=False)
    original=store.one('briefs',result['version_id'])
    store.assess(original['id'],{'brief_hash':original['hash'],'summary':'needs correction','overall':'建议修改',
        'evidence':3,'coverage':3,'analysis':3,'expression':3})
    revised=worker.auto_revise(job,original,worker.folder(job))
    assert revised['revision_status']=='complete'
    assert store.one('briefs',revised['version_id'])['parent_id']==original['id']
    assert store.one('briefs',original['id'])['hash']==original['hash']
    assert len(store.rows('SELECT * FROM assessments'))==2
    turns=len([c for c in engine.calls if c[0]=='turn_start'])
    assert worker.auto_revise(job,original,worker.folder(job))['version_id']==revised['version_id']
    assert len([c for c in engine.calls if c[0]=='turn_start'])==turns


def test_submission_and_receipt_roll_back_together(tmp_path,monkeypatch):
    from briefloop.external_requests import _RequestStore
    store=Store(tmp_path);source=store.add_source('S','Synthetic')
    h=NativeHarness(store,SimpleNamespace(process=None))
    sid=h.create_session(runtime={'model':'fixture/model'})['id']
    config={'native_role':'chat','session_id':sid,'attempt_id':'m1','_harness':h,'model':'fixture/model','permission':'workspace-write','allow_web':False}
    original=_RequestStore.set_meta
    def reject_receipt(self,key,value):
        if key.startswith('native_submit:'):raise RuntimeError('simulated receipt failure')
        return original(self,key,value)
    monkeypatch.setattr(_RequestStore,'set_meta',reject_receipt)
    result=run_tool(store,config,'workspace_action',{'request':{'action':'generate','requirements':{'title':'T','objective':'O','allow_web':False},'source_ids':[source['id']]}})
    assert not result['ok']
    assert not store.rows('SELECT * FROM runs') and not store.rows('SELECT * FROM jobs')


def test_scout_and_analyst_receive_candidate_skill_without_changing_active_skill(tmp_path):
    from briefloop import scout, analyst
    store,source,run,job=setup(tmp_path)
    candidate={'id':'candidate','content':'CANDIDATE_ONLY_SENTINEL','targets':['scout','analyst']}
    task=scout.task(store,run['id'],{'slot_id':'scout-1'},skill_override=candidate)
    assert task['skill']==candidate['content']
    frozen=analyst.packet(store,run['id'],store.root/'candidate',plan={},research={'sources':[],'gaps':[]},skill_override=candidate)
    assert candidate['content'] in (frozen['root']/'writing.md').read_text()
    assert candidate['content'] not in scout.task(store,run['id'],{'slot_id':'scout-1'})['skill']


def test_writer_recovery_refuses_changed_packet_before_model_call(tmp_path):
    from briefloop import analyst
    store,source,run,job=setup(tmp_path)
    class Writer:
        def execute(self,job,prompt,folder,*args,**kw):
            from test_native_analyst import finish
            config={**job['native_packet'],'native_role':'analyst','packet_root':str(folder/'packet'),'attempt_id':'fixture'}
            assert finish(store,config,{'title':'T','editor_document':{'type':'doc','content':[{'type':'paragraph','content':[{'type':'text','text':'Synthetic report'}]}]}})['ok']
    folder=store.root/'writer'
    analyst.run(store,Writer(),job,run['id'],folder,'briefloop-native',plan={},research={'sources':[],'gaps':[]},publish=False)
    (folder/'packet'/'writing.md').write_text('changed outside the task')
    with pytest.raises(ValueError,match='冻结写作材料被修改'):
        analyst.run(store,SimpleNamespace(execute=lambda *a,**k:pytest.fail('must not call model')),job,run['id'],folder,'briefloop-native',plan={},research={'sources':[],'gaps':[]},publish=False)


def test_template_preparation_uses_packet_and_existing_template_admission(tmp_path):
    from io import BytesIO
    from docx import Document
    from briefloop import templates
    from briefloop.native_orchestrator import prepare
    store=Store(tmp_path);doc=Document();doc.add_heading('Summary',1);doc.add_paragraph('Old content')
    buf=BytesIO();doc.save(buf)
    template=templates.import_template(store,'reference.docx',buf.getvalue(),prepare_job=False)
    job=store.enqueue('prepare_template',{'template_id':template['id'],'agent_backend':'briefloop-native','runtime':{'model':'fixture/model'}})
    folder=store.root/'jobs'/job['id'];folder.mkdir(parents=True)
    cfg,prompt=prepare(store,job,folder,templates.preparation_prompt(store,template,folder))
    h=NativeHarness(store,SimpleNamespace(process=None));sid=h.create_session(runtime={'model':'fixture/model'})['id']
    cfg={**cfg,'native_role':cfg['role'],'packet_root':str(folder/'packet'),'session_id':sid,'_harness':h}
    assert (folder/'packet'/'inventory.json').exists()
    result=run_tool(store,cfg,'submit_template',{'sections':[{'section_id':'summary','title':'Summary','index':0,'purpose':'Main result'}],'keep_blocks':[],'paragraph_index':1,'fields':[]})
    assert result['ok'],result
    assert templates.template(store,template['id'])['status']=='ready'
    assert not run_tool(store,cfg,'run_scouts',{'tasks':[]})['ok']


def test_fact_checker_uses_existing_stage_admission_and_rejects_wrong_version(tmp_path):
    from test_fact_check_contract import checked, good_result
    from briefloop import fact_check
    from briefloop.native_orchestrator import prepare
    world=checked(tmp_path);store=world['store'];job=world['job']
    payload={**json.loads(job['payload']),'version_id':world['brief']['id']}
    job={**job,'payload':dump(payload),'allow_web':True}
    folder=store.root/'jobs'/job['id'];folder.mkdir(parents=True)
    prompt=fact_check.fact_check_prompt(store,job,world['brief'],folder,'briefloop-native')
    cfg,_=prepare(store,job,folder,prompt)
    cfg={**cfg,'native_role':cfg['role'],'packet_root':str(folder/'packet')}
    result=good_result(world)
    assert not run_tool(store,cfg,'submit_fact_check',{**result,'version_id':'wrong'})['ok']
    admitted=run_tool(store,cfg,'submit_fact_check',result)
    assert admitted['ok'],admitted
    assert fact_check.records_for(store,world['run']['id'])
    assert not run_tool(store,cfg,'submit_fact_check',result)['ok']


def test_revision_metadata_rejects_number_bindings_before_saving(tmp_path):
    from briefloop.native_orchestrator import revision_metadata
    from briefloop.native_roles import ToolError
    store=Store(tmp_path/'ws');folder=store.root/'jobs/revision';(folder/'packet').mkdir(parents=True)
    (folder/'input.json').write_text(dump({'review_findings':[]}))
    config={'packet_root':str(folder/'packet'),'run_id':'run-test'}
    with pytest.raises(ToolError,match='number_bindings'):
        revision_metadata(store,config,{'responses':[],'bindings':[{'source_id':'source','value':97}]})
    assert not (folder/'revision_bindings.json').exists()
    assert not (folder/'responses.json').exists()
    revision_metadata(store,config,{'responses':[],'bindings':[]})
    assert json.loads((folder/'revision_bindings.json').read_text())==[]


def test_metadata_repair_rejects_source_ids_as_claims_before_settling(tmp_path):
    from briefloop.native_orchestrator import metadata_submit
    from briefloop.native_roles import ToolError
    store=Store(tmp_path/'ws');folder=store.root/'repair';(folder/'packet').mkdir(parents=True)
    packet={'version_id':'v1','brief_hash':'h1','candidate_claims':[],'document':{'type':'doc','content':[]},'findings':[]}
    (folder/'packet/input.json').write_text(dump(packet))
    config={'packet_root':str(folder/'packet')}
    args={'version_id':'v1','brief_hash':'h1','bindings':[{'claim_id':'src_a','block_id':'b','quote':'text'}],'responses':[]}
    with pytest.raises(ToolError,match='source_id'):
        metadata_submit(store,config,args)
    assert not (folder/'metadata.json').exists()
    result=metadata_submit(store,config,{**args,'bindings':[]})
    assert result['settle']
