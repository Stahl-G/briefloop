"""OfficeCLI capability layer: detection, command grammar, checks and renders.

The binary under test is a stub that enforces the real CLI's grammar (validate/
view issues/view screenshot with --page and -o), so a wrong command form — for
example a top-level `screenshot` — fails here on hosts without officecli.
"""
import base64
import json
import os
import subprocess
import threading
from io import BytesIO
from pathlib import Path

import pytest

from briefloop import host_bins, office_cli
from briefloop.store import Store

STUB_MODE_VARIABLE = 'OFFICE_STUB_MODE'
STUB_LOG_VARIABLE = 'OFFICE_STUB_LOG'

_STUB_TEMPLATE = r'''#!/usr/bin/env python3
import base64, json, os, sys, time
from pathlib import Path

PNG = base64.b64decode("__PNG_B64__")

def emit(payload, code=0):
    print(json.dumps(payload, ensure_ascii=False), flush=True)
    raise SystemExit(code)

def fail(message, code_name, code=1):
    emit({"success": False, "error": {"error": message, "code": code_name}}, code)

argv = [str(item) for item in sys.argv[1:]]
mode = os.environ.get("OFFICE_STUB_MODE", "")
log = os.environ.get("OFFICE_STUB_LOG")
if log:
    with open(log, "a", encoding="utf-8") as handle:
        handle.write(json.dumps(argv) + "\n")
if argv == ["--version"]:
    print("9.9.9-stub", flush=True)
    raise SystemExit(0)
if not argv or argv[0] not in ("validate", "view"):
    fail("Unrecognized command or argument " + (argv[0] if argv else "<missing>"), "invalid_command")
if len(argv) < 2 or not Path(argv[1]).is_file():
    fail("File not found: " + (argv[1] if len(argv) > 1 else "<missing>"), "file_not_found")
if mode == "hang":
    time.sleep(30)
if argv[0] == "validate":
    if mode in ("validate_fail", "all_fail"):
        fail("Validation failed: synthetic breakage", "validation_failed")
    if mode == "bad_json":
        print("this is not a json envelope", flush=True)
        raise SystemExit(0)
    emit({"success": True, "data": "Validation passed: no errors found.", "message": "ok"})
if len(argv) < 3 or argv[2] not in ("issues", "screenshot"):
    fail("Unrecognized command or argument " + (argv[2] if len(argv) > 2 else "<missing>"), "invalid_command")
if argv[2] == "issues":
    if mode == "bad_json":
        print("not json", flush=True)
        raise SystemExit(0)
    if mode in ("view_fail", "all_fail"):
        fail("view failed", "view_failed")
    found = [] if mode not in ("issues", "issues_noise", "issues_punct") else [
        {"type": 1, "message": "表格行列不平衡", "path": "/body/table[1]"},
        {"type": 1, "message": "空段落", "path": "/body/paragraph[3]"}]
    if mode == "issues_noise":  # punctuation noise mixed with a real finding
        found = [
            {"type": 1, "message": "Duplicate punctuation", "path": "/body/p[1]"},
            {"type": 1, "message": "Consecutive spaces", "path": "/body/p[2]"},
            {"type": 1, "message": "表格行列不平衡", "path": "/body/table[1]"}]
    if mode == "issues_punct":  # nothing but punctuation noise
        found = [
            {"type": 1, "message": "Duplicate punctuation", "path": "/body/p[1]"},
            {"type": 1, "message": "Consecutive spaces", "path": "/body/p[2]"}]
    emit({"success": True, "data": {"count": len(found), "issues": found}})
# view <file> screenshot --page N -o out.png --json is the only render form.
rest = argv[3:]
if "--page" not in rest or "-o" not in rest or "--json" not in rest:
    fail("screenshot requires --page, -o and --json", "invalid_usage")
if mode == "screenshot_fail":
    fail("render failed", "render_failed")
target = Path(rest[rest.index("-o") + 1])
target.parent.mkdir(parents=True, exist_ok=True)
target.write_bytes(PNG)
emit({"success": True, "data": str(target.resolve()), "message": str(target.resolve())})
'''


def _png_bytes():
    from PIL import Image
    buffer = BytesIO()
    Image.new('RGB', (12, 6), 'green').save(buffer, format='PNG')
    return buffer.getvalue()


def install_stub(tmp_path, monkeypatch, mode=None):
    """Put a grammar-checked officecli stub where host_bins finds it."""
    directory = tmp_path / 'stub-bin'
    directory.mkdir(exist_ok=True)
    binary = directory / 'officecli'
    binary.write_text(_STUB_TEMPLATE.replace('__PNG_B64__', base64.b64encode(_png_bytes()).decode()))
    binary.chmod(0o755)
    monkeypatch.setenv('PATH', '/usr/bin:/bin')
    monkeypatch.setattr(host_bins, 'EXTRA_DIRS', (str(directory),))
    monkeypatch.setenv(STUB_LOG_VARIABLE, str(tmp_path / 'stub-calls.log'))
    if mode:
        monkeypatch.setenv(STUB_MODE_VARIABLE, mode)
    office_cli._clear_caches()
    return binary


def stub_calls(tmp_path):
    """Logged officecli invocations, without the cached --version probes."""
    log = tmp_path / 'stub-calls.log'
    calls = [json.loads(line) for line in log.read_text().splitlines()] if log.exists() else []
    return [call for call in calls if call != ['--version']]


@pytest.fixture(autouse=True)
def _reset_office_caches():
    yield
    office_cli._clear_caches()


def _tiny_docx(tmp_path, name='probe.docx'):
    from docx import Document
    path = tmp_path / name
    document = Document()
    document.add_heading('Probe report', 0)
    document.add_paragraph('Revenue was USD 12 million.')
    document.save(path)
    return path


def _workspace(tmp_path):
    return Store(tmp_path / 'workspace')


def _published_brief(store):
    source = store.add_source('Synthetic', 'Revenue was USD 12 million.')
    run = store.create_run({'title': 'Probe report', 'objective': 'Explain revenue', 'allow_web': False},
                           [source['id']])
    return store.publish(run['id'], {'title': 'Probe report', 'markdown': 'Revenue was USD 12 million.'})


@pytest.mark.skipif(os.name == 'nt', reason='stub uses a POSIX shebang')
def test_find_and_capability_report_the_stub(tmp_path, monkeypatch):
    binary = install_stub(tmp_path, monkeypatch)
    assert Path(office_cli.find()).resolve() == binary.resolve()
    store = _workspace(tmp_path)
    capability = office_cli.capability(store)
    assert capability['id'] == 'officecli' and capability['name'] == 'OfficeCLI'
    assert capability['installed'] is True and capability['version'] == '9.9.9-stub'
    assert capability['enabled'] is False and capability['available'] is False
    assert '已检测到' in capability['diagnostic'] and '未检测到' not in capability['diagnostic']
    store.update_settings({'officecli_enabled': True})
    capability = office_cli.capability(store)
    assert capability['enabled'] is True and capability['available'] is True


def test_missing_binary_degrades_and_hint_never_spawns(tmp_path, monkeypatch):
    monkeypatch.setenv('PATH', '/usr/bin:/bin')
    monkeypatch.setattr(host_bins, 'EXTRA_DIRS', ())
    office_cli._clear_caches()
    store = _workspace(tmp_path)
    capability = office_cli.capability(store)
    assert capability == {'id': 'officecli', 'name': 'OfficeCLI', 'installed': False, 'path': None,
                          'version': None, 'enabled': False, 'available': False,
                          'diagnostic': '未检测到 OfficeCLI；' + host_bins.SEARCH_HINT}
    store.update_settings({'officecli_enabled': True})
    assert office_cli.capability(store)['available'] is False  # switch alone is not availability

    def explode(*args, **kwargs):
        raise AssertionError('the state hint must not spawn processes')

    monkeypatch.setattr(subprocess, 'run', explode)
    assert office_cli.snapshot_hint(store) == {'installed': False, 'enabled': True}
    assert store.snapshot()['office'] == {'installed': False, 'enabled': True}


def test_settings_default_off_and_backfilled_for_old_workspaces(tmp_path):
    store = _workspace(tmp_path)
    assert store.settings()['officecli_enabled'] is False
    legacy = {key: value for key, value in store.settings().items() if key != 'officecli_enabled'}
    store.set_meta('settings', legacy)
    reopened = Store(store.root)
    assert reopened.settings()['officecli_enabled'] is False  # model_validate backfills
    assert reopened.update_settings({'officecli_enabled': True})['officecli_enabled'] is True


@pytest.mark.skipif(os.name == 'nt', reason='stub uses a POSIX shebang')
def test_disabled_switch_is_a_silent_noop(tmp_path, monkeypatch):
    install_stub(tmp_path, monkeypatch)
    store = _workspace(tmp_path)
    assert store.settings()['officecli_enabled'] is False

    def explode(*args, **kwargs):
        raise AssertionError('a disabled switch must not run officecli')

    monkeypatch.setattr(office_cli, 'run_json', explode)
    docx = _tiny_docx(tmp_path)
    assert office_cli.check_file(store, docx) is None
    assert store.rows('SELECT * FROM office_checks') == []
    assert store.rows("SELECT * FROM events WHERE kind='office_check'") == []


@pytest.mark.skipif(os.name == 'nt', reason='stub uses a POSIX shebang')
def test_check_file_uses_the_real_command_grammar(tmp_path, monkeypatch):
    binary = install_stub(tmp_path, monkeypatch)
    store = _workspace(tmp_path)
    store.update_settings({'officecli_enabled': True})
    docx = _tiny_docx(tmp_path)
    summary = office_cli.check_file(store, docx, job_id='job_probe', version_id='brief_probe')
    assert summary['tool'] == 'officecli' and summary['tool_version'] == '9.9.9-stub'
    assert summary['validate']['status'] == 'ok'
    assert summary['validate']['summary'].startswith('Validation passed')
    assert summary['issues']['status'] == 'ok' and summary['issues']['count'] == 0
    assert {(row['kind'], row['status']) for row in store.rows('SELECT kind,status FROM office_checks')} == {
        ('validate', 'ok'), ('issues', 'ok')}
    events = store.rows("SELECT * FROM events WHERE kind='office_check'")
    assert len(events) == 1 and events[0]['job_id'] == 'job_probe'
    assert json.loads(events[0]['data'])['validate']['status'] == 'ok'
    # Exactly the two full command forms the shipped CLI speaks (the stub log
    # records its own argv, without the binary path):
    assert stub_calls(tmp_path) == [['validate', str(docx), '--json'],
                                    ['view', str(docx), 'issues', '--json']]


@pytest.mark.skipif(os.name == 'nt', reason='stub uses a POSIX shebang')
def test_punctuation_noise_is_filtered_and_keeps_both_counts(tmp_path, monkeypatch):
    """Findings are neutral observations, and pure punctuation noise (duplicate
    punctuation, consecutive spaces — normal in CJK prose) only survives as a
    count, never as red rows."""
    install_stub(tmp_path, monkeypatch, mode='issues_noise')
    store = _workspace(tmp_path)
    store.update_settings({'officecli_enabled': True})
    docx = _tiny_docx(tmp_path)
    summary = office_cli.check_file(store, docx)
    issues = summary['issues']
    assert issues['status'] == 'observed'  # neutral observation label, not an alarm state
    assert [item['message'] for item in issues['items']] == ['表格行列不平衡']
    assert issues['total'] == 3 and issues['count'] == 1 and issues['noise_filtered'] == 2
    row = store.rows("SELECT * FROM office_checks WHERE kind='issues'")[0]
    assert row['status'] == 'observed' and row['summary'] == '观察 1 项（另过滤 2 项纯标点类噪音）'
    # The stored row reads back with the same before/after counts.
    view = office_cli._view_for_digest(store, office_cli._file_sha256(docx))
    assert view['issues']['count'] == 1 and view['issues']['noise_filtered'] == 2 and view['issues']['total'] == 3

    monkeypatch.setenv(STUB_MODE_VARIABLE, 'issues_punct')  # noise only
    summary = office_cli.check_file(store, docx)
    assert summary['issues']['status'] == 'ok'
    assert summary['issues']['count'] == 0 and summary['issues']['noise_filtered'] == 2
    assert summary['issues']['total'] == 2 and summary['issues']['items'] == []
    assert store.rows("SELECT summary FROM office_checks WHERE kind='issues'")[0]['summary'] \
        == '未发现质检问题（过滤 2 项纯标点类噪音）'


@pytest.mark.skipif(os.name == 'nt', reason='stub uses a POSIX shebang')
def test_wrong_command_forms_fail_against_real_grammar(tmp_path, monkeypatch):
    install_stub(tmp_path, monkeypatch)
    docx = _tiny_docx(tmp_path)
    binary = office_cli.find()
    top_level = office_cli.run_json([binary, 'screenshot', str(docx), '--page', '1',
                                     '-o', str(tmp_path / 'out.png'), '--json'], timeout=30)
    assert top_level['ok'] is False and 'Unrecognized command' in top_level['reason']
    missing = office_cli.run_json([binary, 'validate', str(tmp_path / 'absent.docx'), '--json'], timeout=30)
    assert missing['ok'] is False and 'File not found' in missing['reason']


@pytest.mark.skipif(os.name == 'nt', reason='stub uses a POSIX shebang')
def test_check_file_records_every_failure_branch_without_raising(tmp_path, monkeypatch):
    install_stub(tmp_path, monkeypatch)
    store = _workspace(tmp_path)
    store.update_settings({'officecli_enabled': True})
    docx = _tiny_docx(tmp_path)
    branches = (('issues', {'validate': 'ok', 'issues': 'observed'}),
                ('validate_fail', {'validate': 'error', 'issues': 'ok'}),
                ('bad_json', {'validate': 'error', 'issues': 'error'}),
                ('view_fail', {'validate': 'ok', 'issues': 'error'}),
                ('all_fail', {'validate': 'error', 'issues': 'error'}))
    for mode, expected in branches:
        monkeypatch.setenv(STUB_MODE_VARIABLE, mode)
        summary = office_cli.check_file(store, docx)
        assert {kind: summary[kind]['status'] for kind in expected} == expected, mode
        if mode == 'issues':
            assert summary['issues']['count'] == 2 and len(summary['issues']['items']) == 2
        if mode == 'bad_json':
            assert 'JSON' in summary['validate']['reason']
        # A refresh keeps one newest row per (file_sha256, kind).
        assert len(store.rows('SELECT * FROM office_checks')) == 2, mode
    monkeypatch.setenv(STUB_MODE_VARIABLE, 'hang')
    monkeypatch.setattr(office_cli, 'VALIDATE_TIMEOUT', 0.4)
    monkeypatch.setattr(office_cli, 'ISSUES_TIMEOUT', 0.4)
    summary = office_cli.check_file(store, docx)
    assert summary['validate']['status'] == 'error' and '超时' in summary['validate']['reason']
    assert summary['issues']['status'] == 'error' and '超时' in summary['issues']['reason']


@pytest.mark.skipif(os.name == 'nt', reason='stub uses a POSIX shebang')
def test_request_budget_skips_remaining_steps(tmp_path, monkeypatch):
    install_stub(tmp_path, monkeypatch, mode='hang')
    store = _workspace(tmp_path)
    store.update_settings({'officecli_enabled': True})
    docx = _tiny_docx(tmp_path)
    monkeypatch.setattr(office_cli, 'REQUEST_BUDGET_SECONDS', 1.5)
    monkeypatch.setattr(office_cli, 'MIN_STEP_SECONDS', 0.2)
    summary = office_cli.check_file(store, docx)
    assert summary['validate']['status'] == 'error'  # first step burns the budget
    assert summary['issues']['status'] == 'error'
    assert summary['issues']['reason'] == '请求时间预算用尽'
    assert len(stub_calls(tmp_path)) == 1  # the second step never spawned


def _completed_export(store, brief):
    from briefloop.export_jobs import enqueue_export, generate_word
    queued = enqueue_export(store, brief['id'])
    job = store.one('jobs', queued['id'])
    result = generate_word(store, job, threading.Event())
    store.update_job(job['id'], 'complete', result=result)
    return job, result


@pytest.mark.skipif(os.name == 'nt', reason='stub uses a POSIX shebang')
def test_generate_word_survives_permanent_officecli_failure(tmp_path, monkeypatch):
    install_stub(tmp_path, monkeypatch, mode='all_fail')
    store = _workspace(tmp_path)
    store.update_settings({'officecli_enabled': True})
    brief = _published_brief(store)
    from briefloop.export_jobs import enqueue_export, generate_word, output_path
    queued = enqueue_export(store, brief['id'])
    job = store.one('jobs', queued['id'])
    result = generate_word(store, job, threading.Event())
    assert output_path(store, job).is_file()
    assert result['office']['validate']['status'] == 'error'
    assert result['office']['issues']['status'] == 'error'
    progress = [json.loads(row['data']) for row in
                store.rows("SELECT data FROM events WHERE kind='export_progress' ORDER BY seq")]
    assert len(progress) == 4 and {item['step'] for item in progress} == {1, 2, 3, 4}
    assert all(item['total'] == 4 for item in progress)
    assert store.rows("SELECT * FROM events WHERE kind='office_check'")
    # The fingerprint reuse path is untouched by the quality gate.
    assert enqueue_export(store, brief['id'])['id'] == job['id']


@pytest.mark.skipif(os.name == 'nt', reason='stub uses a POSIX shebang')
def test_generate_word_without_the_switch_matches_current_behavior(tmp_path, monkeypatch):
    install_stub(tmp_path, monkeypatch)
    store = _workspace(tmp_path)
    assert store.settings()['officecli_enabled'] is False
    brief = _published_brief(store)
    from briefloop.export_jobs import enqueue_export, generate_word
    job = store.one('jobs', enqueue_export(store, brief['id'])['id'])
    result = generate_word(store, job, threading.Event())
    assert 'office' not in result
    assert store.rows('SELECT * FROM office_checks') == []
    assert store.rows("SELECT * FROM events WHERE kind='office_check'") == []


@pytest.mark.skipif(os.name == 'nt', reason='stub uses a POSIX shebang')
def test_run_check_for_job_validates_then_reruns(tmp_path, monkeypatch):
    install_stub(tmp_path, monkeypatch)
    store = _workspace(tmp_path)
    brief = _published_brief(store)
    job, result = _completed_export(store, brief)  # switch off: no rows from the hook
    assert store.rows('SELECT * FROM office_checks') == []
    with pytest.raises(ValueError, match='未检测到 OfficeCLI 或未开启'):
        office_cli.run_check_for_job(store, job['id'])
    store.update_settings({'officecli_enabled': True})
    view = office_cli.run_check_for_job(store, job['id'])
    assert view['file_sha256'] == result['sha256']
    assert view['job_kind'] == 'export_docx' and view['job_id'] == job['id']
    assert view['tool_version'] == '9.9.9-stub'
    assert view['validate']['status'] == 'ok' and view['issues']['status'] == 'ok'
    assert office_cli.version_office_view(store, brief['id'])['file_sha256'] == result['sha256']
    office_cli.run_check_for_job(store, job['id'])  # a refresh replaces, not accumulates
    assert len(store.rows('SELECT * FROM office_checks')) == 2
    from briefloop.export_jobs import output_path
    output_path(store, job).write_bytes(b'changed')
    with pytest.raises(ValueError, match='工件已变化'):
        office_cli.run_check_for_job(store, job['id'])
    with pytest.raises(ValueError, match='只能对'):
        office_cli.run_check_for_job(store, store.enqueue('assess', {'version_id': brief['id']})['id'])


@pytest.mark.skipif(os.name == 'nt', reason='stub uses a POSIX shebang')
def test_generate_word_check_stage_stops_between_substeps_when_cancelled(tmp_path, monkeypatch):
    """A stop request during the post-export quality gate lands the job as
    cancelled between the two sub-steps instead of blocking on the second
    subprocess for its full budget."""
    install_stub(tmp_path, monkeypatch)
    store = _workspace(tmp_path)
    store.update_settings({'officecli_enabled': True})
    brief = _published_brief(store)
    from briefloop.export_jobs import enqueue_export, generate_word
    job = store.one('jobs', enqueue_export(store, brief['id'])['id'])
    cancelled = threading.Event()
    real_run_json = office_cli.run_json

    def run_json_then_cancel(args, **kwargs):
        outcome = real_run_json(args, **kwargs)
        cancelled.set()  # the stop arrives while the gate is running
        return outcome

    monkeypatch.setattr(office_cli, 'run_json', run_json_then_cancel)
    with pytest.raises(InterruptedError):
        generate_word(store, job, cancelled)
    # validate completed; the issues subprocess never started.
    assert [call[0] for call in stub_calls(tmp_path)] == ['validate']
    assert {row['kind'] for row in store.rows('SELECT kind FROM office_checks')} == {'validate'}
    assert not store.rows("SELECT * FROM events WHERE kind='office_check'")


@pytest.mark.skipif(os.name == 'nt', reason='stub uses a POSIX shebang')
def test_render_preview_caches_pages_and_serves_bound_images(tmp_path, monkeypatch):
    install_stub(tmp_path, monkeypatch)
    store = _workspace(tmp_path)
    store.update_settings({'officecli_enabled': True})
    from briefloop.sources import upload
    source = upload(store, 'sample.docx', _tiny_docx(tmp_path, 'upload.docx').read_bytes())
    result = office_cli.render_preview(store, {'source_id': source['id']})
    assert result['target'] == {'kind': 'source', 'id': source['id'], 'name': 'sample.docx'}
    assert result['tool'] == 'officecli' and result['tool_version'] == '9.9.9-stub'
    assert result['cached'] is False and result['incomplete'] is False and result['reason'] is None
    page = result['pages'][0]
    assert page['page'] == 1 and page['width'] == 12 and page['height'] == 6
    digest = page['url'].split('digest=')[1].split('&')[0]
    assert office_cli.office_image(store, digest, 1) == _png_bytes()
    assert len(stub_calls(tmp_path)) == 1  # validate/issues never run for previews
    again = office_cli.render_preview(store, {'source_id': source['id'], 'pages': [1, 1, 2]})
    assert again['cached'] is False  # page 2 was rendered fresh
    assert [item['page'] for item in again['pages']] == [1, 2]
    assert len(stub_calls(tmp_path)) == 2  # only page 2 spawned a new render
    assert office_cli.office_image(store, digest, 2) == _png_bytes()
    repeat = office_cli.render_preview(store, {'source_id': source['id'], 'pages': [1, 2]})
    assert repeat['cached'] is True and len(stub_calls(tmp_path)) == 2
    # Different sources also stage into separate per-request files.
    def distinct_docx(name, marker):
        from docx import Document
        path = tmp_path / name
        document = Document()
        document.add_heading('Probe ' + marker, 0)
        document.save(path)
        return path

    probe_a = distinct_docx('probe-a.docx', 'A')
    probe_b = distinct_docx('probe-b.docx', 'B')
    office_cli.render_page(store, probe_a, 1)
    office_cli.render_page(store, probe_b, 1)
    shots = [call for call in stub_calls(tmp_path)
             if call[0] == 'view' and 'screenshot' in call
             and (str(probe_a) in call or str(probe_b) in call)]
    assert len(shots) == 2, 'one fresh render per probe file'
    outs = {call[call.index('-o') + 1] for call in shots}
    assert len(outs) == 2, 'the same page rendered twice must not share one staging file'
    # …and a failed render only removes its own staging file, never another
    # request's. The decoy is what the old shared name would have deleted.
    monkeypatch.setenv(STUB_MODE_VARIABLE, 'screenshot_fail')
    probe_c = distinct_docx('probe-c.docx', 'C')
    digest_c = office_cli._file_sha256(probe_c)
    directory = office_cli._render_directory(store, digest_c)
    decoy = directory / '.render-1.tmp.png'  # the pre-fix shared staging name
    decoy.write_bytes(b'another request in flight')
    with pytest.raises(ValueError, match='预览失败'):
        office_cli.render_page(store, probe_c, 1)
    assert decoy.is_file(), 'cleanup must only delete this request\'s staging file'
    assert list(directory.glob('.render-*.tmp.png')) == [decoy]
    with pytest.raises(ValueError):
        office_cli.office_image(store, 'zz' * 32, 1)  # not a 64-hex digest
    with pytest.raises(ValueError):
        office_cli.office_image(store, digest, 99)  # nothing cached for that page
    note = store.add_source('note.txt', 'plain text only')
    with pytest.raises(ValueError):  # no renderable Office original at all
        office_cli.render_preview(store, {'source_id': note['id']})


def test_same_page_concurrent_publication_keeps_each_image_bound(tmp_path, monkeypatch):
    """Force two renders and both manifest publications to overlap, with
    different valid PNGs so a mixed image/metadata pair cannot pass by chance."""
    from concurrent.futures import ThreadPoolExecutor
    import hashlib
    from PIL import Image

    store = _workspace(tmp_path)
    target = tmp_path / 'same.docx'
    target.write_bytes(b'one frozen source')
    render_barrier, publish_barrier = threading.Barrier(2), threading.Barrier(2)
    payloads = []
    for color in ('red', 'blue'):
        buf = BytesIO()
        Image.new('RGB', (12, 6), color).save(buf, format='PNG')
        payloads.append(buf.getvalue())
    assignments, assignment_lock, manifests = [], threading.Lock(), []
    real_replace = os.replace

    def render(args, **kwargs):
        with assignment_lock:
            payload = payloads[len(assignments)]
            assignments.append(args[args.index('-o') + 1])
        Path(args[args.index('-o') + 1]).write_bytes(payload)
        render_barrier.wait(timeout=5)
        return {'ok': True, 'data': {}, 'reason': None}

    def replace(src, dst):
        if str(dst).endswith('.json'):
            manifests.append(str(src))
            publish_barrier.wait(timeout=5)
        real_replace(src, dst)
        if str(dst).endswith('.json'):
            # Even while another request publishes, the visible manifest must
            # refer to complete immutable bytes, never a half-updated pair.
            assert office_cli._cached_render(Path(dst).parent, digest, 1)

    monkeypatch.setattr(office_cli, 'find', lambda: 'synthetic-renderer')
    monkeypatch.setattr(office_cli, 'version', lambda *_: 'synthetic')
    monkeypatch.setattr(office_cli, 'run_json', render)
    monkeypatch.setattr(office_cli.os, 'replace', replace)
    digest = office_cli._file_sha256(target)
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(office_cli.render_page, store, target, 1) for _ in range(2)]
        results = [future.result(timeout=10) for future in futures]
    assert len(set(assignments)) == len(set(manifests)) == 2
    assert {row['image_sha256'] for row in results} == {hashlib.sha256(p).hexdigest() for p in payloads}
    for row in results:
        assert hashlib.sha256(Path(row['path']).read_bytes()).hexdigest() == row['image_sha256']
    assert office_cli.office_image(store, digest, 1) in payloads
    assert not list(Path(results[0]['path']).parent.glob('*.tmp'))


@pytest.mark.skipif(os.name == 'nt', reason='stub uses a POSIX shebang')
def test_partial_preview_names_real_failure_apart_from_budget(tmp_path, monkeypatch):
    """A page that actually failed while the budget was also spent must be
    reported as a render failure, not relabelled as budget exhaustion."""
    install_stub(tmp_path, monkeypatch)
    store = _workspace(tmp_path)
    store.update_settings({'officecli_enabled': True})
    from briefloop.sources import upload
    source = upload(store, 'sample.docx', _tiny_docx(tmp_path, 'upload.docx').read_bytes())

    def render_page1_then_fail(args, *, timeout, deadline=None):
        # Page 1 succeeds (written straight to the staging path); page 2 dies
        # with a real timeout while the shared budget is already gone.
        if args[args.index('--page') + 1] == '2':
            return {'ok': False, 'data': None, 'reason': 'officecli 执行超时（180 秒）'}
        Path(args[args.index('-o') + 1]).write_bytes(_png_bytes())
        return {'ok': True, 'data': 'ok', 'reason': None}

    monkeypatch.setattr(office_cli, 'run_json', render_page1_then_fail)
    monkeypatch.setattr(office_cli, 'REQUEST_BUDGET_SECONDS', -1)  # budget already spent
    result = office_cli.render_preview(store, {'source_id': source['id'], 'pages': [1, 2]})
    assert result['incomplete'] is True
    assert [page['page'] for page in result['pages']] == [1]
    assert '渲染失败' in result['reason'] and 'officecli 执行超时' in result['reason']
    assert result['reason'] != '请求时间预算用尽，仅返回已完成页面'

    def budget_skips_every_page(args, *, timeout, deadline=None):
        return {'ok': False, 'data': None, 'reason': '请求时间预算用尽'}

    monkeypatch.setattr(office_cli, 'run_json', budget_skips_every_page)
    monkeypatch.setattr(office_cli, 'REQUEST_BUDGET_SECONDS', 180)
    result = office_cli.render_preview(store, {'source_id': source['id'], 'pages': [3, 4]})
    assert result['incomplete'] is True and result['pages'] == []
    assert result['reason'] == '请求时间预算用尽，仅返回已完成页面'


@pytest.mark.skipif(os.name == 'nt', reason='stub uses a POSIX shebang')
def test_preview_and_check_gates(tmp_path, monkeypatch):
    install_stub(tmp_path, monkeypatch)
    store = _workspace(tmp_path)
    store.update_settings({'officecli_enabled': True})
    brief = _published_brief(store)
    job, _result = _completed_export(store, brief)
    store.update_settings({'officecli_enabled': False})
    with pytest.raises(ValueError, match='未检测到 OfficeCLI 或未开启'):
        office_cli.render_preview(store, {'job_id': job['id']})
    store.update_settings({'officecli_enabled': True})
    for body in ({}, {'job_id': job['id'], 'source_id': 'src_x'},
                 {'job_id': job['id'], 'pages': [0]},
                 {'job_id': job['id'], 'pages': [1] * 5},
                 {'job_id': 'job_missing'},
                 {'release_id': 'rel_missing'}):
        with pytest.raises(ValueError):
            office_cli.render_preview(store, body)
    ok = office_cli.render_preview(store, {'job_id': job['id'], 'pages': [1]})
    assert ok['target']['kind'] == 'export' and ok['target']['name'] == 'report.docx'


def test_preview_requires_the_binary_even_when_enabled(tmp_path, monkeypatch):
    monkeypatch.setenv('PATH', '/usr/bin:/bin')
    monkeypatch.setattr(host_bins, 'EXTRA_DIRS', ())
    office_cli._clear_caches()
    store = _workspace(tmp_path)
    store.update_settings({'officecli_enabled': True})
    with pytest.raises(ValueError, match='未检测到 OfficeCLI 或未开启'):
        office_cli.render_preview(store, {'source_id': 'src_x'})
    with pytest.raises(ValueError, match='未检测到 OfficeCLI 或未开启'):
        office_cli.render_page(store, _tiny_docx(tmp_path), 1)


@pytest.mark.skipif(not office_cli.find(), reason='officecli is not installed on this host')
def test_real_officecli_validate_issues_and_screenshot(tmp_path):
    store = _workspace(tmp_path)
    store.update_settings({'officecli_enabled': True})
    docx = _tiny_docx(tmp_path)
    summary = office_cli.check_file(store, docx, job_id='job_probe')
    assert summary['tool'] == 'officecli' and summary['tool_version']
    assert summary['validate']['status'] == 'ok'
    assert summary['issues']['status'] in ('ok', 'issues')
    rendered = office_cli.render_page(store, docx, 1)
    payload = Path(rendered['path']).read_bytes()
    assert payload[:8] == b'\x89PNG\r\n\x1a\n' and rendered['width'] > 0 and rendered['height'] > 0
    from PIL import Image
    with Image.open(BytesIO(payload)) as image:
        assert image.format == 'PNG'
    assert office_cli.office_image(store, rendered['digest'], 1) == payload


def test_every_invocation_skips_the_tools_self_update(tmp_path, monkeypatch):
    """Without these variables each call runs officecli's daily update check:
    it writes ~/.officecli, contacts the vendor mirror in a background process
    and refreshes skill files it installed into other AI tools."""
    binary = tmp_path / 'officecli'
    binary.write_text('')
    seen = []

    def record(command, **kwargs):
        seen.append(kwargs.get('env') or {})
        return subprocess.CompletedProcess(command, 0, stdout='{"success":true,"data":"ok"}', stderr='')

    monkeypatch.setattr(subprocess, 'run', record)
    monkeypatch.delenv('OFFICECLI_SKIP_UPDATE', raising=False)
    office_cli._clear_caches()
    office_cli.version(str(binary))
    office_cli.run_json([str(binary), 'validate', 'x.docx', '--json'], timeout=5)
    assert len(seen) == 2
    for env in seen:
        assert env.get('OFFICECLI_SKIP_UPDATE') == '1' and env.get('OFFICECLI_NO_AUTO_INSTALL') == '1'
        assert env.get('PATH') == os.environ.get('PATH')  # the rest of the environment is kept
