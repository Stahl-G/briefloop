from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook

from briefloop.evidence import create_span
from briefloop.media import source_files
from briefloop.sources import upload
from briefloop.store import Store


def office_bytes(document):
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def cell_text(text, path):
    return text.split('[DOCX ' + path + ' columns=', 1)[1].split('[DOCX end ' + path + ']', 1)[0]


def test_xlsx_empty_inline_string_upload_keeps_other_cells_and_original(tmp_path):
    book = Workbook(); sheet = book.active
    sheet.append(['备注', '']); sheet.append(['金额', 12]); sheet.append(['计算', '=1+1'])
    data = office_bytes(book)
    assert load_workbook(BytesIO(data)).active['B2'].value == 12
    store = Store(tmp_path)
    source = upload(store, 'values.xlsx', data)
    assert source['status'] == 'ready'
    text = store.source_text(source['id'])
    assert 'A1: 备注 | B1: ' in text and 'A2: 金额 | B2: 12' in text
    assert 'B3: [公式没有保存值]' in text
    _, provenance, original = source_files(store, source['id'])
    assert original.read_bytes() == data and provenance['extraction_status'] == 'ready'


def test_docx_source_distinguishes_cell_paragraphs_and_keeps_line_evidence(tmp_path):
    document = Document(); document.add_paragraph('正文在表格之前')
    layouts = [
        [[['指标'], ['2025'], ['2026']], [['收入'], ['11', '22'], ['33']]],
        [[['指标'], ['2025'], ['2026']], [['收入'], ['11'], ['22', '33']]],
    ]
    for layout in layouts:
        table = document.add_table(rows=2, cols=3)
        for ri, row in enumerate(layout):
            for ci, paragraphs in enumerate(row):
                cell = table.cell(ri, ci); cell.text = paragraphs[0]
                for text in paragraphs[1:]: cell.add_paragraph(text)
        document.add_paragraph('正文在表格之后')
    store = Store(tmp_path); data = office_bytes(document)
    source = upload(store, 'tables.docx', data)
    assert source['status'] == 'ready'
    text = store.source_text(source['id'])
    assert text.index('正文在表格之前') < text.index('[DOCX table 1]') < text.index('正文在表格之后') < text.index('[DOCX table 2]')
    first = cell_text(text, 'table 1/row 2/cell 2')
    second = cell_text(text, 'table 2/row 2/cell 2')
    assert '\n11\n' in first and '\n22\n' in first
    assert '\n11\n' in second and '\n22\n' not in second
    assert '\n22\n' in cell_text(text, 'table 2/row 2/cell 3')
    lines = text.splitlines(); line = lines.index('22') + 1
    span = create_span(store, {'source_id': source['id'], 'locator': {'kind': 'text', 'start_line': line, 'end_line': line}, 'excerpt': '22'})
    assert span['data']['located_text'] == '22'
    _, provenance, original = source_files(store, source['id'])
    assert original.read_bytes() == data and 'table cell coordinates' in provenance['extractor']


def test_docx_source_keeps_merge_markers_nested_scope_and_omitted_columns(tmp_path):
    document = Document(); table = document.add_table(rows=4, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = '横向合并'
    table.cell(1, 0).merge(table.cell(2, 0)).text = '纵向合并'
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = '内表唯一文本'
    nested.cell(0, 0).paragraphs[0].add_run().add_break()
    nested.cell(0, 0).paragraphs[0].add_run('换行后的文字')
    row = table.rows[3]; row.cells[1].text = '省略两侧的中列'
    for cell in (row.cells[0]._tc, row.cells[2]._tc): row._tr.remove(cell)
    for name in ('gridBefore', 'gridAfter'):
        value = OxmlElement('w:' + name); value.set(qn('w:val'), '1'); row._tr.get_or_add_trPr().append(value)
    store = Store(tmp_path); source = upload(store, 'merged.docx', office_bytes(document))
    assert source['status'] == 'ready'
    text = store.source_text(source['id'])
    assert 'table 1/row 1/cell 1 columns=1:2 gridSpan=2' in text
    assert 'table 1/row 2/cell 1 columns=1:1 gridSpan=1 vMerge=restart' in text
    continuation = cell_text(text, 'table 1/row 3/cell 1')
    assert 'vMerge=continue' in continuation and ' empty]' in continuation and '纵向合并' not in continuation
    assert '[DOCX table 1/row 2/cell 2/table 1]' in text
    assert text.count('内表唯一文本') == 1 and '\n内表唯一文本\n换行后的文字\n' in text
    assert 'row 4 gridBefore=1 gridAfter=1' in text
    assert 'table 1/row 4/cell 1 columns=2:2' in text


def test_empty_docx_table_still_has_visible_failed_source(tmp_path):
    document = Document(); document.add_table(rows=2, cols=2)
    store = Store(tmp_path); data = office_bytes(document)
    source = upload(store, 'empty.docx', data)
    assert source['status'] == 'failed' and '未能读取正文' in source['error']
    assert store.source_text(source['id']) == ''
    _, provenance, original = source_files(store, source['id'])
    assert original.read_bytes() == data and provenance['extraction_status'] == 'failed'
