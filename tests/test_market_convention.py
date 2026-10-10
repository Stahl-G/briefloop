"""Market colors are report-specific, semantic, and preserved in actual Word XML."""
from io import BytesIO
import json

import pytest
from docx import Document
from docx.shared import RGBColor

from briefloop.market_convention import (
    apply_docx_market_colors, resolve_market, save_report_settings, semantic_delta_spans,
)
from briefloop.models import Requirements
from briefloop.store import Store


def test_language_defaults_and_report_setting_survive_reload(tmp_path):
    assert Requirements(title='报告', objective='说明').market_convention == 'cn'
    assert Requirements(title='Report', objective='Explain', language='English').market_convention == 'intl'
    assert Requirements(title='Report', objective='Explain', language='en', market_convention='cn').market_convention == 'cn'
    with pytest.raises(ValueError):
        Requirements(title='报告', objective='说明', market_convention='unknown')
    assert resolve_market({'language': 'English'}) == 'intl'
    store = Store(tmp_path / 'workspace')
    chinese = store.create_run({'title': '报告', 'objective': '说明'}, [])
    english = store.create_run({'title': 'Report', 'objective': 'Explain', 'language': 'en'}, [])
    before = json.loads(chinese['requirements'])
    result = save_report_settings(store, {'workspace_id': store.meta('workspace_id'), 'run_id': chinese['id'], 'market_convention': 'intl'})
    reloaded = Store(tmp_path / 'workspace')
    after = json.loads(reloaded.one('runs', chinese['id'])['requirements'])
    assert result['requirements'] == after == {**before, 'market_convention': 'intl'}
    assert json.loads(reloaded.one('runs', english['id'])['requirements'])['market_convention'] == 'intl'
    save_report_settings(reloaded, {'workspace_id': reloaded.meta('workspace_id'), 'run_id': english['id'], 'market_convention': 'cn'})
    assert json.loads(reloaded.one('runs', chinese['id'])['requirements'])['market_convention'] == 'intl'
    assert json.loads(reloaded.one('runs', english['id'])['requirements'])['market_convention'] == 'cn'
    assert reloaded.meta('requirements')['market_convention'] == 'intl'  # saved report choices never rewrite the next-report defaults
    with pytest.raises(ValueError, match='工作区'):
        save_report_settings(store, {'workspace_id': 'different', 'run_id': chinese['id'], 'market_convention': 'cn'})
    with pytest.raises(ValueError, match='只支持'):
        save_report_settings(store, {'workspace_id': store.meta('workspace_id'), 'run_id': chinese['id'], 'market_convention': 'cn', 'objective': 'overwrite'})


@pytest.mark.parametrize('text,context,expected', [
    ('营收 100，利润 -20，利润率 -5%，存款 +25', '', []),
    ('revenue 100; net margin -5%; balance +25', '', []),
    ('+5%', '', []),
    ('同比 +5%，环比 −2.0%，增长 0%', '', [('+5%', 'up'), ('−2.0%', 'down'), ('0%', 'flat')]),
    ('Revenue grew 5%; margin decreased by 1.2 percentage points', '', [('5%', 'up'), ('1.2 percentage points', 'down')]),
    ('↑ 4.5% / ▼2%', '', [('↑ 4.5%', 'up'), ('▼2%', 'down')]),
    ('-1,200', 'YoY change', [('-1,200', 'down')]),
    ('  +1.5%  ', '环比', [('+1.5%', 'up')]),
    ('-25', '余额', []),
    ('-5%', 'Profit margin', []),
])
def test_only_explicit_change_spans_receive_direction(text, context, expected):
    assert [(text[start:end], direction) for start, end, direction in semantic_delta_spans(text, context=context)] == expected


def test_word_colors_only_semantic_numbers_and_preserves_authored_formatting():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('收入 -200，利润率 -5%；同比 ').bold = True
    paragraph.add_run('+5%').italic = True
    paragraph.add_run('，下跌 2%。')
    authored = document.add_paragraph().add_run('增长 7%')
    authored.font.color.rgb = RGBColor.from_string('123456')
    table = document.add_table(rows=2, cols=3)
    for index, value in enumerate(('营收', '同比', '净利率')):
        table.cell(0, index).text = value
    for index, value in enumerate(('-200', '-3%', '-4%')):
        table.cell(1, index).text = value
    original = paragraph.text
    apply_docx_market_colors(document, {'market_convention': 'intl'})
    data = BytesIO(); document.save(data); data.seek(0)
    saved = Document(data)
    assert saved.paragraphs[0].text == original
    colored = [(run.text, str(run.font.color.rgb)) for run in saved.paragraphs[0].runs if run.font.color.rgb]
    assert colored == [('+5%', '1E8E4F'), ('2%', 'D9363E')]
    assert next(run for run in saved.paragraphs[0].runs if run.text == '+5%').italic is True
    assert saved.paragraphs[0].runs[0].bold is True
    assert str(saved.paragraphs[1].runs[0].font.color.rgb) == '123456'
    assert str(saved.tables[0].cell(1, 1).paragraphs[0].runs[0].font.color.rgb) == 'D9363E'
    assert saved.tables[0].cell(1, 0).paragraphs[0].runs[0].font.color.rgb is None
    assert saved.tables[0].cell(1, 2).paragraphs[0].runs[0].font.color.rgb is None
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    assert saved.tables[0].cell(1, 1).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert saved.tables[0].cell(1, 0).paragraphs[0].alignment is None
    assert '<w:shd' not in saved._element.xml
    # Repeating a post-processing pass does not duplicate text or runs.
    before = saved._element.xml
    apply_docx_market_colors(saved, {'market_convention': 'intl'})
    assert saved._element.xml == before


def test_template_defaults_can_be_colored_without_overriding_authored_colors():
    from briefloop.document_export import render_document
    from briefloop.document_model import markdown_document
    from briefloop.templates import table_defaults
    template = Document()
    table = template.add_table(rows=2, cols=3)
    for row in table.rows:
        for cell in row.cells:
            run = cell.paragraphs[0].add_run('Example')
            run.font.color.rgb = RGBColor.from_string('334455')
    styles = table_defaults(template)
    source = markdown_document('| YoY | Revenue | Margin |\n|---|---|---|\n| -3% | -200 | -5% |\n\nGrowth 4%')
    source['content'][-1]['content'][0]['marks'] = [{'type': 'textStyle', 'attrs': {'color': '#123456'}}]
    doc, protected = Document(), set()
    render_document(doc, source, styles=styles, protected_runs=protected)
    assert str(doc.tables[0].cell(1, 0).paragraphs[0].runs[0].font.color.rgb) == '334455'
    apply_docx_market_colors(doc, {'market_convention': 'cn'}, protected_runs=protected)
    output = BytesIO(); doc.save(output); output.seek(0); saved = Document(output)
    assert str(next(run for run in saved.tables[0].cell(1, 0).paragraphs[0].runs if run.text).font.color.rgb) == '1E8E4F'
    assert str(saved.tables[0].cell(1, 1).paragraphs[0].runs[0].font.color.rgb) == '334455'
    assert str(saved.tables[0].cell(1, 2).paragraphs[0].runs[0].font.color.rgb) == '334455'
    assert str(saved.paragraphs[0].runs[0].font.color.rgb) == '123456'


def test_report_settings_route_validates_and_persists(tmp_path):
    import http.client
    import threading
    from briefloop.server import make_server, _close_service
    server = make_server(tmp_path / 'api-workspace', port=0, paused=True)
    thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
    base = f'http://127.0.0.1:{server.server_port}'
    def request(path, body=None, token=None):
        connection = http.client.HTTPConnection('127.0.0.1', server.server_port)
        headers = {'Origin': base}
        if token: headers['X-BriefLoop-Token'] = token
        connection.request('POST' if body is not None else 'GET', path, json.dumps(body) if body is not None else None, headers)
        response = connection.getresponse(); result = response.status, json.loads(response.read()); connection.close(); return result
    try:
        run = server.store.create_run({'title': '报告', 'objective': '说明'}, [])
        _, session = request('/api/session')
        payload = {'workspace_id': server.store.meta('workspace_id'), 'run_id': run['id'], 'market_convention': 'intl'}
        status, saved = request('/api/report-settings', payload, session['token'])
        assert status == 200 and saved['requirements']['market_convention'] == 'intl'
        assert json.loads(server.store.one('runs', run['id'])['requirements'])['market_convention'] == 'intl'
        status, _ = request('/api/report-settings', {**payload, 'market_convention': 'invalid'}, session['token'])
        assert status == 400
        assert json.loads(server.store.one('runs', run['id'])['requirements'])['market_convention'] == 'intl'
    finally:
        server.shutdown(); thread.join(); _close_service(server); server.server_close(); server.workspace_lock.close()


def test_word_merged_columns_skip_inference_and_explicit_alignment_survives():
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document(); table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = 'YoY'; table.cell(0, 1).text = 'Balance'
    table.cell(1, 0).merge(table.cell(2, 0)).text = '5%'
    table.cell(1, 1).text = '100'; table.cell(2, 1).text = '-200'
    apply_docx_market_colors(doc)
    assert all(run.font.color.rgb is None for row in table.rows for cell in row.cells for paragraph in cell.paragraphs for run in paragraph.runs)
    simple = doc.add_table(rows=2, cols=1)
    simple.cell(0, 0).text = '同比'; simple.cell(1, 0).text = '-3%'
    simple.cell(1, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_docx_market_colors(doc)
    assert simple.cell(1, 0).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert str(simple.cell(1, 0).paragraphs[0].runs[0].font.color.rgb) == '1E8E4F'
