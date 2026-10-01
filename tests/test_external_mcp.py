"""Official SDK stdio subprocess -> real local HTTP service -> real file Worker.

Synthetic materials and a seeded saved draft; generation slots are reserved. No model,
user reports, model credentials, remote servers or cloud access are needed.
"""
import copy
import json
import os
from pathlib import Path
import shutil
import sys
import threading

import anyio
from docx import Document
from mcp import Client as MCPClient, StdioServerParameters
import pytest

from briefloop.external_mcp import WorkspaceAdapter
from briefloop.store import Store


def parameters(root):
    repo = Path(__file__).resolve().parents[1]
    return StdioServerParameters(command=sys.executable,
                                args=['-m', 'briefloop', 'mcp', '--workspace', str(root)],
                                env={**os.environ, 'PYTHONPATH': str(repo / 'src')}, cwd=repo)


async def call(client, action, arguments=None, *, error=False):
    result = await client.call_tool('briefloop_' + action, arguments or {})
    assert bool(result.is_error) == error, result
    assert json.loads(result.content[0].text) == result.structured_content
    return result.structured_content


def test_unavailable_workspace_stdio_discovery_does_not_create_or_start(tmp_path):
    missing = tmp_path / 'absent'

    async def check():
        async with MCPClient(parameters(missing), cache=None) as client:
            assert (await call(client, 'discover'))['status'] == 'not_workspace'
            await call(client, 'inspect', error=True)
            await call(client, 'discover', {'workspace_id': 'replacement'}, error=True)
        assert not missing.exists()

    anyio.run(check)
    with pytest.raises(ValueError, match='绝对路径'):
        WorkspaceAdapter('relative-workspace')
    store = Store(tmp_path / 'offline')
    adapter = WorkspaceAdapter(store.root)
    assert adapter.discovery()['status'] == 'service_unavailable'
    assert not (store.root / 'server.json').exists()


def test_real_stdio_saved_draft_revision_word_and_idempotent_reconnection(tmp_path):
    from briefloop.document_model import markdown_document
    from briefloop.server import make_server

    root = tmp_path / 'synthetic-workspace'
    store = Store(root)
    # Admission requires a model selection; generation slots are held, so this
    # synthetic name is never invoked and no credential is configured or read.
    store.update_settings({'model': 'synthetic-not-invoked', 'model_selection_required': False})
    source = store.add_source('Synthetic MCP source', 'Order count 17; delivery in two batches.')
    requirements = {'title': 'Synthetic MCP report', 'objective': 'Summarize the supplied record',
                    'writing_mode': 'general', 'allow_web': False}
    run = store.create_run(requirements, [source['id']])
    document = markdown_document('Existing synthetic report.\n\n**Keep this emphasis.**\n\n| Metric | Count |\n| --- | --- |\n| Orders | 17 |')
    brief = store.publish(run['id'], {'title': requirements['title'], 'editor_document': document}, author='example')
    server = make_server(root, port=0, paused=True)
    # --paused only prevents recovery of OLD queued jobs. Hold the real model
    # budget before starting threads so NEW submissions also cannot run models.
    assert server.worker.budget.reserve('synthetic-no-model', server.worker.budget.limit)
    def forbidden_model_runtime():
        raise AssertionError('No model runtime may be constructed in this test')
    server.worker._report_runtime_factory = forbidden_model_runtime
    server.worker.start()
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    wid = store.meta('workspace_id')
    (root / 'server.json').write_text(json.dumps({'pid': os.getpid(),
        'url': f'http://127.0.0.1:{server.server_port}', 'workspace_id': wid}))

    async def check():
        submission = {'request_id': 'mcp-submit-1', 'requirements': requirements, 'source_ids': [source['id']]}
        # A true copy preserves both DB UUID and server.json. The original
        # server must not be mistaken for a service opened at the clone path.
        clone = tmp_path / 'copied-workspace'
        shutil.copytree(root, clone)
        cloned_store = Store(clone)
        assert cloned_store.meta('workspace_id') == wid
        original_only = store.add_source('Original-only source', 'Only saved in the original workspace.')
        original_briefs = store.rows('SELECT id FROM briefs')
        clone_briefs = cloned_store.rows('SELECT id FROM briefs')
        async with MCPClient(parameters(clone), cache=None) as client:
            discovery = await call(client, 'discover')
            assert not discovery['ready'] and discovery['status'] == 'identity_changed'
            assert '路径或身份不匹配' in discovery['message']
            for action, arguments in (
                ('inspect', {}), ('source', {'source_id': original_only['id']}),
                ('read', {'version_id': brief['id']}), ('submit', submission),
                ('revise', {'request_id': 'clone-revise', 'base_version': brief['id'], 'editor_document': document}),
                ('export', {'request_id': 'clone-export', 'version_id': brief['id']}),
                ('query', {'job_id': 'not-admitted'}),
                ('download', {'job_id': 'not-admitted', 'output': str(tmp_path / 'clone.docx')}),
            ):
                assert (await call(client, action, arguments, error=True))['code'] == 'workspace_unavailable'
        assert store.rows('SELECT id FROM briefs') == original_briefs
        assert cloned_store.rows('SELECT id FROM briefs') == clone_briefs
        assert not store.rows("SELECT id FROM jobs WHERE kind='generate'")
        # A real-directory alias is allowed after resolving both paths.
        alias = tmp_path / 'same-directory-alias'
        alias.symlink_to(root, target_is_directory=True)
        async with MCPClient(parameters(alias), cache=None) as client:
            assert (await call(client, 'discover'))['ready']
            assert (await call(client, 'source', {'source_id': original_only['id']}))['text'].startswith('Only saved')
        async with MCPClient(parameters(root), cache=None, read_timeout_seconds=15) as client:
            catalog = {tool.name: tool for tool in (await client.list_tools()).tools}
            assert len(catalog) == 9
            for name, tool in catalog.items():
                write = name in {'briefloop_submit', 'briefloop_revise', 'briefloop_export', 'briefloop_download'}
                assert tool.annotations.read_only_hint == (not write)
                assert tool.annotations.destructive_hint is False
                assert tool.annotations.idempotent_hint is True
                assert tool.annotations.open_world_hint == (name == 'briefloop_submit')
                assert tool.input_schema['additionalProperties'] is False
                assert 'workspace_id' not in tool.input_schema['properties']
            assert (await call(client, 'discover'))['ready']
            from briefloop.external_client import discover
            caps = discover(root)['capabilities']
            assert caps['workspace_path'] == str(root.resolve()) and caps['workspace_id'] == wid
            assert (await call(client, 'inspect'))['reports'][0]['version_id'] == brief['id']
            assert (await call(client, 'source', {'source_id': source['id']}))['text'].startswith('Order count 17')
            # The report task is admitted, but no model session is available.
            accepted = await call(client, 'submit', submission)
            assert accepted['status'] == 'accepted' and not accepted['replayed']
            state = await call(client, 'query', {'job_id': accepted['job_id']})
            assert state['status'] == 'queued' and state['run_id'] == accepted['run_id']
            assert (await call(client, 'submit', submission))['job_id'] == accepted['job_id']
            conflict = await call(client, 'submit', {**submission, 'source_ids': []}, error=True)
            assert conflict['code'] == 'conflict' and 'request_id' in conflict['message']
            before = await call(client, 'read', {'version_id': brief['id']})
            edited = copy.deepcopy(before['editor_document'])
            edited['content'][0]['content'][0]['text'] = 'Revised synthetic report.'
            revision = {'request_id': 'mcp-revise-1', 'base_version': brief['id'], 'editor_document': edited}
            saved = await call(client, 'revise', revision)
            assert saved['version_id'] != brief['id']
            assert (await call(client, 'revise', revision))['version_id'] == saved['version_id']
            conflict = await call(client, 'revise', {**revision, 'request_id': 'mcp-stale-base'}, error=True)
            assert conflict['code'] == 'conflict' and '稿件已有更新' in conflict['message']
            after = await call(client, 'read', {'version_id': saved['version_id']})
            assert after['parent_id'] == brief['id'] and after['editor_document']['content'][1:] == before['editor_document']['content'][1:]
            assert (await call(client, 'read', {'version_id': brief['id']})) == before
            export = {'request_id': 'mcp-export-1', 'version_id': saved['version_id']}
            exported = await call(client, 'export', export)
            assert exported['export_kind'] == 'working_draft'
            with anyio.fail_after(15):
                while True:
                    state = await call(client, 'query', {'job_id': exported['job_id']})
                    if state['terminal']:
                        break
                    await anyio.sleep(.1)
            assert state['status'] == 'complete' and state['artifact_available'], state
            target = tmp_path / 'synthetic-report.docx'
            downloaded = await call(client, 'download', {'job_id': exported['job_id'], 'output': str(target)})
            assert downloaded['version_id'] == saved['version_id'] and downloaded['sha256'] == state['sha256']
            word = Document(target)
            assert any('Revised synthetic report.' in p.text for p in word.paragraphs)
            assert word.tables[0].cell(1, 1).text == '17'
            assert (await call(client, 'download', {'job_id': exported['job_id'], 'output': str(target)}))['reused']
            occupied = tmp_path / 'occupied.docx'
            occupied.write_bytes(b'keep existing file')
            await call(client, 'download', {'job_id': exported['job_id'], 'output': str(occupied)}, error=True)
            assert occupied.read_bytes() == b'keep existing file'
            # Validation does not echo credential-shaped invalid inputs.
            invalid = await call(client, 'query', {'job_id': {'api_key': 'synthetic-secret'}}, error=True)
            assert 'synthetic-secret' not in json.dumps(invalid)
            # Inputs inside allowed object fields can reach Pydantic/server
            # errors. Never forward input_value or raw exception strings.
            for field, secret in (('api_key', 'synthetic-secret-api'),
                                  ('password', 'synthetic-secret-password'),
                                  ('authorization', 'Basic synthetic-secret-basic')):
                rejected = await call(client, 'submit', {
                    **submission, 'request_id': 'invalid-' + field,
                    'requirements': {**requirements, field: secret}}, error=True)
                assert rejected['code'] == 'invalid_request'
                assert secret not in json.dumps(rejected) and 'input_value' not in json.dumps(rejected)
                malformed = {'type': 'doc', 'content': [{'type': secret, 'attrs': {field: secret}}]}
                rejected = await call(client, 'revise', {
                    'request_id': 'invalid-editor-' + field, 'base_version': saved['version_id'],
                    'editor_document': malformed}, error=True)
                assert rejected['code'] == 'invalid_request'
                assert secret not in json.dumps(rejected) and 'input_value' not in json.dumps(rejected)
            server.draining = True
            assert not (await call(client, 'discover'))['ready']
            assert '尚未就绪' in (await call(client, 'inspect', error=True))['message']
            server.draining = False
            # Pin durable identity throughout this MCP process's lifetime.
            store.set_meta('workspace_id', 'replaced-workspace')
            assert (await call(client, 'discover'))['status'] == 'identity_changed'
            await call(client, 'inspect', error=True)
            store.set_meta('workspace_id', wid)
        # MCP teardown does not stop the backend or forget write receipts.
        assert server.worker.thread.is_alive()
        async with MCPClient(parameters(root), cache=None, mode='legacy') as client:
            replay = await call(client, 'submit', submission)
            assert replay['replayed'] and replay['job_id'] == accepted['job_id']
            assert (await call(client, 'revise', revision))['version_id'] == saved['version_id']
            assert (await call(client, 'export', export))['job_id'] == exported['job_id']
            # Safe error projection, without launching a model to manufacture failure.
            with store.tx() as connection:
                connection.execute("UPDATE jobs SET status='failed',error=? WHERE id=?",
                                   ('Authorization: Bearer synthetic-secret', accepted['job_id']))
            failed = await call(client, 'query', {'job_id': accepted['job_id']})
            assert failed['terminal'] and failed['status'] == 'failed'
            assert 'synthetic-secret' not in failed['error']
            # Missing file receipts are not silently replaced with new exports.
            from briefloop.export_jobs import output_path
            output_path(store, store.one('jobs', exported['job_id'])).unlink()
            assert (await call(client, 'export', export))['job_id'] == exported['job_id']
            assert not (await call(client, 'query', {'job_id': exported['job_id']}))['artifact_available']
            await call(client, 'download', {'job_id': exported['job_id'], 'output': str(tmp_path / 'missing.docx')}, error=True)
        assert not server.worker._generation_jobs
        assert not store.rows("SELECT id FROM jobs WHERE kind='generate' AND status='running'")

    try:
        anyio.run(check)
    finally:
        server.draining = False
        store.set_meta('workspace_id', wid)
        server.shutdown()
        thread.join()
        server.worker.close()
        server.harness.close()
        server.opencode_harness.close()
        server.runtime_bridge.close()
        server.server_close()
        server.workspace_lock.close()
