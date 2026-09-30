"""#858: a previous report's tracked changes become the first edits to classify."""
import json
from io import BytesIO

from docx import Document
from docx.oxml import parse_xml

from briefloop import revision_edits
from briefloop.learning import _eligible_cases
from briefloop.previous_report import docx_versions, import_previous
from briefloop.store import Store

NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def tracked_docx():
    doc = Document()
    doc.add_heading('组件价格周报', level=1)
    doc.add_heading('价格', level=2)
    p = doc.add_paragraph()
    p._p.append(parse_xml(f'<w:r {NS}><w:t xml:space="preserve">美国组件均价约每瓦</w:t></w:r>'))
    p._p.append(parse_xml(f'<w:del {NS} w:id="1" w:author="主管" w:date="2026-09-01T00:00:00Z"><w:r><w:delText>0.290</w:delText></w:r></w:del>'))
    p._p.append(parse_xml(f'<w:ins {NS} w:id="2" w:author="主管" w:date="2026-09-01T00:00:00Z"><w:r><w:t>0.310</w:t></w:r></w:ins>'))
    p._p.append(parse_xml(f'<w:r {NS}><w:t xml:space="preserve">美元。</w:t></w:r>'))
    doc.add_paragraph('东南亚价格保持平稳。')
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_tracked_changes_split_into_before_and_after():
    before, after, changed = docx_versions(tracked_docx())
    assert changed
    assert '# 组件价格周报' in before and '## 价格' in before
    assert '每瓦0.290美元' in before and '0.310' not in before
    assert '每瓦0.310美元' in after and '0.290' not in after


def test_import_creates_original_and_revision_feedback_but_no_comparison_case(tmp_path):
    store = Store(tmp_path / 'ws')
    result = import_previous(store, '周报.docx', tracked_docx())
    assert result['tracked_changes'] and result['title'] == '组件价格周报'
    briefs = store.rows('SELECT author,markdown FROM briefs WHERE run_id=? ORDER BY rowid', (result['run_id'],))
    assert [b['author'] for b in briefs] == ['import', 'user']
    feedback = store.rows("SELECT * FROM feedback WHERE kind='revision'")
    assert len(feedback) == 1
    items = revision_edits.pending(store, [feedback[0]['id']])
    assert [(e['before'], e['after']) for e in items[0]['edits']] == [('美国组件均价约每瓦0.290美元。', '美国组件均价约每瓦0.310美元。')]
    provenance = json.loads((store.root / 'sources' / (result['source_id'] + '.provenance.json')).read_text())
    assert provenance['usage'] == 'previous_report'
    cases, skipped = _eligible_cases(store, [result['run_id']])
    assert cases == [] and '往期导入' in skipped[0]['reason']
    assert store.meta('requirements') is None


def test_markdown_import_has_no_revision(tmp_path):
    store = Store(tmp_path / 'ws')
    result = import_previous(store, 'weekly.md', '# 周报\n\n正文。'.encode())
    assert not result['tracked_changes']
    assert not store.rows("SELECT * FROM feedback WHERE kind='revision'")


def test_source_text_of_a_tracked_docx_reads_the_accepted_document(tmp_path):
    store = Store(tmp_path / 'ws')
    result = import_previous(store, '周报.docx', tracked_docx())
    text = store.source_text(result['source_id'])
    assert '每瓦0.310美元' in text and '0.290' not in text
