"""Excel export behavior. Everything here passes with OfficeCLI absent: the
openpyxl base path is independently testable, and the optional enhancement
layer only ever degrades and records — it never blocks the artifact."""
import http.client
import inspect
import json
import threading

import pytest

from briefloop import host_bins, office_cli, xlsx_export
from briefloop.store import Store, dump
from briefloop.task_labels import LABELS, REPORTED

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


@pytest.fixture(autouse=True)
def _no_officecli(tmp_path, monkeypatch):
    """A faked PATH with no officecli anywhere: the base path must not care."""
    monkeypatch.setenv('PATH', str(tmp_path / 'nowhere-bin'))
    monkeypatch.setattr(host_bins, 'EXTRA_DIRS', ())
    office_cli._clear_caches()
    yield
    office_cli._clear_caches()


def _para(text):
    return {'type': 'paragraph', 'content': [{'type': 'text', 'text': text}]}


def _cell(text, kind='tableCell', **attrs):
    node = {'type': kind, 'content': [_para(text)]}
    if attrs:
        node['attrs'] = attrs
    return node


def _row(*cells):
    return {'type': 'tableRow', 'content': list(cells)}


def _table(*rows):
    return {'type': 'table', 'content': list(rows)}


def _heading(text, level=2):
    return {'type': 'heading', 'attrs': {'level': level}, 'content': [{'type': 'text', 'text': text}]}


def _store(tmp_path):
    return Store(tmp_path / 'workspace')


def _brief(store, document, *, title='示例报告', requirements=None):
    source = store.add_source('Synthetic', 'Revenue was USD 12 million.')
    run = store.create_run({'title': title, 'objective': 'Explain revenue', 'allow_web': False,
                            **(requirements or {})}, [source['id']])
    return store.publish(run['id'], {'title': title, 'markdown': 'Revenue.',
                                     'editor_document': {'type': 'doc', 'content': list(document)}})


def _generate(store, job):
    result = xlsx_export.generate_xlsx(store, job, threading.Event())
    with store.tx() as c:
        c.execute("UPDATE jobs SET status='complete',result=? WHERE id=?", (dump(result), job['id']))
    return result


def _queued_xlsx(store, brief, layout='sheets'):
    return store.one('jobs', xlsx_export.enqueue_export_xlsx(store, brief['id'], layout)['id'])


def _sales_document():
    # 合计行有两个数值格（1254、7.5），满足合计行判定；千分位/前导零/小数覆盖数值化守则。
    return [_heading('销售明细'),
            _table(_row(_cell('品名', 'tableHeader'), _cell('销量', 'tableHeader'), _cell('编号', 'tableHeader')),
                   _row(_cell('甲'), _cell('1,234'), _cell('007')),
                   _row(_cell('乙'), _cell('20'), _cell('0.5')),
                   _row(_cell('合计'), _cell('1,254'), _cell('7.5')))]


def _open(tmp_path, result, data_only=False):
    from openpyxl import load_workbook
    return load_workbook(tmp_path / 'workspace' / result['path'], data_only=data_only)


def test_base_render_typing_merges_styles_and_result_record(tmp_path):
    store = _store(tmp_path)
    document = [_heading('区域对比：华东'),
                _table(_row(_cell('地区', 'tableHeader', rowspan=2, backgroundColor='#e8f0f8'),
                            _cell('指标', 'tableHeader')),
                       _row(_cell('销量', textAlign='center'))),
                _table(_row(_cell('列', 'tableHeader', colwidth=[140]), _cell('值', 'tableHeader')),
                       _row(_cell('长串'), _cell('12345678901234567890')),
                       _row(_cell('多行', 'tableHeader'), _cell('第一行\n第二行')))]
    brief = _brief(store, document)
    job = _queued_xlsx(store, brief)
    result = _generate(store, job)
    assert result['path'] == f'exports/{job["id"]}/report.xlsx'
    assert result['download_url'] == '/api/export-file?job=' + job['id']
    assert result['layout'] == 'sheets' and result['enhanced'] is False
    import hashlib
    assert result['sha256'] == hashlib.sha256((tmp_path / 'workspace' / result['path']).read_bytes()).hexdigest()
    assert 'office' not in result  # switch off, binary absent: no rows, no events
    workbook = _open(tmp_path, result)
    assert workbook.sheetnames == ['目录', '区域对比-华东', '区域对比-华东-2']  # both tables share the heading
    merged = workbook['区域对比-华东']
    assert [str(span) for span in merged.merged_cells.ranges] == ['A2:A3']
    header = merged['A2']
    assert header.font.bold is True and header.fill.fill_type == 'solid'
    assert 'E8F0F8' in str(header.fill.start_color.rgb).upper()
    assert merged['B2'].font.bold is True
    assert merged['B3'].alignment.horizontal == 'center' and merged['B3'].alignment.vertical == 'top'
    second = workbook['区域对比-华东-2']
    assert second['B3'].value == '12345678901234567890'  # 20 significant digits stay text
    assert second['B3'].number_format == 'General'
    assert second['A4'].font.bold is True  # a tableHeader outside the first row still renders bold
    multiline = second['B4']
    assert multiline.value == '第一行\n第二行' and multiline.alignment.wrap_text is True
    assert second.column_dimensions['A'].width == round(140 / 7, 2)
    sales = _open(tmp_path, _generate(store, _queued_xlsx(store, _brief(store, _sales_document(), title='销售'))))['销售明细']
    assert sales['B3'].value == 1234 and sales['B3'].number_format == 'General'  # 基础态：值保留、显示丢千分位
    assert sales['C3'].value == '007' and sales['C4'].value == 0.5
    assert sales['B5'].value == 1254  # base artifact: the reported value, never a formula


def test_sheet_names_are_sanitized_deduped_and_indexed_exactly(tmp_path):
    store = _store(tmp_path)
    long_title = '市场分析' * 10  # 40 characters
    pair = (_row(_cell('a', 'tableHeader')), _row(_cell('1')))
    document = [_table(*pair),                                   # no heading → 表1
                _heading('市场分析：华东区'), _table(*pair), _table(*pair),  # same heading twice
                _heading(long_title), _table(*pair)]
    brief = _brief(store, document)
    result = _generate(store, _queued_xlsx(store, brief))
    workbook = _open(tmp_path, result)  # openpyxl itself rejects illegal titles on load
    names = workbook.sheetnames
    assert names == ['目录', '表1', '市场分析-华东区', '市场分析-华东区-2', long_title[:28]]
    assert all(len(name) <= 31 for name in names)
    index = workbook['目录']
    assert index['A1'].value == '示例报告' and index['A2'].value is None  # no report_date: left empty
    assert [index.cell(row=3, column=c).value for c in (1, 2, 3)] == ['序号', '表格标题', '工作表']
    listed = [[index.cell(row=r, column=c).value for c in (1, 2, 3)] for r in range(4, 8)]
    assert [row[0] for row in listed] == [1, 2, 3, 4]
    assert [row[1] for row in listed] == ['表1', '市场分析：华东区', '市场分析：华东区', long_title]
    assert [row[2] for row in listed] == names[1:]  # index matches the real sheets exactly


def test_report_date_lands_in_the_index(tmp_path):
    store = _store(tmp_path)
    brief = _brief(store, _sales_document(), requirements={'report_date': '2026-09-01'})
    result = _generate(store, _queued_xlsx(store, brief))
    assert _open(tmp_path, result)['目录']['A2'].value == '2026-09-01'


@pytest.mark.parametrize('layout', ['sheets', 'single'])
def test_base_export_fits_cjk_and_wraps_prose_without_changing_cells(tmp_path, layout):
    store = _store(tmp_path)
    prose = '本季度收入同比增长，仍需结合回款与客户结构判断增长质量。' * 5
    title = '合成经营简报：指标与后续观察'
    document = [_heading(title),
                _table(_row(_cell('业务指标', 'tableHeader'), _cell('观察说明', 'tableHeader')),
                       _row(_cell('经营活动现金流量净额'), _cell(prose)),
                       _row(_cell('手动宽度', colwidth=[140]), _cell('第一行\n第二行')))]
    result = _generate(store, _queued_xlsx(store, _brief(store, document), layout))
    workbook = _open(tmp_path, result)
    sheet = workbook[title.replace('：', '-')] if layout == 'sheets' else workbook['示例报告']
    assert sheet['A3'].value == '经营活动现金流量净额'
    assert sheet['B3'].value == prose
    assert sheet.column_dimensions['A'].width == 20  # authored width retained
    assert 30 <= sheet.column_dimensions['B'].width <= 60
    assert sheet['B3'].alignment.wrap_text and sheet.row_dimensions[3].height > 30
    assert sheet['B4'].value == '第一行\n第二行' and sheet.row_dimensions[4].height >= 34
    if layout == 'sheets':
        index = workbook['目录']
        assert index['B4'].value == title and index['B4'].alignment.wrap_text
        assert index.column_dimensions['B'].width == 30
        assert index.row_dimensions[4].height >= 34  # WPS wraps the last CJK glyph


@pytest.mark.parametrize('layout', ['sheets', 'single'])
def test_colspan_and_mixed_header_merges_export_in_both_layouts(tmp_path, layout):
    store = _store(tmp_path)
    document = [_heading('横向表头'),
                _table(_row(_cell('指标', 'tableHeader', colspan=2)),
                       _row(_cell('a'), _cell('1'))),
                _heading('混合表头'),
                _table(_row(_cell('地区', 'tableHeader', rowspan=2), _cell('指标', 'tableHeader', colspan=2)),
                       _row(_cell('量', 'tableHeader'), _cell('额', 'tableHeader')),
                       _row(_cell('华东'), _cell('2'), _cell('3')))]
    result = _generate(store, _queued_xlsx(store, _brief(store, document), layout))
    workbook = _open(tmp_path, result)
    first = workbook['横向表头'] if layout == 'sheets' else workbook['示例报告']
    mixed = workbook['混合表头'] if layout == 'sheets' else first
    assert 'A2:B2' in {str(span) for span in first.merged_cells.ranges}
    top = 2 if layout == 'sheets' else 6
    assert {f'A{top}:A{top+1}', f'B{top}:C{top}'} <= {str(span) for span in mixed.merged_cells.ranges}
    assert first['A2'].value == '指标' and first['A2'].font.bold
    assert mixed.cell(top, 2).value == '指标' and mixed.cell(top, 2).fill.fill_type == 'solid'
    assert mixed.cell(top+2, 3).value == 3


@pytest.mark.parametrize('layout', ['sheets', 'single'])
def test_equals_prefixed_report_text_is_not_an_excel_formula(tmp_path, layout):
    store = _store(tmp_path)
    document = [_heading('=Data'), _table(_row(_cell('=Header', 'tableHeader')),
                                        _row(_cell('=1+1')), _row(_cell('#N/A')))]
    result = _generate(store, _queued_xlsx(store, _brief(store, document, title='=Report'), layout))
    workbook = _open(tmp_path, result)
    sheet = workbook['=Data'] if layout == 'sheets' else workbook['=Report']
    for coordinate, text in [('A1', '=Data'), ('A2', '=Header'), ('A3', '=1+1'), ('A4', '#N/A')]:
        assert sheet[coordinate].value == text and sheet[coordinate].data_type == 's'
    if layout == 'sheets':
        for coordinate, text in [('A1', '=Report'), ('B4', '=Data'), ('C4', '=Data')]:
            assert workbook['目录'][coordinate].value == text and workbook['目录'][coordinate].data_type == 's'


def test_sheet_name_case_collisions_match_index_and_enhancement_paths(tmp_path):
    store = _store(tmp_path)
    table = _table(_row(_cell('值', 'tableHeader')), _row(_cell('1,234')))
    document = [_heading('Sales'), table, _heading('sales'), table, _heading('SALES-2'), table]
    result = _generate(store, _queued_xlsx(store, _brief(store, document)))
    workbook = _open(tmp_path, result)
    assert workbook.sheetnames == ['目录', 'Sales', 'sales-2', 'SALES-2-2']
    assert [workbook['目录'].cell(row, 3).value for row in range(4, 7)] == workbook.sheetnames[1:]
    plan = xlsx_export.plan_enhancements(xlsx_export._layout({'type':'doc','content':document}, 'sheets'))
    assert {item['path'].split('/')[1] for item in plan['commands']} == set(workbook.sheetnames[1:])


@pytest.mark.parametrize('layout', ['sheets', 'single'])
def test_english_excel_labels_preserve_authored_text(tmp_path, layout):
    store = _store(tmp_path)
    document = [_table(_row(_cell('原文列名', 'tableHeader')), _row(_cell('1,234'))),
                _heading('原文表名'), _table(_row(_cell('Value', 'tableHeader')), _row(_cell('25%')))]
    brief = _brief(store, document, title='原文报告标题', requirements={'language':'en'})
    result = _generate(store, _queued_xlsx(store, brief, layout))
    workbook = _open(tmp_path, result)
    if layout == 'sheets':
        assert workbook.sheetnames == ['Contents', 'Table 1', '原文表名']
        index = workbook['Contents']
        assert index['A1'].value == '原文报告标题'
        assert [index.cell(3, c).value for c in (1, 2, 3)] == ['No.', 'Table title', 'Worksheet']
        assert [index.cell(r, 3).value for r in (4, 5)] == ['Table 1', '原文表名']
        assert workbook['Table 1']['A2'].value == '原文列名'
    else:
        assert workbook.sheetnames == ['原文报告标题']
        sheet = workbook.active
        assert sheet['A1'].value == 'Table 1' and sheet['A2'].value == '原文列名'
        assert sheet['A5'].value == '原文表名' and sheet['A6'].value == 'Value'
    models = xlsx_export._layout({'type':'doc','content':document}, layout,
                                 report_title='原文报告标题', language='en')
    plan = xlsx_export.plan_enhancements(models)
    assert {command['path'].split('/')[1] for command in plan['commands']} <= set(workbook.sheetnames)


def test_freeze_panes_follow_the_layout(tmp_path):
    store = _store(tmp_path)
    header_table = _table(_row(_cell('a', 'tableHeader')), _row(_cell('1')))
    plain_table = _table(_row(_cell('a')), _row(_cell('1')))
    brief = _brief(store, [_heading('有表头'), header_table, plain_table])
    workbook = _open(tmp_path, _generate(store, _queued_xlsx(store, brief)))
    assert workbook['有表头'].freeze_panes == 'A3'    # title row 1 + header row 2
    assert workbook['有表头-2'].freeze_panes == 'A2'  # no header row: below the title
    single = _generate(store, _queued_xlsx(
        store, _brief(store, [_heading('单表'), header_table, plain_table], title='单表版式'), 'single'))
    workbook = _open(tmp_path, single)
    assert workbook.sheetnames == ['单表版式']
    sheet = workbook['单表版式']
    assert sheet.freeze_panes is None               # one pane per sheet: nothing frozen
    assert sheet['A1'].value == '单表' and sheet['A2'].value == 'a'   # title row before each table
    assert sheet['A5'].value == '单表' and sheet['A6'].value == 'a'


def test_report_without_tables_is_rejected_before_queueing(tmp_path):
    store = _store(tmp_path)
    brief = _brief(store, [_para('纯文字简报，没有任何表格。')])
    with pytest.raises(ValueError, match='报告没有可导出的表格'):
        xlsx_export.enqueue_export_xlsx(store, brief['id'])
    assert store.rows("SELECT * FROM jobs WHERE kind='export_xlsx'") == []
    assert list((tmp_path / 'workspace' / 'exports').rglob('*.xlsx')) == []
    # Generate-time guard: a document whose tables vanished in place after
    # queueing gets the explicit message, not a bare fingerprint complaint.
    tabular = _brief(store, [_table(_row(_cell('a', 'tableHeader')), _row(_cell('1')))])
    job = _queued_xlsx(store, tabular)
    with store.tx() as c:
        c.execute("UPDATE briefs SET editor_document=? WHERE id=?",
                  (dump({'type': 'doc', 'content': [_para('表格已删除。')]}), tabular['id']))
    with pytest.raises(ValueError, match='报告没有可导出的表格'):
        xlsx_export.generate_xlsx(store, job, threading.Event())


def test_changed_inputs_fail_the_fingerprint_check(tmp_path):
    store = _store(tmp_path)
    brief = _brief(store, _sales_document())
    job = _queued_xlsx(store, brief)
    with store.tx() as c:  # the queued version row itself moved under the job
        c.execute("UPDATE briefs SET editor_document=? WHERE id=?",
                  (dump({'type': 'doc', 'content': _sales_document() + [_para('补充说明。')]}), brief['id']))
    with pytest.raises(ValueError, match='导出输入已变化'):
        xlsx_export.generate_xlsx(store, job, threading.Event())


def test_fingerprint_dedup_covers_layout_enhanced_flag_and_stale_files(tmp_path, monkeypatch):
    store = _store(tmp_path)
    brief = _brief(store, _sales_document())
    first = xlsx_export.enqueue_export_xlsx(store, brief['id'])
    assert json.loads(first['payload'])['enhanced'] is False
    assert xlsx_export.enqueue_export_xlsx(store, brief['id'])['id'] == first['id']  # queued reuse
    _generate(store, store.one('jobs', first['id']))
    assert xlsx_export.enqueue_export_xlsx(store, brief['id'])['id'] == first['id']  # complete + intact
    assert xlsx_export.enqueue_export_xlsx(store, brief['id'], 'single')['id'] != first['id']
    monkeypatch.setattr(office_cli, 'enhancement_ready', lambda store: True)
    enhanced = xlsx_export.enqueue_export_xlsx(store, brief['id'])
    assert json.loads(enhanced['payload'])['enhanced'] is True
    assert enhanced['id'] != first['id']  # flipping the switch rebuilds, never reuses the base cache
    monkeypatch.setattr(office_cli, 'enhancement_ready', lambda store: False)
    xlsx_export.output_path_xlsx(store, store.one('jobs', first['id'])).unlink()
    rebuilt = xlsx_export.enqueue_export_xlsx(store, brief['id'])
    assert rebuilt['id'] != first['id']  # complete cache with no file cannot be reused
    _generate(store, store.one('jobs', rebuilt['id']))
    xlsx_export.output_path_xlsx(store, store.one('jobs', rebuilt['id'])).write_bytes(b'changed on disk')
    assert xlsx_export.enqueue_export_xlsx(store, brief['id'])['id'] != rebuilt['id']


def test_output_path_xlsx_rejects_wrong_kind_and_suffix(tmp_path):
    store = _store(tmp_path)
    brief = _brief(store, _sales_document())
    job = _queued_xlsx(store, brief)
    _generate(store, job)
    assert xlsx_export.output_path_xlsx(store, job).name == 'report.xlsx'
    with pytest.raises(ValueError, match='不是 Excel 文件任务'):
        xlsx_export.output_path_xlsx(store, {'kind': 'export_docx', 'id': job['id'], 'result': None})
    for path in (f'exports/{job["id"]}/report.docx', f'../outside/{job["id"]}/report.xlsx'):
        with store.tx() as c:
            c.execute("UPDATE jobs SET result=? WHERE id=?", (dump({'path': path}), job['id']))
        with pytest.raises(ValueError, match='无效导出结果路径'):
            xlsx_export.output_path_xlsx(store, store.one('jobs', job['id']))


def test_plan_enhancements_is_pure_and_picky(tmp_path):
    document = [_heading('销售明细'),
                _table(_row(_cell('品名', 'tableHeader'), _cell('销量', 'tableHeader'), _cell('金额', 'tableHeader')),
                       _row(_cell('甲'), _cell('10'), _cell('1,234')),
                       _row(_cell('乙'), _cell('20'), _cell('266')),
                       _row(_cell('合计'), _cell('30'), _cell('$1,500'))),
                _table(_row(_cell('x', 'tableHeader'), _cell('y', 'tableHeader')),
                       _row(_cell('1'), _cell('2')),
                       _row(_cell('3'), _cell('4')))]
    models = xlsx_export._layout({'type': 'doc', 'content': document}, 'sheets', report_title='销售')
    plan = xlsx_export.plan_enhancements(models)
    sets = {item['path']: item['props'] for item in plan['commands'] if item['command'] == 'set'}
    assert sets['/销售明细/B5'] == {'formula': 'SUM(B3:B4)'}  # plain totals keep General
    assert sets['/销售明细/C5'] == {'formula': 'SUM(C3:C4)', 'numberformat': '"$"#,##0'}
    assert sets['/销售明细/C3'] == {'numberformat': '#,##0'}
    assert [f['cell'] for f in plan['formulas']] == ['B5', 'C5']
    assert [f['expected'] for f in plan['formulas']] == [30, 1500]
    bars = {(item['parent'], item['props']['ref']) for item in plan['commands'] if item.get('command') == 'add'}
    assert ('/销售明细', 'B3:B5') in bars and ('/销售明细', 'C3:C5') in bars
    assert all(parent == '/销售明细' for parent, _ in bars)  # 2 numeric cells in 表2: no data bar
    assert not any(path.startswith('/表2/') for path in sets)
    guarded = xlsx_export.plan_enhancements(xlsx_export._layout(
        {'type': 'doc', 'content': [_table(_row(_cell('编号', 'tableHeader'), _cell('值', 'tableHeader')),
                                           _row(_cell('007'), _cell('12345678901234567890')),
                                           _row(_cell('小计'), _cell('008')))]}, 'sheets'))
    assert guarded['commands'] == []  # leading zeros and over-long digits never enter the numeric set


def test_enhancement_and_check_degrade_when_officecli_cannot_run(tmp_path, monkeypatch):
    store = _store(tmp_path)
    store.update_settings({'officecli_enabled': True})
    assert office_cli.enhancement_ready(store) is False  # switch on, binary absent
    brief = _brief(store, _sales_document())
    job = _queued_xlsx(store, brief)
    assert json.loads(job['payload'])['enhanced'] is False
    result = _generate(store, job)
    assert result['enhanced'] is False and 'enhance_reason' not in result
    assert (tmp_path / 'workspace' / result['path']).is_file()
    assert store.rows('SELECT * FROM office_checks') == []

    # Binary "found" but every subprocess fails: still complete, still recorded.
    monkeypatch.setattr(office_cli, 'find', lambda: '/nonexistent/officecli')
    monkeypatch.setattr(office_cli, 'run_json',
                        lambda *args, **kwargs: {'ok': False, 'data': None, 'reason': 'synthetic officecli outage'})
    outage_job = _queued_xlsx(store, _brief(store, _sales_document(), title='第二次'), 'single')
    outage = _generate(store, outage_job)
    assert outage['enhanced'] is False and 'synthetic officecli outage' in outage['enhance_reason']
    assert (tmp_path / 'workspace' / outage['path']).is_file()
    assert outage['office']['validate']['status'] == 'error'
    assert outage['office']['issues']['status'] == 'error'
    assert [row['kind'] for row in store.rows(
        "SELECT kind FROM events WHERE job_id=? AND kind='export_xlsx_enhancement_failed'", (outage_job['id'],))] \
        == ['export_xlsx_enhancement_failed']


def _apply_planned_sets(path, plan):
    """Controlled replacement for the optional CLI; real XLSX read/write."""
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    for item in plan['commands']:
        if item['command'] != 'set': continue
        sheet, coordinate = item['path'].strip('/').split('/')
        cell, props = workbook[sheet][coordinate], item['props']
        if 'value' in props: cell.value = props['value']
        if 'formula' in props: cell.value = '=' + props['formula']
        if 'numberformat' in props: cell.number_format = props['numberformat']
    workbook.save(path)
    workbook.close()


@pytest.mark.parametrize('failure,with_formulas', [
    ('exception', False), ('corrupt', False), ('corrupt', True),
    ('missing_sheet', False), ('changed_text', False), ('replace_locked', False),
])
def test_invalid_enhancement_preserves_complete_downloadable_base(tmp_path, monkeypatch, failure, with_formulas):
    from openpyxl import load_workbook
    store = _store(tmp_path)
    document = _sales_document() if with_formulas else [
        _heading('比例'), _table(_row(_cell('类别', 'tableHeader'), _cell('值', 'tableHeader')),
                              _row(_cell('甲'), _cell('25%')))]
    monkeypatch.setattr(office_cli, 'enhancement_ready', lambda _: True)
    job = _queued_xlsx(store, _brief(store, document))
    captured = {}

    def enhance(_store, path, plan, **kwargs):
        captured['base'] = path.read_bytes()
        assert bool(plan['formulas']) == with_formulas
        if failure == 'exception': raise OSError('controlled renderer failure')
        if failure == 'corrupt': path.write_bytes(b'not a workbook')
        else:
            _apply_planned_sets(path, plan)
            workbook = load_workbook(path)
            if failure == 'missing_sheet': workbook.remove(workbook[workbook.sheetnames[-1]])
            if failure == 'changed_text': workbook[workbook.sheetnames[-1]]['A3'] = 'unauthorized rewrite'
            workbook.save(path)
            workbook.close()
        return {'applied': True, 'reason': None}

    monkeypatch.setattr(office_cli, 'enhance_workbook', enhance)
    real_replace = xlsx_export.os.replace
    def replace(src, dst):
        if failure == 'replace_locked' and str(src).endswith('.staging.xlsx'):
            raise PermissionError('controlled file lock')
        return real_replace(src, dst)
    monkeypatch.setattr(xlsx_export.os, 'replace', replace)
    result = _generate(store, job)
    assert result['enhanced'] is False and result['enhance_reason']
    assert store.one('jobs', job['id'])['status'] == 'complete'
    path = store.root / result['path']
    assert path.read_bytes() == captured['base']
    assert _open(tmp_path, result).sheetnames == (['目录', '销售明细'] if with_formulas else ['目录', '比例'])
    assert not path.with_name('report.staging.xlsx').exists()
    assert store.rows("SELECT * FROM events WHERE job_id=? AND kind='export_xlsx_enhancement_failed'", (job['id'],))


def test_valid_enhancement_without_formulas_is_verified_and_published(tmp_path, monkeypatch):
    store = _store(tmp_path)
    document = [_heading('比例'), _table(_row(_cell('类别', 'tableHeader'), _cell('值', 'tableHeader')),
                                      _row(_cell('甲'), _cell('25%')))]
    monkeypatch.setattr(office_cli, 'enhancement_ready', lambda _: True)
    def enhance(_store, path, plan, **kwargs):
        assert plan['formulas'] == []
        _apply_planned_sets(path, plan)
        return {'applied': True, 'reason': None}
    monkeypatch.setattr(office_cli, 'enhance_workbook', enhance)
    result = _generate(store, _queued_xlsx(store, _brief(store, document)))
    assert result['enhanced'] is True
    cell = _open(tmp_path, result)['比例']['B3']
    assert cell.value == .25 and cell.number_format == '0%'


@pytest.mark.parametrize('wrong_cache', [False, True])
def test_planned_formula_and_cached_result_are_both_checked(tmp_path, monkeypatch, wrong_cache):
    from io import BytesIO
    from zipfile import ZipFile
    import xml.etree.ElementTree as ET
    store = _store(tmp_path)
    document = [_heading('合计检查'),
                _table(_row(_cell('项目', 'tableHeader'), _cell('量', 'tableHeader'), _cell('额', 'tableHeader')),
                       _row(_cell('甲'), _cell('10'), _cell('100')),
                       _row(_cell('乙'), _cell('20'), _cell('200')),
                       _row(_cell('合计'), _cell('30'), _cell('300')))]
    monkeypatch.setattr(office_cli, 'enhancement_ready', lambda _: True)
    captured = {}
    def enhance(_store, path, plan, **kwargs):
        captured['base'] = path.read_bytes()
        _apply_planned_sets(path, plan)
        # openpyxl deliberately does not evaluate formulas. Populate the two
        # known synthetic sums in OOXML, as a renderer would, without changing
        # formulas; the bad fixture changes one cached value only.
        src, dst = BytesIO(path.read_bytes()), BytesIO()
        ns = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        with ZipFile(src) as zin, ZipFile(dst, 'w') as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == 'xl/worksheets/sheet2.xml':
                    tree = ET.fromstring(data)
                    for address, value in [('B5', 31 if wrong_cache else 30), ('C5', 300)]:
                        cell = tree.find(f'.//s:c[@r="{address}"]', ns)
                        assert cell.find('s:f', ns) is not None
                        cell.find('s:v', ns).text = str(value)
                    data = ET.tostring(tree)
                zout.writestr(item, data)
        path.write_bytes(dst.getvalue())
        return {'applied': True, 'reason': None}
    monkeypatch.setattr(office_cli, 'enhance_workbook', enhance)
    job = _queued_xlsx(store, _brief(store, document))
    result = _generate(store, job)
    assert result['enhanced'] is not wrong_cache
    if wrong_cache:
        assert result['enhance_reason'] == '公式结果与报告数值不一致'
        assert (store.root/result['path']).read_bytes() == captured['base']
        assert store.rows("SELECT * FROM events WHERE job_id=? AND kind='export_xlsx_formula_mismatch'", (job['id'],))
    else:
        assert _open(tmp_path, result)['合计检查']['B5'].value == '=SUM(B3:B4)'
        assert _open(tmp_path, result, data_only=True)['合计检查']['B5'].value == 30


def test_excel_cancellation_is_preserved_during_enhancement_and_quality_check(tmp_path, monkeypatch):
    store = _store(tmp_path)
    store.update_settings({'officecli_enabled': True})
    cancelled, calls = threading.Event(), []
    monkeypatch.setattr(office_cli, 'find', lambda: 'synthetic-officecli')
    def run(args, **kwargs):
        calls.append(args[1])
        cancelled.set()
        return {'ok': True, 'data': {'summary': {'failed': 0}}, 'reason': None}
    monkeypatch.setattr(office_cli, 'run_json', run)
    with pytest.raises(InterruptedError):
        office_cli.enhance_workbook(store, tmp_path/'staging.xlsx', {'commands':[{'command':'set'}]}, cancelled=cancelled)
    assert calls == ['batch']  # do not start close or silently accept cancellation

    cancelled.clear()
    monkeypatch.setattr(office_cli, 'enhancement_ready', lambda _: False)
    def check(_store, _path, **kwargs):
        assert kwargs['cancelled'] is cancelled
        cancelled.set()
        return None
    monkeypatch.setattr(office_cli, 'check_file', check)
    with pytest.raises(InterruptedError):
        xlsx_export.generate_xlsx(store, _queued_xlsx(store, _brief(store, _sales_document())), cancelled)


def test_labels_and_notifications_cover_the_new_kind(tmp_path):
    assert LABELS['export_xlsx'] == '生成报表 Excel' and 'export_xlsx' in REPORTED
    store = _store(tmp_path)
    job = _queued_xlsx(store, _brief(store, _sales_document()))
    _generate(store, job)
    from briefloop.notifications import job_status
    job_status(store, store.one('jobs', job['id']), 'complete')
    assert any('生成报表 Excel 已完成' in row['title'] for row in store.rows('SELECT * FROM notifications'))


def test_runtime_dispatch_and_placeholders(tmp_path):
    from briefloop import runtime
    kinds = runtime.FILE_JOB_KINDS
    assert 'export_xlsx' in kinds
    for name in ('loop', 'file_loop'):
        source = inspect.getsource(getattr(runtime.Worker, name))
        # Placeholders are derived from the tuple itself, never hand-counted:
        # a fixed "?,?,?" once desynced from the tuple and killed the worker
        # threads with sqlite3.ProgrammingError.
        assert '",".join("?"*len(FILE_JOB_KINDS))' in source, name
        assert '(?,?,?)' not in source, name

    store = _store(tmp_path)
    brief = _brief(store, _sales_document())
    worker = runtime.Worker(store)
    errors = {}

    def run(name, target):
        try:
            target()
        except BaseException as exc:  # a dead dispatch thread is the failure mode
            errors[name] = exc

    threads = [threading.Thread(target=run, args=('loop', worker.loop), daemon=True),
               threading.Thread(target=run, args=('file', worker.file_loop), daemon=True)]
    queued = xlsx_export.enqueue_export_xlsx(store, brief['id'])
    for thread in threads:
        thread.start()
    try:
        for _ in range(100):
            job = store.one('jobs', queued['id'])
            if job['status'] == 'complete':
                break
            threading.Event().wait(0.1)
        assert job['status'] == 'complete', job['error']  # ran as generate_xlsx, not audit_bundle
        assert json.loads(job['result'])['layout'] == 'sheets'
        assert (tmp_path / 'workspace' / 'exports' / job['id'] / 'report.xlsx').is_file()
        progress = [json.loads(row['data']) for row in store.rows(
            "SELECT data FROM events WHERE kind='export_progress' AND job_id=? ORDER BY seq", (job['id'],))]
        assert progress == [{'step': 1, 'total': 4, 'message': '读取已保存的报告版本'},
                            {'step': 2, 'total': 4, 'message': '展开表格并生成工作簿'},
                            {'step': 3, 'total': 4, 'message': '写入 Excel 文件'},
                            {'step': 4, 'total': 4, 'message': 'Excel 已生成，可以下载'}]
    finally:
        worker.stopping.set()
        worker.wake()
        for thread in threads:
            thread.join(timeout=10)
        if store._job_wakeup == worker.wake:
            store._job_wakeup = None
    assert errors == {}


def _server(tmp_path):
    from briefloop.server import make_server
    server = make_server(tmp_path / 'http-workspace', port=0, paused=True)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()

    def request(path, body=None):
        connection = http.client.HTTPConnection('127.0.0.1', server.server_port)
        headers = {}
        if body is not None:  # POSTs carry the page session token
            connection.request('GET', '/api/session')
            token = json.loads(connection.getresponse().read())['token']
            headers = {'X-BriefLoop-Token': token, 'Content-Type': 'application/json'}
        connection.request('POST' if body is not None else 'GET', path,
                           body=json.dumps(body) if body is not None else None, headers=headers)
        response = connection.getresponse()
        payload = response.read()
        response_headers = dict(response.getheaders())
        connection.close()
        return response.status, payload, response_headers

    return server, thread, request


def test_server_surface_accepts_xlsx_end_to_end(tmp_path):
    server, thread, request = _server(tmp_path)
    try:
        store = server.store
        brief = _brief(store, _sales_document())
        status, payload, _ = request('/api/export-xlsx', {'version_id': brief['id']})
        assert status == 200
        job = json.loads(payload)
        assert job['kind'] == 'export_xlsx' and job['status'] == 'queued'
        assert json.loads(job['payload'])['layout'] == 'sheets'
        for body, message in (({'version_id': brief['id'], 'layout': 'grid'}, '未知 Excel 版式'),
                              ({'version_id': 'brief_missing'}, 'Record not found')):
            status, payload, _ = request('/api/export-xlsx', body)
            assert status == 400 and message in payload.decode(), (status, payload)
        plain = _brief(store, [_para('无表格简报。')])
        status, payload, _ = request('/api/export-xlsx', {'version_id': plain['id']})
        assert status == 400 and '报告没有可导出的表格' in payload.decode()

        result = _generate(store, store.one('jobs', job['id']))
        status, payload, _ = request(f'/api/export-status?job={job["id"]}')
        assert status == 200 and json.loads(payload)['kind'] == 'export_xlsx'
        other = store.enqueue('assess', {'version_id': brief['id']})
        status, payload, _ = request(f'/api/export-status?job={other["id"]}')
        assert status == 400 and '不是导出任务' in payload.decode()

        workspace_id = store.meta('workspace_id')
        status, payload, headers = request(f'/api/export-file?job={job["id"]}&workspace_id={workspace_id}')
        assert status == 200
        assert headers['Content-Type'] == XLSX_MIME
        assert headers['Content-Disposition'].endswith('.xlsx')
        assert payload[:2] == b'PK'  # a real xlsx zip
        status, _, _ = request(f'/api/export-file?job={job["id"]}&workspace_id=other')
        assert status == 409
        from briefloop.export_jobs import enqueue_export
        docx_job = store.one('jobs', enqueue_export(store, brief['id'])['id'])
        status, payload, _ = request(f'/api/export-file?job={docx_job["id"]}&workspace_id={workspace_id}')
        assert status == 400 and 'Word 尚未制作完成' in payload.decode()

        # Explicit office endpoints: the kind gate lets export_xlsx through and
        # the request then stops at the binary gate (no officecli in this PATH).
        store.update_settings({'officecli_enabled': True})
        for path in ('/api/office-check', '/api/office-preview'):
            status, payload, _ = request(path, {'job_id': job['id'], 'pages': [1]})
            assert status == 400 and '未检测到 OfficeCLI 或未开启' in payload.decode(), (path, payload)

        with store.tx() as c:  # sha no longer matches the bytes on disk
            c.execute("UPDATE jobs SET result=? WHERE id=?",
                      (dump({**result, 'sha256': '0' * 64}), job['id']))
        status, payload, _ = request(f'/api/export-file?job={job["id"]}&workspace_id={workspace_id}')
        assert status == 400 and 'Excel 文件已变化' in payload.decode()
    finally:
        server.shutdown()
        thread.join()
        server.harness.close()
        server.opencode_harness.close()
        server.server_close()
        server.workspace_lock.close()
