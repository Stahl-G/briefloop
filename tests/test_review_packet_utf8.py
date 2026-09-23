"""Fixed internal JSON is UTF-8 even in a non-UTF-8 API host."""
from copy import deepcopy
import json
import os
from pathlib import Path
import subprocess
import sys

import pytest

from briefloop.store import Store, dump


TEXT = '合成修订记录 🧭 café'


def _review_case(root):
    from briefloop.review import accept_review, respond
    from test_review_learning import save_review

    store = Store(root)
    source = store.add_source(TEXT, 'Revenue was 12 million USD. ' + TEXT)
    run = store.create_run({'title': TEXT, 'objective': 'Explain revenue'}, [source['id']])
    before = store.publish(run['id'], {'title': TEXT, 'markdown': 'Revenue 120 million USD. ' + TEXT})
    first = save_review(store, before, 'review_before')
    first['findings'] = [{'kind': 'insufficient_evidence', 'severity': 'major',
                          'description': TEXT, 'evidence': 'Source states 12, not 120.'}]
    accept_review(store, 'review_before', first)
    assert len(store.rows('SELECT id FROM assessments')) == 1
    assert not store.rows('SELECT id FROM feedback')

    after = store.publish(run['id'], {'title': TEXT, 'markdown': 'Revenue 12 million USD. ' + TEXT},
                          parent_id=before['id'])
    finding = store.rows('SELECT id FROM review_findings')[0]
    response = respond(store, finding['id'], after['id'], 'corrected', TEXT)
    job = store.enqueue('revise', {'version_id': after['id']})
    store.event(job['id'], 'revision_progress', {'message': TEXT})
    final = save_review(store, after, 'review_after')
    final['response_checks'] = [{'response_id': response['id'], 'decision': 'resolved', 'reason': TEXT}]
    accept_review(store, 'review_after', final)
    rows = store.rows("SELECT * FROM feedback WHERE kind='review_correction'")
    assert len(rows) == 1
    feedback = rows[0]
    data = json.loads(feedback['data'])
    assert TEXT in data['before_text'] and TEXT in data['after_text']
    assert TEXT in dump(data['response']) and TEXT in dump(data['execution_records'])
    assessments = store.rows('SELECT * FROM assessments ORDER BY id')
    assert len(assessments) == 2

    # Simulate a crash after Review/assessment commit but before feedback commit.
    # The same accepted output must recover exactly once without another model.
    with store.tx() as connection:
        connection.execute('DELETE FROM feedback WHERE id=?', (feedback['id'],))
    accept_review(store, 'review_after', final)
    accept_review(store, 'review_after', final)
    restored = store.rows("SELECT * FROM feedback WHERE kind='review_correction'")
    assert len(restored) == 1 and restored[0]['id'] == feedback['id']
    assert restored[0]['data'] == feedback['data']
    assert store.rows('SELECT * FROM assessments ORDER BY id') == assessments
    assert len(store.rows('SELECT id FROM reviews')) == 2


def _release_case(root):
    # The existing fixture labels its transport and judgments synthetic; no model.
    from briefloop import backends
    from briefloop.release import release_file, validate_release
    from test_release import reviewed_report, complete_release

    backends.CAPABILITIES['codex'] |= {'restricted_review'}
    store, _, brief, _ = reviewed_report(root)
    release, job = complete_release(store, brief)
    assert store.one('jobs', job['id'])['status'] == 'complete'
    assert validate_release(store, release)['version_id'] == brief['id']
    original = release_file(store, release['id']).read_bytes()
    assert release_file(store, release['id']).read_bytes() == original
    records = store.root / release['result']['manifest_path']
    assert '达到要求' in (records.parent / 'records.json').read_text(encoding='utf-8')


def _manifest_case(root):
    from briefloop import backends
    from briefloop.release import sha, validate_release
    from test_release import reviewed_report, complete_release

    backends.CAPABILITIES['codex'] |= {'restricted_review'}
    store, _, brief, _ = reviewed_report(root)
    release, _ = complete_release(store, brief)
    original_path = store.root / release['result']['manifest_path']
    manifest = json.loads(original_path.read_text(encoding='utf-8'))
    # Current generated identifiers are ASCII. A valid JSON extension must still
    # obey the fixed UTF-8 protocol, independently of the host's locale.
    manifest['producer_note'] = TEXT
    path = original_path.with_name('unicode-manifest.json')
    path.write_bytes(dump(manifest).encode('utf-8'))
    extended = deepcopy(release)
    extended['result'].update(manifest_path=str(path.relative_to(store.root)),
                              manifest_hash=sha(path.read_bytes()))
    assert validate_release(store, extended)['producer_note'] == TEXT

    # Hash-valid bytes are not permission to reinterpret malformed UTF-8 as GBK.
    # A1 A1 is a valid GBK ideographic space, but is invalid standalone UTF-8.
    broken = dump({**manifest, 'producer_note': 'BAD_UTF8'}).encode('utf-8').replace(b'BAD_UTF8', b'\xa1\xa1')
    path.write_bytes(broken)
    extended['result']['manifest_hash'] = sha(broken)
    with pytest.raises(UnicodeDecodeError):
        validate_release(store, extended)


def _company_case(root):
    from briefloop.company_context import prepare_review
    from briefloop.native_orchestrator import prepare
    store = Store(root)
    store.update_settings({'company_context_enabled': True, 'auto_learn': False})
    source = store.add_source(TEXT, TEXT)
    run = store.create_run({'title': TEXT, 'objective': TEXT, 'writing_mode': 'internal_report'}, [source['id']])
    job = store.enqueue('generate', {'run_id': run['id'], 'agent_backend': 'briefloop-native',
                                    'runtime': {'model': 'fake/unused'}})
    folder = store.root/'jobs'/job['id']; folder.mkdir(parents=True)
    class NoModel(Exception): pass
    class Runtime:
        reached = False
        def execute(self, staged, prompt, review_folder, **kwargs):
            payload = json.loads((review_folder/'input.json').read_text(encoding='utf-8'))
            assert payload['requirements']['objective'] == TEXT
            assert payload['sources'][0]['hash'] == source['hash']
            config, _ = prepare(store, staged, review_folder, prompt)
            assert config['task_kind'] == 'company_review'
            assert json.loads((review_folder/'packet/input.json').read_text(encoding='utf-8')) == payload
            self.reached = True
            raise NoModel()
    runtime = Runtime()
    with pytest.raises(NoModel):
        prepare_review(store, runtime, job, run, folder, 'briefloop-native')
    assert runtime.reached
    assert store.one('runs', run['id'])['requirements'] == run['requirements']
    assert store.source_text(source['id']) == TEXT


def _template_prepare_case(root):
    import hashlib
    from briefloop.native_orchestrator import prepare
    assert sys.flags.utf8_mode == 0
    store = Store(root)
    info = json.loads((root/'template-fixture.json').read_text(encoding='utf-8'))
    original = root/'templates'/info['id']/'original.docx'
    inventory = original.with_name('inventory.json')
    before = json.loads(inventory.read_text(encoding='utf-8'))
    job = store.enqueue('prepare_template', {'template_id': info['id']})
    folder = root/'prepare'; folder.mkdir()
    config, _ = prepare(store, job, folder, 'Read the frozen inventory')
    assert config['task_kind'] == 'prepare_template'
    assert json.loads((folder/'packet/inventory.json').read_text(encoding='utf-8')) == before
    assert before['blocks'][0]['text'] == '合成产品名 🧪'
    assert before['headers'] == [[TEXT]] and before['footers'] == [[TEXT]]
    assert hashlib.sha256(original.read_bytes()).hexdigest() == info['source_hash']
    # Valid GBK is not valid UTF-8; never silently reinterpret internal JSON.
    inventory.write_bytes(b'{"note":"\xa1\xa1"}')
    with pytest.raises(UnicodeDecodeError):
        prepare(store, job, folder, 'No model')
    assert hashlib.sha256(original.read_bytes()).hexdigest() == info['source_hash']


def _template_case(root):
    import hashlib
    from io import BytesIO
    from docx import Document
    from briefloop.templates import import_template
    store = Store(root)
    document = Document(); document.add_paragraph('合成产品名 🧪')
    document.sections[0].header.paragraphs[0].text = TEXT
    document.sections[0].footer.paragraphs[0].text = TEXT
    output = BytesIO(); document.save(output); data = output.getvalue()
    row = import_template(store, '合成模板 🧪.docx', data, prepare_job=False)
    assert row['source_hash'] == hashlib.sha256(data).hexdigest()
    assert (root/'templates'/row['id']/'original.docx').read_bytes() == data
    (root/'template-fixture.json').write_text(dump(row), encoding='utf-8')
    code = 'from pathlib import Path; import sys; from test_review_packet_utf8 import _template_prepare_case; _template_prepare_case(Path(sys.argv[1]))'
    prepared = subprocess.run([sys.executable, '-X', 'utf8=0', '-c', code, str(root)],
                              capture_output=True, text=True, encoding='utf-8', timeout=30)
    assert prepared.returncode == 0, prepared.stderr


def _analyst_case(root):
    from briefloop import analyst, analyst_drafts as drafts, writer_input as writer
    from briefloop.native_roles import _source_ids
    store = Store(root)
    source = store.add_source(TEXT, TEXT)
    run = store.create_run({'title': TEXT, 'objective': TEXT, 'allow_web': False}, [source['id']])
    inputs = {'plan': {'draft_structure': [TEXT]}, 'research': {'sources': [{
        'source_id': source['id'], 'locator': 'line 1', 'excerpt': TEXT,
        'facts': [TEXT], 'coverage_status': 'complete'}], 'gaps': []}}
    job = store.enqueue('generate', {'run_id': run['id'], 'writer_input_protocol': 'writer_input_v1',
                                    'agent_backend': 'briefloop-native', 'runtime': {'model': 'fake/unused'}})
    folder = root/'writer'
    class Runtime:
        calls = 0
        def execute(self, staged, prompt, destination):
            self.calls += 1
            self.config = {**staged['native_packet'], 'native_role': 'analyst',
                           'packet_root': str(destination/'packet'), 'attempt_id': staged['id']}
            cfg = self.config
            assert writer.protocol(cfg) == 'writer_input_v1'
            assert _source_ids(store, cfg) == {source['id']}
            writer.ensure_revision_base(store, cfg)
            if self.calls == 1:
                writer.write_sections(store, cfg, {'sections': [
                    {'section_id': 'body', 'markdown': f'{TEXT}[@{source["id"]}]'},
                    {'section_id': 'next', 'markdown': TEXT}]})
                saved = writer.assemble_report(store, cfg, {'title': TEXT, 'section_ids': ['body', 'next']})
                saved = writer.assemble_evidence(store, cfg, {'base_revision': saved['revision'],
                    'citations': [{'source_id': source['id'], 'excerpt': TEXT}]})
                self.revision = saved['revision']
            drafts.check(store, cfg, {'revision': self.revision})
            drafts.submit(store, cfg, {'revision': self.revision})
    runtime = Runtime()
    first = analyst.run(store, runtime, job, run['id'], folder, 'briefloop-native', publish=False, **inputs)
    frozen = {path.relative_to(folder/'packet').as_posix(): path.read_bytes()
              for path in (folder/'packet').rglob('*') if path.is_file()}
    task = json.loads(frozen['input.json'].decode('utf-8'))
    assert task['requirements']['objective'] == TEXT
    assert json.loads(frozen['document-guide.json'].decode('utf-8'))['protocol'] == 'writer_input_v1'
    assert frozen[f'sources/{source["id"]}.txt'].decode('utf-8') == TEXT
    repeated = analyst.run(store, runtime, job, run['id'], folder, 'briefloop-native', publish=False,
                           expected_fingerprint=first['packet_fingerprint'], **inputs)
    assert repeated == first and runtime.calls == 2
    assert all((folder/'packet'/name).read_bytes() == content for name, content in frozen.items())
    original = store.publish(run['id'], drafts.submitted(store, runtime.config))
    revised = analyst.packet(store, run['id'], root/'revision', base_version=original['id'],
                             feedback=[TEXT], writer_protocol='writer_input_v1', **inputs)
    cfg = {'run_id': run['id'], 'packet_root': str(revised['root']),
           'result_file': str(root/'revision/draft.json'), 'attempt_id': 'revision'}
    writer.ensure_revision_base(store, cfg)
    current = drafts._read(drafts._root(store, cfg)/'current.json')
    writer.ensure_revision_base(store, cfg)
    assert drafts._read(drafts._root(store, cfg)/'current.json') == current
    assert TEXT in drafts._candidate(store, cfg, current)['draft']['markdown']


@pytest.mark.parametrize('case', ['review', 'release', 'manifest', 'company', 'template', 'analyst'])
def test_fixed_packets_in_non_utf8_subprocess(tmp_path, case):
    import briefloop

    # Use whichever package pytest imported, including an installed-wheel run.
    paths = [str(Path(briefloop.__file__).resolve().parent.parent), str(Path(__file__).parent)]
    env = {**os.environ, 'PYTHONPATH': os.pathsep.join(paths), 'PYTHONUTF8': '0',
           'PYTHONCOERCECLOCALE': '0', 'LC_ALL': 'C', 'PYTHONIOENCODING': 'utf-8'}
    code = '''import json, locale, sys
from pathlib import Path
import test_review_packet_utf8 as cases
assert sys.flags.utf8_mode == 0
encoding = locale.getpreferredencoding(False)
if encoding.lower().replace('-', '') == 'utf8':
    print(json.dumps({'skip': 'host locale is already UTF-8'}))
else:
    getattr(cases, '_' + sys.argv[1] + '_case')(Path(sys.argv[2]))
    print(json.dumps({'case': sys.argv[1], 'encoding': encoding, 'utf8_mode': 0}))
'''
    result = subprocess.run([sys.executable, '-X', 'utf8=0', '-c', code, case, str(tmp_path)],
                            env=env, capture_output=True, text=True, encoding='utf-8', timeout=60)
    assert result.returncode == 0, result.stderr
    outcome = json.loads(result.stdout)
    if outcome.get('skip'):
        pytest.skip(outcome['skip'])
    assert outcome['case'] == case and outcome['utf8_mode'] == 0
