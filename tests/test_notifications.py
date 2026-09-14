from briefloop.store import Store
from briefloop.notifications import post,snapshot,mark_read,job_status,wiki_changed,version_available


def test_read_survives_reopen_and_cannot_swallow_new_activity(tmp_path):
    store=Store(tmp_path)
    post(store,'one','reports','已开始')
    before=snapshot(store)
    post(store,'two','reports','已完成')
    post(store,'two','reports','重放完成')
    mark_read(store,before['through'],'reports')
    result=snapshot(Store(tmp_path))
    assert result['counts']['reports']==1
    assert len(result['items'])==2
    assert result['items'][0]['title']=='已完成'


def test_task_without_chat_still_notifies_and_attempts_stay_distinct(tmp_path):
    store=Store(tmp_path)
    job=store.enqueue('generate',{'run_id': store.create_run({'title':'测试','objective':'测试任务绑定','allow_web':True},[])['id']})
    assert snapshot(store)['unread']==0  # queued is not running
    job_status(store,job,'running');job_status(store,job,'complete')
    job_status(store,job,'complete')
    assert snapshot(store)['counts']['reports']==2
    job['payload']='{"run_id":"fixture","attempt":2}'
    job['error']='HTTP 503 sk-exampleSecret123456'
    job_status(store,job,'failed')
    latest=snapshot(store)['items'][0]
    assert '503' in latest['body'] and 'exampleSecret' not in latest['body']
    assert snapshot(store)['counts']['reports']==3


def test_wiki_content_and_stable_version_deduplicate(tmp_path):
    store=Store(tmp_path)
    wiki_changed(store,'经验A');wiki_changed(store,'经验A')
    wiki_changed(store,'经验B')
    version_available(store,'0.20.0','0.21.0');version_available(store,'0.20.0','0.21.0')
    version_available(store,'0.20.0','0.18.0')
    value=snapshot(store)
    assert value['counts']['learning']==2
    assert value['counts']['updates']==1
    mark_read(store,value['through'])
    assert snapshot(Store(tmp_path))['unread']==0


def test_prepared_user_template_emits_actual_update(tmp_path):
    from io import BytesIO
    from docx import Document
    from briefloop.templates import import_template,prepare
    store=Store(tmp_path)
    doc=Document();doc.add_paragraph('本月报告');doc.add_paragraph('一、进展');doc.add_paragraph('往期正文')
    stream=BytesIO();doc.save(stream)
    record=import_template(store,'用户月报.docx',stream.getvalue(),prepare_job=False)
    assert snapshot(store)['counts']['templates']==0
    prepared=prepare(store,record['id'],{'keep_blocks':[0],'paragraph_index':2,
        'sections':[{'section_id':'progress','title':'一、进展','index':1}],
        'fields':[{'old':'本月报告','field':'title'}]})
    result=snapshot(store)
    assert prepared['status']=='ready' and result['counts']['templates']==1
    assert result['items'][0]['target']['template_id']==record['id']
