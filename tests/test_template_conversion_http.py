"""Template-first output runs through the real HTTP and file-job path, without a model."""
import base64
import http.client
from io import BytesIO
import json
import threading
import time

from docx import Document

from briefloop.server import make_server, _close_service
from briefloop.templates import import_builtin


def test_template_conversion_http_produces_downloadable_word_without_rewriting(tmp_path):
    server = make_server(tmp_path / 'workspace', port=0, paused=True)
    import_builtin(server.store)
    server.worker.start()
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    authority = f'127.0.0.1:{server.server_port}'
    token = ''

    def request(path, body=None):
        connection = http.client.HTTPConnection('127.0.0.1', server.server_port, timeout=10)
        try:
            headers = {'X-BriefLoop-Token': token, 'Origin': 'http://' + authority}
            if body is not None:
                headers['Content-Type'] = 'application/json'
            connection.request('POST' if body is not None else 'GET', '/api/' + path,
                               json.dumps(body) if body is not None else None, headers)
            response = connection.getresponse()
            data = response.read()
            return response.status, dict(response.getheaders()), data
        finally:
            connection.close()

    try:
        status, _, data = request('session')
        assert status == 200
        token = json.loads(data)['token']
        workspace = server.store.meta('workspace_id')
        template = server.store.rows("SELECT id FROM templates WHERE name='商业报告·极简蓝'")[0]['id']
        text = '# 原稿标题\n\n这一段必须保留：收入 12，成本 8。\n\n| 指标 | 数值 |\n| --- | --- |\n| 收入 | 12 |\n'
        body = {'name': '合成原稿.md', 'data': base64.b64encode(text.encode()).decode(),
                'template_id': template, 'workspace_id': workspace, 'request_id': 'conversion-1'}
        status, _, data = request('template-convert', {**body, 'workspace_id': 'different-workspace'})
        assert status == 400 and '工作区已切换' in data.decode()
        assert server.store.rows('SELECT id FROM briefs') == []
        status, _, data = request('template-convert', body)
        assert status == 200, data.decode()
        accepted = json.loads(data)
        status, _, data = request('template-convert', body)
        assert status == 200, data.decode()
        converted = json.loads(data)
        assert converted['version']['id'] == accepted['version']['id']
        assert converted['job']['id'] == accepted['job']['id']
        assert len(server.store.rows('SELECT id FROM sources')) == 1
        assert len(server.store.rows('SELECT id FROM briefs')) == 1
        version, job = converted['version'], converted['job']
        assert version['author'] == 'user'
        deadline = time.monotonic() + 15
        while job['status'] in ('queued', 'running') and time.monotonic() < deadline:
            time.sleep(.05)
            status, _, data = request('export-status?job=' + job['id'])
            assert status == 200
            job = json.loads(data)
        assert job['status'] == 'complete', job
        status, headers, data = request('export-file?job=' + job['id'] + '&workspace_id=' + workspace)
        assert status == 200
        assert headers['Content-Disposition'].startswith('attachment;')
        document = Document(BytesIO(data))
        assert any('这一段必须保留：收入 12，成本 8。' in p.text for p in document.paragraphs)
        assert any([cell.text for cell in row.cells] == ['收入', '12'] for table in document.tables for row in table.rows)
        assert server.store.one('briefs', version['id']) == version
        assert {row['kind'] for row in server.store.rows('SELECT kind FROM jobs')} == {'export_docx'}
        assert server.store.rows('SELECT id FROM feedback') == []
        assert server.store.rows('SELECT id FROM assessments') == []
        status, _, data = request('reports')
        assert status == 200 and version['id'] in data.decode()
        status, _, data = request('export', {'version_id': version['id'], 'template_id': template, 'workspace_id': 'wrong'})
        assert status == 400 and '工作区已切换' in data.decode()
    finally:
        server.shutdown()
        thread.join(timeout=5)
        _close_service(server)
