"""Focused export provenance, visible labels, and native package metadata checks."""
from io import BytesIO
import json
import threading
from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

from briefloop.export_labeling import (brief_label, office_properties, read_office_properties,
                                      pdf_properties, markdown_label, FIELDS, NOTICE_ZH)
from briefloop.document_model import markdown_document
from briefloop.store import Store, dump
from briefloop.exports import docx_bytes
from briefloop.export_jobs import export_input, enqueue_export, generate_word, output_path
from briefloop.xlsx_export import xlsx_bytes, _enhanced_mismatches


def sample(tmp_path, *, author='agent'):
    store = Store(tmp_path / 'workspace')
    source = store.add_source('合成验收数据', '收入同比增长5%，成本同比下降3%。')
    run = store.create_run({'title': '标识验收', 'objective': '核对导出', 'allow_web': False}, [source['id']])
    document = markdown_document('## 经营指标\n\n收入同比增长 5%，成本同比下降 3%。\n\n| 指标 | 同比 |\n| --- | --- |\n| 收入 | 5% |\n| 成本 | -3% |')
    brief = store.publish(run['id'], {'title': '标识验收', 'editor_document': document}, author=author)
    return store, run, brief, document


def test_positive_provenance_survives_user_revision_without_labeling_user_only(tmp_path):
    store, run, original, document = sample(tmp_path, author='user')
    assert brief_label(store, original) is None
    human = store.revise(original['id'], editor_document=markdown_document(original['markdown']+'\n\n人工补充'))
    assert brief_label(store, human) is None
    agent = store.revise(human['id'], editor_document=markdown_document(human['markdown']+'\n\n模型补充'), author='agent')
    mixed = store.revise(agent['id'], editor_document=markdown_document(agent['markdown']+'\n\n用户修订'))
    label = brief_label(store, mixed)
    assert label['metadata']['Label'] == '1'
    assert label['metadata']['ProduceID'] == mixed['id']
    assert '人工核验' not in label['notice'] and 'reviewed' not in label['notice']
    assert brief_label(store, mixed, language='English')['notice'].startswith('This file contains AI-generated')


@pytest.mark.parametrize('template', [False, True])
def test_word_has_body_and_all_footer_labels_registered_once_without_mutating_report(tmp_path, template):
    store, run, brief, document = sample(tmp_path)
    original = store.one('briefs', brief['id'])
    override = None
    if template:
        from briefloop.templates import import_builtin
        import_builtin(store)
        override = next(row['id'] for row in store.rows('SELECT * FROM templates') if row['name'] == '通用报告·品牌黛蓝')
    job = enqueue_export(store, brief['id'], override)
    generate_word(store, job, threading.Event())
    blob = output_path(store, job).read_bytes()
    doc = Document(BytesIO(blob))
    assert sum(p.text == NOTICE_ZH for p in doc.paragraphs) == 1
    for section in doc.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            assert sum(p.text == NOTICE_ZH for p in footer.paragraphs) == 1
    delta_table = next(table for table in doc.tables if table.rows[0].cells[-1].text == '同比')
    assert str(delta_table.rows[1].cells[-1].paragraphs[0].runs[-1].font.color.rgb) == 'D9363E'
    assert str(delta_table.rows[2].cells[-1].paragraphs[0].runs[-1].font.color.rgb) == '1E8E4F'
    values = read_office_properties(blob)
    metadata = json.loads(values['AIGC'])
    assert set(metadata) == set(FIELDS) and metadata['Label'] == '1'
    assert metadata['ContentProducer'] == metadata['ContentPropagator'] == 'BriefLoop'
    assert values['BriefLoopMarketConvention'] == 'cn'
    with ZipFile(BytesIO(blob)) as archive:
        custom = etree.fromstring(archive.read('docProps/custom.xml'))
        assert len(custom.xpath('//*[local-name()="property"][@name="AIGC"]')) == 1
        assert b'docProps/custom.xml' in archive.read('_rels/.rels')
        assert b'/docProps/custom.xml' in archive.read('[Content_Types].xml')
    assert store.one('briefs', brief['id']) == original


def test_custom_properties_keep_existing_fields_and_are_idempotent(tmp_path):
    store, _, brief, _ = sample(tmp_path)
    blob = docx_bytes('Body', label=brief_label(store, brief))
    blob = office_properties(blob, {'BusinessOwner': 'Finance'})
    blob = office_properties(blob, {'AIGC': read_office_properties(blob)['AIGC']})
    values = read_office_properties(blob)
    assert values['BusinessOwner'] == 'Finance'
    with ZipFile(BytesIO(blob)) as archive:
        root = etree.fromstring(archive.read('docProps/custom.xml'))
        assert len({p.get('pid') for p in root}) == len(root)
        assert len(root.xpath('//*[local-name()="property"][@name="AIGC"]')) == 1
    Document(BytesIO(blob))


@pytest.mark.parametrize('layout', ['sheets', 'single'])
def test_excel_labels_each_sheet_and_print_footer_without_shifting_data(tmp_path, layout):
    store, run, brief, document = sample(tmp_path)
    blob = xlsx_bytes(document, json.loads(brief['detail']), json.loads(run['requirements']), layout, label=brief_label(store, brief))
    workbook = load_workbook(BytesIO(blob))
    for sheet in workbook:
        assert sum(cell.value == NOTICE_ZH for row in sheet for cell in row) == 1
        assert sheet.oddFooter.center.text == sheet.evenFooter.center.text == sheet.firstFooter.center.text == NOTICE_ZH
    table = workbook[workbook.sheetnames[-1]]
    assert table['B3'].value == '5%' and table['B4'].value == '-3%'
    assert table['B3'].font.color.rgb == '00D9363E'
    assert table['B4'].font.color.rgb == '001E8E4F'
    assert table['B3'].alignment.horizontal == table['B4'].alignment.horizontal == 'right'
    assert json.loads(read_office_properties(blob)['AIGC'])['Label'] == '1'


def test_human_only_office_exports_have_no_ai_claim(tmp_path):
    store, run, brief, document = sample(tmp_path, author='user')
    label = brief_label(store, brief)
    blob = docx_bytes(document=document, label=label)
    assert 'AIGC' not in read_office_properties(blob)
    assert NOTICE_ZH not in '\n'.join(p.text for p in Document(BytesIO(blob)).paragraphs)
    blob = xlsx_bytes(document, {}, {}, 'sheets', label=label)
    assert 'AIGC' not in read_office_properties(blob)
    assert all(NOTICE_ZH != cell.value for sheet in load_workbook(BytesIO(blob)) for row in sheet for cell in row)


def test_excel_enhancement_that_drops_print_labels_is_rejected(tmp_path):
    store, run, brief, document = sample(tmp_path)
    blob = xlsx_bytes(document, {}, {}, 'sheets', label=brief_label(store, brief))
    workbook = load_workbook(BytesIO(blob)); workbook.worksheets[-1].oddFooter.center.text = ''
    staging = tmp_path / 'changed.xlsx'; workbook.save(staging)
    with pytest.raises(ValueError, match='打印标识'):
        _enhanced_mismatches(staging, {'commands': [], 'formulas': []}, blob)


def test_pdf_metadata_and_markdown_header_use_format_specific_schema(tmp_path):
    store, _, brief, _ = sample(tmp_path)
    label = brief_label(store, brief)
    buffer = BytesIO(); writer = PdfWriter(); writer.add_blank_page(width=100,height=100); writer.write(buffer)
    pdf = PdfReader(BytesIO(pdf_properties(buffer.getvalue(), label, {'market_convention': 'intl'})))
    assert len(pdf.pages) == 1
    assert json.loads(pdf.metadata['/AIGC']) == label['metadata']
    assert pdf.metadata['/BriefLoopMarketConvention'] == 'intl'
    markdown = markdown_label('# Body', label)
    assert markdown.startswith('---\nAIGC:\n  Label: "1"\n') and markdown.endswith(NOTICE_ZH+'\n')
    assert markdown_label('# Human', None) == '# Human'


def test_queued_word_convention_change_rejects_stale_fingerprint(tmp_path):
    store, run, brief, _ = sample(tmp_path)
    queued = enqueue_export(store, brief['id'])
    req = json.loads(run['requirements']); req['market_convention'] = 'intl'
    with store.tx() as connection:
        connection.execute('UPDATE runs SET requirements=? WHERE id=?', (dump(req), run['id']))
    with pytest.raises(ValueError, match='导出输入已变化'):
        generate_word(store, queued, threading.Event())
    assert export_input(store, brief)[0]['requirements']['market_convention'] == 'intl'


def test_excel_mixed_direction_prose_is_not_colored_as_a_single_delta(tmp_path):
    store, run, brief, _ = sample(tmp_path)
    document = markdown_document('| 指标 | 说明 |\n| --- | --- |\n| 经营 | 营收同比 +5%，利润同比 -2% |')
    workbook = load_workbook(BytesIO(xlsx_bytes(document, {}, {}, 'single')))
    cell = workbook.active['B3']
    assert cell.value == '营收同比 +5%，利润同比 -2%'
    assert cell.font.color is None or cell.font.color.type != 'rgb'


def test_local_pdf_finalization_binds_workspace_version_and_market(tmp_path):
    import http.client
    from briefloop.server import make_server
    server = make_server(tmp_path / 'http-workspace', port=0, paused=True)
    thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
    try:
        store = server.store
        source = store.add_source('Fixture', '5% growth')
        run = store.create_run({'title': 'PDF', 'objective': 'test', 'allow_web': False}, [source['id']])
        brief = store.publish(run['id'], {'title': 'PDF', 'markdown': 'AI-generated fixture'})
        workspace = store.meta('workspace_id')
        def request(method, path, data=None, headers=None):
            connection = http.client.HTTPConnection('127.0.0.1', server.server_port)
            connection.request(method, path, body=data, headers=headers or {})
            response = connection.getresponse(); result = response.status, response.read()
            connection.close(); return result
        _, response = request('GET', '/api/session'); token = json.loads(response)['token']
        _, response = request('GET', '/api/export-label?version='+brief['id']+'&workspace_id='+workspace)
        info = json.loads(response)
        assert info['label']['metadata']['Label'] == '1' and info['market_convention'] == 'cn'
        writer = PdfWriter(); writer.add_blank_page(width=100, height=100); buf = BytesIO(); writer.write(buf)
        route = '/api/export-pdf-label?version='+brief['id']+'&workspace_id='+workspace+'&market_convention=cn'
        assert request('POST', route, buf.getvalue())[0] == 403
        headers = {'Content-Type': 'application/pdf', 'X-BriefLoop-Token': token}
        status, result = request('POST', route, buf.getvalue(), headers)
        assert status == 200
        assert json.loads(PdfReader(BytesIO(result)).metadata['/AIGC'])['ProduceID'] == brief['id']
        status, result = request('POST', route.replace('market_convention=cn', 'market_convention=intl'), buf.getvalue(), headers)
        assert status == 400 and '配色设置已变化' in result.decode()
        status, result = request('POST', route.replace(workspace, 'wrong'), buf.getvalue(), headers)
        assert status == 400 and '工作区已切换' in result.decode()
    finally:
        server.shutdown(); server.server_close(); thread.join()


def test_legacy_word_code_fences_are_not_semantic_changes():
    doc = Document(BytesIO(docx_bytes('growth 5%\n\n```text\ngrowth 5%\n```\n\n`growth 5%`')))
    assert any(str(run.font.color.rgb) == 'D9363E' for run in doc.paragraphs[0].runs)
    assert all(run.font.color.rgb is None for run in doc.paragraphs[1].runs)
    assert all(run.font.color.rgb is None for run in doc.paragraphs[2].runs)
