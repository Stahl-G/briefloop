"""Native XLSX file IO, locks and junctions; no model judgments are simulated."""
from io import BytesIO
import hashlib
import json
import os
import threading

from openpyxl import load_workbook
import pytest

from briefloop import office_cli, xlsx_export
from briefloop.platform_support import filesystem_path
from briefloop.store import Store
from test_export_xlsx import _sales_document


pytestmark = pytest.mark.skipif(os.name != 'nt', reason='Windows native file paths, locks and junctions')


@pytest.fixture(params=[False, True], ids=['short', 'long-unicode'])
def case(tmp_path, request):
    name='中文表格工作区-'
    if request.param:name+='w'*(224-len(str(tmp_path))-1-len(name))
    store=Store(tmp_path/name)
    store.update_settings({'officecli_enabled':False, 'auto_learn':False})
    source=store.add_source('Synthetic fixture', 'Manually entered synthetic table; no model output.')
    run=store.create_run({'title':'合成表格', 'objective':'Export the manual table', 'allow_web':False}, [source['id']])
    brief=store.publish(run['id'], {'title':'合成表格','markdown':'Manual synthetic table',
        'editor_document':{'type':'doc','content':_sales_document()}}, author='user')
    return store, brief


def generate(store, brief, layout='sheets'):
    job=xlsx_export.enqueue_export_xlsx(store, brief['id'], layout)
    result=xlsx_export.generate_xlsx(store, job, threading.Event())
    store.update_job(job['id'], 'complete', result=result)
    return store.one('jobs', job['id']), result


def test_long_workbooks_cache_hashes_layouts_and_stage_reads_preserve_cells(case):
    store, brief=case
    for layout in ('sheets','single'):
        job,result=generate(store, brief, layout)
        path=xlsx_export.output_path_xlsx(store, job);blob=filesystem_path(path).read_bytes()
        assert result['sha256']==hashlib.sha256(blob).hexdigest()
        assert '\\' not in result['path'] and not str(path).startswith('\\\\?\\')
        assert xlsx_export.enqueue_export_xlsx(store, brief['id'], layout)['id']==job['id']
        with BytesIO(blob) as stream:
            workbook=load_workbook(stream)
            sheet=workbook[workbook.sheetnames[-1]]
            assert sheet['B3'].value==1234 and sheet['C3'].value=='007'
            assert sheet['B5'].value==1254 and sheet['C5'].value==7.5
            workbook.close()
        staging=path.with_name('report.staging.xlsx');filesystem_path(staging).write_bytes(blob)
        assert xlsx_export._enhanced_mismatches(staging, {'commands':[],'formulas':[]}, blob)==[]
        changed=load_workbook(filesystem_path(staging));changed[changed.sheetnames[-1]]['A3']='unauthorized rewrite'
        changed.save(filesystem_path(staging));changed.close()
        with pytest.raises(ValueError, match='未授权内容'):
            xlsx_export._enhanced_mismatches(staging, {'commands':[],'formulas':[]}, blob)
        filesystem_path(staging).unlink()
        if len(str(store.root))==224:assert len(str(path))>260
    filesystem_path(path).write_bytes(b'changed artifact')
    assert xlsx_export.enqueue_export_xlsx(store, brief['id'], layout)['id']!=job['id']
    assert store.one('briefs', brief['id'])['hash']==brief['hash']


def test_native_read_handle_keeps_original_and_saves_an_alternate_workbook(case):
    store, brief=case
    job,first=generate(store, brief)
    path=xlsx_export.output_path_xlsx(store, job);original=filesystem_path(path).read_bytes()
    with filesystem_path(path).open('rb'):
        result=xlsx_export.generate_xlsx(store, job, threading.Event())
    assert result['path']!=first['path'] and result['path'].endswith('.xlsx')
    assert filesystem_path(path).read_bytes()==original
    assert hashlib.sha256(filesystem_path(store.root/result['path']).read_bytes()).hexdigest()==result['sha256']
    assert not filesystem_path(path.with_suffix('.tmp')).exists()
    saved=json.loads(store.rows("SELECT data FROM events WHERE job_id=? AND kind='export_saved_as' ORDER BY seq DESC",(job['id'],))[0]['data'])
    assert saved['path']==result['path']


@pytest.mark.parametrize('location', ['directory','temporary','staging'])
def test_junction_cannot_redirect_workbook_or_enhancement_writes(case, tmp_path, monkeypatch, location):
    import _winapi
    store, brief=case
    if location=='staging':monkeypatch.setattr(office_cli,'enhancement_ready',lambda store:True)
    job=xlsx_export.enqueue_export_xlsx(store, brief['id'])
    folder=store.root/'exports'/job['id']
    link=folder if location=='directory' else folder/('report.tmp' if location=='temporary' else 'report.staging.xlsx')
    filesystem_path(link.parent).mkdir(parents=True,exist_ok=True)
    target=tmp_path/'owned unrelated workbook';target.mkdir()
    marker=target/'keep.txt';marker.write_bytes(b'unchanged target')
    _winapi.CreateJunction(str(target),str(filesystem_path(link)))
    try:
        with pytest.raises(ValueError,match='链接'):
            xlsx_export.generate_xlsx(store, job, threading.Event())
        assert {p.name for p in target.iterdir()}=={'keep.txt'} and marker.read_bytes()==b'unchanged target'
    finally:
        os.rmdir(filesystem_path(link))  # Remove the owned junction, never its target.
