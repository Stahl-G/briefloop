"""Saved citation locations survive the actual Word export entry points."""
from io import BytesIO
import http.client
import threading

from docx import Document
from docx.oxml.ns import qn
import pytest

from briefloop.document_model import markdown_document
from briefloop.export_jobs import enqueue_export, generate_word, output_path
from briefloop.exports import reader_markdown
from briefloop.store import Store


def _download(store, version_id):
    from briefloop.server import make_server
    server = make_server(store.root, port=0, paused=True)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    connection = http.client.HTTPConnection('127.0.0.1', server.server_port)
    try:
        connection.request('GET', '/api/download?format=docx&version=' + version_id)
        response = connection.getresponse()
        blob = response.read()
        assert response.status == 200
        return blob
    finally:
        connection.close()
        server.shutdown(); thread.join()
        server.harness.close(); server.opencode_harness.close(); server.native_harness.close()
        server.runtime_bridge.close(); server.native_engine.close()
        server.server_close(); server.workspace_lock.close()


@pytest.mark.parametrize('route', ['queue', 'template', 'download', 'legacy_download'])
def test_word_source_locations_match_saved_version_and_keep_links(tmp_path, route):
    store = Store(tmp_path / 'workspace')
    web = store.add_source('Public evidence', 'First line.\nSecond line.\nThird line.', url='https://example.test/report')
    local = store.add_source('Local evidence', 'First line.\nSecond line.\nThird line.')
    unused = store.add_source('Uncited evidence', 'Not cited in the document.')
    run = store.create_run({'title': 'Source locations', 'objective': 'Preserve saved locations'},
                           [web['id'], local['id'], unused['id']])
    markdown = 'Web finding [@' + web['id'] + ']. Local finding [@' + local['id'] + '].'
    citations = [{'source_id': web['id'], 'locator': value} for value in ['line 1', 'lines 2-3', 'line 1', '']]
    citations += [{'source_id': local['id'], 'locator': 'line 3'},
                  {'source_id': unused['id'], 'locator': 'unused-location'}]
    content = {'markdown': markdown} if route == 'legacy_download' else {'editor_document': markdown_document(markdown)}
    brief = store.publish(run['id'], {'title': 'Source locations', **content, 'citations': citations})
    saved_brief = store.one('briefs', brief['id'])
    saved_sources = store.rows('SELECT * FROM sources ORDER BY id')
    saved_markdown = reader_markdown(store, brief)
    assert 'Public evidence](https://example.test/report) · line 1；lines 2-3' in saved_markdown

    if route.endswith('download'):
        blob = _download(store, brief['id'])
    else:
        template_id = None
        if route == 'template':
            from briefloop.templates import import_template, prepare
            original = Document()
            original.add_paragraph('Template title'); original.add_heading('Summary', level=1)
            original.add_paragraph('Sample body.')
            stream = BytesIO(); original.save(stream)
            template = import_template(store, 'Synthetic.docx', stream.getvalue(), prepare_job=False)
            prepare(store, template['id'], {'keep_blocks': [0], 'paragraph_index': 2,
                    'sections': [{'section_id': 'summary', 'title': 'Summary', 'index': 1}],
                    'fields': [{'old': 'Template title', 'field': 'title'}]})
            template_id = template['id']
        job = enqueue_export(store, brief['id'], template_override=template_id)
        result = generate_word(store, job, threading.Event())
        assert result['version_id'] == brief['id']
        blob = output_path(store, job).read_bytes()

    document = Document(BytesIO(blob))
    paragraphs = [''.join(p.xpath('.//w:t/text()')) for p in document.element.xpath('.//w:p')]
    assert sum(p.endswith('Public evidence · line 1；lines 2-3') for p in paragraphs) == 1
    assert sum(p.endswith('Local evidence · line 3') for p in paragraphs) == 1
    assert 'unused-location' not in document.element.xml
    assert 'Uncited evidence' not in document.element.xml
    web_links = [link for link in document.element.xpath('.//w:hyperlink')
                 if link.get(qn('r:id')) and document.part.rels[link.get(qn('r:id'))].target_ref == web['url']]
    assert any(''.join(link.xpath('.//w:t/text()')) == 'Public evidence' for link in web_links)
    if route != 'legacy_download':
        bookmarks = {node.get(qn('w:name')) for node in document.element.xpath('.//w:bookmarkStart')}
        local_links = document.element.xpath('.//w:hyperlink[@w:anchor]')
        assert local_links and all(link.get(qn('w:anchor')) in bookmarks for link in local_links)
    assert store.one('briefs', brief['id']) == saved_brief
    assert store.rows('SELECT * FROM sources ORDER BY id') == saved_sources
    assert reader_markdown(store, brief) == saved_markdown
