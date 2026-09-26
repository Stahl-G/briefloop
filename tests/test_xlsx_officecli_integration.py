"""Real-OfficeCLI integration for the Excel export enhancement layer.

These run only where the binary exists (locally verified against 1.0.152);
any other host skips the whole module. The base artifact never needs it."""
import contextlib
import hashlib
import json
import threading
import warnings

import pytest

from briefloop import office_cli, xlsx_export
from briefloop.store import Store, dump

pytestmark = pytest.mark.skipif(office_cli.find() is None,
                                reason='officecli is not installed on this host')


def _para(text):
    return {'type': 'paragraph', 'content': [{'type': 'text', 'text': text}]}


def _row(cells, kind='tableCell'):
    return {'type': 'tableRow', 'content': [
        {'type': kind, 'content': [_para(str(cell))]} for cell in cells]}


def _sales_document(total_amount='$1,500'):
    # Column sums match the totals exactly, so the enhanced workbook's cached
    # formula values equal the numbers the report shows.
    return {'type': 'doc', 'content': [
        {'type': 'heading', 'attrs': {'level': 2}, 'content': [{'type': 'text', 'text': '销售明细'}]},
        {'type': 'table', 'content': [
            _row(['品名', '销量', '金额', '占比'], 'tableHeader'),
            _row(['甲', '10', '1,234', '45%']),
            _row(['乙', '20', '266', '55%']),
            _row(['合计', '30', total_amount, '100%'])]}]}


def _brief(store, document, title='销售月报'):
    source = store.add_source('Synthetic', 'Sales figures.')
    run = store.create_run({'title': title, 'objective': 'Explain sales', 'allow_web': False,
                            'report_date': '2026-09-01'}, [source['id']])
    return store.publish(run['id'], {'title': title, 'markdown': 'Sales.',
                                     'editor_document': document})


def _run_job(store, brief, layout='sheets'):
    job = store.one('jobs', xlsx_export.enqueue_export_xlsx(store, brief['id'], layout)['id'])
    result = xlsx_export.generate_xlsx(store, job, threading.Event())
    store.update_job(job['id'], 'complete', result=result)
    return store.one('jobs', job['id']), result


@contextlib.contextmanager
def _sheets(path):
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')  # x14 data-bar extension warns on read
        formulas = load_workbook(path)
        values = load_workbook(path, data_only=True)
        yield formulas, values
        formulas.close()
        values.close()


def test_enhanced_workbook_formulas_formats_bars_and_checks(tmp_path):
    store = Store(tmp_path / 'workspace')
    store.update_settings({'officecli_enabled': True})
    assert office_cli.enhancement_ready(store) is True
    job, result = _run_job(store, _brief(store, _sales_document()))
    assert result['enhanced'] is True and 'enhance_reason' not in result
    path = tmp_path / 'workspace' / result['path']
    assert result['sha256'] == hashlib.sha256(path.read_bytes()).hexdigest()
    with _sheets(path) as (formulas, values):
        sheet = formulas['销售明细']
        cached = values['销售明细']
        assert sheet['B5'].value == '=SUM(B3:B4)' and cached['B5'].value == 30
        assert sheet['C5'].value == '=SUM(C3:C4)' and cached['C5'].value == 1500
        assert sheet['D5'].value == '=SUM(D3:D4)' and cached['D5'].value == 1
        assert sheet['C3'].number_format == '#,##0'      # thousands display restored
        assert sheet['C5'].number_format == '"$"#,##0'   # currency total: formula + format
        assert sheet['D4'].number_format == '0%' and sheet['D4'].value == 0.55
        assert {'B3:B5', 'C3:C5', 'D3:D5'} <= {str(rule.sqref) for rule in sheet.conditional_formatting}
    checks = {(row['kind'], row['status']) for row in
              store.rows('SELECT kind,status FROM office_checks WHERE file_sha256=?', (result['sha256'],))}
    assert checks == {('validate', 'ok'), ('issues', 'ok')}
    assert any(row['kind'] == 'office_check' for row in
               store.rows("SELECT kind FROM events WHERE job_id=?", (job['id'],)))
    # No staging leftovers next to the delivered artifact.
    assert [p.name for p in path.parent.iterdir()] == ['report.xlsx']


def test_report_total_disagreeing_with_the_column_degrades_to_base(tmp_path):
    store = Store(tmp_path / 'workspace')
    store.update_settings({'officecli_enabled': True})
    job, result = _run_job(store, _brief(store, _sales_document(total_amount='999')))
    assert result['enhanced'] is False
    assert result['enhance_reason'] == '公式结果与报告数值不一致'
    path = tmp_path / 'workspace' / result['path']
    with _sheets(path) as (formulas, _values):
        sheet = formulas['销售明细']
        assert sheet['B5'].value == 30 and sheet['C5'].value == 999  # base values, no formulas
        assert [str(rule.sqref) for rule in sheet.conditional_formatting] == []
    mismatched = store.rows("SELECT data FROM events WHERE kind='export_xlsx_formula_mismatch' AND job_id=?",
                            (job['id'],))
    cells = json.loads(mismatched[0]['data'])['cells']
    assert cells == [{'sheet': '销售明细', 'cell': 'C5', 'expected': 999, 'found': 1500}]
    assert job['status'] == 'complete' and [p.name for p in path.parent.iterdir()] == ['report.xlsx']


def test_exhausted_request_budget_delivers_the_base_workbook(tmp_path, monkeypatch):
    store = Store(tmp_path / 'workspace')
    store.update_settings({'officecli_enabled': True})
    monkeypatch.setattr(office_cli, 'REQUEST_BUDGET_SECONDS', 0.01)  # no subprocess may start
    job, result = _run_job(store, _brief(store, _sales_document(), title='预算耗尽'))
    assert result['enhanced'] is False and result['enhance_reason'] == '请求时间预算用尽'
    path = tmp_path / 'workspace' / result['path']
    with _sheets(path) as (formulas, _):
        assert formulas['销售明细']['B5'].value == 30  # reported value kept, no formula attempted
    failed = store.rows("SELECT data FROM events WHERE kind='export_xlsx_enhancement_failed' AND job_id=?",
                        (job['id'],))
    assert json.loads(failed[0]['data']) == {'reason': '请求时间预算用尽'}
    assert job['status'] == 'complete'
