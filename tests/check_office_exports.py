"""Release smoke against the installed wheel; creates only synthetic Office files."""
from io import BytesIO
import json
from pathlib import Path
import sys
import threading

from docx import Document
from openpyxl import load_workbook
import briefloop
from briefloop import __version__
from briefloop.document_model import markdown_document
from briefloop.export_jobs import enqueue_export, generate_word, output_path
from briefloop.export_labeling import brief_label, read_office_properties, NOTICE_ZH
from briefloop.native_engine import NativeEngine
from briefloop.store import Store
from briefloop.xlsx_export import xlsx_bytes


def main():
    assert Path(briefloop.__file__).resolve().is_relative_to(Path(sys.prefix).resolve())
    out = Path(sys.argv[1]).resolve()
    out.mkdir(parents=True, exist_ok=True)
    results = []
    for market in ('cn', 'intl'):
        store = Store(out / ('synthetic-' + market))
        source = store.add_source('合成验收数据', '收入同比增长5%，成本同比下降3%。')
        run = store.create_run({'title': '合成验收', 'objective': '验证实际导出', 'allow_web': False,
                               'market_convention': market}, [source['id']])
        document = markdown_document('## 经营指标\n\n收入同比增长 5%，成本同比下降 3%。\n\n'
                                     '| 指标 | 同比 |\n| --- | --- |\n| 收入 | 5% |\n| 成本 | -3% |')
        for author in ('agent', 'user'):
            brief = store.publish(run['id'], {'title': '合成验收 ' + author, 'editor_document': document}, author=author)
            label = brief_label(store, brief)
            job = enqueue_export(store, brief['id'])
            generate_word(store, job, threading.Event())
            blob = output_path(store, job).read_bytes()
            doc = Document(BytesIO(blob))
            assert ('AIGC' in read_office_properties(blob)) == (author == 'agent')
            assert any(p.text == NOTICE_ZH for p in doc.paragraphs) == (author == 'agent')
            if author == 'agent':
                assert all(any(p.text == NOTICE_ZH for p in footer.paragraphs) for section in doc.sections
                           for footer in (section.footer, section.first_page_footer, section.even_page_footer))
            table = next(t for t in doc.tables if t.rows[0].cells[-1].text == '同比')
            colors = [str(table.rows[n].cells[-1].paragraphs[0].runs[-1].font.color.rgb) for n in (1, 2)]
            assert colors == (['D9363E', '1E8E4F'] if market == 'cn' else ['1E8E4F', 'D9363E'])
            (out / f'{market}-{author}.docx').write_bytes(blob)
            xblob = xlsx_bytes(document, json.loads(brief['detail']), json.loads(run['requirements']), 'sheets', label=label)
            wb = load_workbook(BytesIO(xblob))
            assert ('AIGC' in read_office_properties(xblob)) == (author == 'agent')
            sheet = wb[wb.sheetnames[-1]]
            assert [sheet['B3'].font.color.rgb, sheet['B4'].font.color.rgb] == ['00' + x for x in colors]
            if author == 'agent':
                assert all(s.oddFooter.center.text == NOTICE_ZH for s in wb)
            (out / f'{market}-{author}.xlsx').write_bytes(xblob)
            results.append({'market': market, 'author_fixture': author, 'docx_xlsx': 'passed', 'colors': colors})
    engine = NativeEngine()
    try:
        ping = engine.call('ping')
        assert ping['pi'] == '1.0.0'
    finally:
        engine.close()
    print(json.dumps({'version': __version__, 'pi': ping['pi'], 'checks': results, 'model_calls': 0}))


if __name__ == '__main__':
    main()
