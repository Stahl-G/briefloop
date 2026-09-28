from io import BytesIO
import json
import threading

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

from briefloop.export_jobs import generate_word, output_path
from briefloop.media import source_files
from briefloop.store import Store
from briefloop.template_conversion import ConversionError, convert_file
from briefloop.templates import import_template, prepare


def word_bytes(doc):
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


@pytest.fixture
def workspace(tmp_path):
    store = Store(tmp_path)
    doc = Document()
    doc.add_paragraph('Template title', 'Title')
    doc.add_heading('Summary', 1)
    doc.add_paragraph('Template body')
    row = import_template(store, 'Template.docx', word_bytes(doc), prepare_job=False)
    prepare(store, row['id'], {'keep_blocks': [0], 'paragraph_index': 2,
                              'sections': [{'section_id': 'summary', 'title': 'Summary', 'index': 1}],
                              'fields': [{'old': 'Template title', 'field': 'title'}]})
    return store, row['id']


def nodes(value):
    yield value
    for child in value.get('content', []):
        yield from nodes(child)


def assert_no_research(store):
    assert {j['kind'] for j in store.rows('SELECT * FROM jobs')} <= {'export_docx'}
    assert store.rows('SELECT * FROM feedback') == []
    assert store.rows('SELECT * FROM assessments') == []


def test_omitted_cells_do_not_shift_values_under_other_headers(workspace):
    store, template_id = workspace
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ['Metric', 'Current', '']):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ['', '120', '130']):
        cell.text = value
    # Both rows have two physical cells but different grid offsets. Flattening
    # them into two columns would put 120 under Metric instead of Current.
    for row, edge, name in ((table.rows[0], -1, 'w:gridAfter'),
                            (table.rows[1], 0, 'w:gridBefore')):
        row._tr.remove(row._tr.tc_lst[edge])
        omitted = OxmlElement(name)
        omitted.set(qn('w:val'), '1')
        row._tr.get_or_add_trPr().append(omitted)
    original = word_bytes(doc)
    with pytest.raises(ConversionError, match='列位置') as failure:
        convert_file(store, 'offset-table.docx', original, template_id)
    _, _, retained = source_files(store, failure.value.source_id)
    assert retained.read_bytes() == original
    assert store.rows('SELECT * FROM briefs') == []
    assert store.rows('SELECT * FROM jobs') == []


def test_original_word_to_template_preserves_body_and_table_without_model(workspace, monkeypatch):
    store, template_id = workspace
    monkeypatch.setattr(store, 'runtime_config', lambda: pytest.fail('conversion must not require a model'))
    previous = {'title': '原来的新报告配置', 'allow_web': True, 'template_id': 'unchanged'}
    store.set_meta('requirements', previous)
    doc = Document()
    doc.add_heading('业务原文', 1)
    p = doc.add_paragraph()
    p.add_run('本期增长 12%。').bold = True
    p.add_run(' [1] 和 [@src_original] 是原文。').italic = True
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), doc.part.relate_to('https://example.com/evidence', RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = '原文链接'
    run.append(text)
    link.append(run)
    p._p.append(link)
    doc.add_paragraph('第一项', 'List Number')
    doc.add_paragraph('第二项', 'List Number')
    doc.add_paragraph('项目符号', 'List Bullet')
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = '合并表头'
    table.cell(1, 0).merge(table.cell(2, 0)).text = '跨行内容'
    table.cell(1, 1).text = '12'
    table.cell(1, 2).text = '万元'
    table.cell(2, 1).text = '14'
    table.cell(2, 2).text = '万元'
    doc.add_heading('来源', 2)
    doc.add_paragraph('这一节以及后面的反馈文字必须保留。')
    doc.add_heading('反馈', 2)
    doc.add_paragraph('用户写下的结尾。')
    original = word_bytes(doc)
    result = convert_file(store, '原稿.docx', original, template_id)
    assert store.meta('requirements') == previous
    assert store.one('runs', result['version']['run_id'])['mode'] == 'normal'
    version = result['version']
    assert version['author'] == 'user' and version['parent_id'] is None
    document = json.loads(version['editor_document'])
    all_nodes = list(nodes(document))
    assert '这一节以及后面的反馈文字必须保留。' in version['markdown']
    assert '用户写下的结尾。' in version['markdown']
    assert all(n['type'] != 'citation' for n in all_nodes)
    assert any(n.get('text') == '1. ' for n in all_nodes)
    assert any(n.get('text') == '2. ' for n in all_nodes)
    assert any(n.get('attrs', {}).get('colspan') == 3 for n in all_nodes)
    assert any(n.get('attrs', {}).get('rowspan') == 2 for n in all_nodes)
    assert any(m['type'] == 'link' for n in all_nodes for m in n.get('marks', []))
    assert any(m['type'] == 'bold' for n in all_nodes for m in n.get('marks', []))
    _, _, retained = source_files(store, result['source_id'])
    assert retained.read_bytes() == original
    req = json.loads(store.one('runs', version['run_id'])['requirements'])
    assert req['allow_web'] is False and req['fact_check'] is False
    assert (req['workflow_id'], req['workflow_variant']) == ('general_report', 'general')
    assert json.loads(version['detail'])['citations'] == []
    generate_word(store, result['job'], threading.Event())
    rendered = Document(output_path(store, result['job']))
    body = '\n'.join(p.text for p in rendered.paragraphs)
    assert body.count('业务原文') == 1
    assert '1. 第一项' in body and '2. 第二项' in body
    assert '[1] 和 [@src_original] 是原文。' in body
    assert '用户写下的结尾。' in body
    assert rendered.tables[0].cell(0, 2).text == '合并表头'
    assert rendered.tables[0].cell(2, 0).text == '跨行内容'
    assert_no_research(store)


@pytest.mark.parametrize('suffix', ['.md', '.txt'])
def test_markdown_and_txt_keep_literal_source_markers_without_feedback(workspace, suffix):
    store, template_id = workspace
    text = '# 原文\n\n**保留内容** [@src_original]\n\n4. 第四项\n5. 第五项\n'
    result = convert_file(store, 'input' + suffix, text.encode(), template_id)
    document = json.loads(result['version']['editor_document'])
    all_nodes = list(nodes(document))
    assert not any(n['type'] == 'citation' for n in all_nodes)
    generate_word(store, result['job'], threading.Event())
    rendered = Document(output_path(store, result['job']))
    body = '\n'.join(p.text for p in rendered.paragraphs)
    assert '[@src_original]' in body and '4. 第四项' in body and '5. 第五项' in body
    if suffix == '.txt':
        assert all(n['type'] in ('doc', 'paragraph', 'text') for n in all_nodes)
        assert '# 原文' in body and '**保留内容**' in body
    else:
        assert any(n['type'] == 'heading' for n in all_nodes)
        assert any(m['type'] == 'bold' for n in all_nodes for m in n.get('marks', []))
    assert_no_research(store)


@pytest.mark.parametrize('case', ['html', 'image', 'markdown_footnote', 'nested_list', 'control', 'field', 'tracked', 'footnote', 'header',
                                'inherited_superscript', 'inherited_hidden', 'wrapped_table_row', 'page_break', 'column_break',
                                'horizontal_merge', 'multiblock_list', 'section_break'])
def test_unsupported_objects_retain_original_without_partial_version(workspace, case):
    store, template_id = workspace
    if case == 'control':
        name, data = 'source.txt', b'Cannot render this: \x00'
    elif case == 'markdown_footnote':
        name, data = 'source.md', '正文[^1]\n\n[^1]: 不能丢失'.encode()
    elif case == 'multiblock_list':
        name, data = 'source.md', '1. 第一项\n\n   属于第一项的续段\n\n2. 第二项\n'.encode()
    elif case == 'nested_list':
        name, data = 'source.md', '1. 一级事项\n   1. 从属于一级事项的内容\n'.encode()
    elif case in ('html', 'image'):
        name = 'source.md'
        data = ('正文\n\n<div>不可丢失</div>' if case == 'html' else '正文\n\n![不可丢失](photo.png)').encode()
    else:
        name = 'source.docx'
        doc = Document()
        p = doc.add_paragraph('可读取正文')
        if case == 'section_break':
            from docx.enum.section import WD_SECTION
            doc.add_section(WD_SECTION.NEW_PAGE)
            doc.add_paragraph('新节正文')
        elif case in ('page_break', 'column_break'):
            element = OxmlElement('w:br')
            element.set(qn('w:type'), case.removesuffix('_break'))
            p.add_run()._r.append(element)
        elif case == 'horizontal_merge':
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = '横向合并内容'
            for cell, value in zip(table.rows[0].cells, ['restart', 'continue']):
                merge = OxmlElement('w:hMerge')
                merge.set(qn('w:val'), value)
                cell._tc.get_or_add_tcPr().append(merge)
        elif case.startswith('inherited_'):
            base = doc.styles.add_style('Imported Base', WD_STYLE_TYPE.CHARACTER)
            setattr(base.font, case.removeprefix('inherited_'), True)
            derived = doc.styles.add_style('Imported Derived', WD_STYLE_TYPE.CHARACTER)
            derived.base_style = base
            p.add_run('不能改成普通正文').style = derived
        elif case == 'wrapped_table_row':
            table = doc.add_table(rows=2, cols=1)
            table.cell(0, 0).text = '普通表格行'
            table.cell(1, 0).text = '不能丢失的第二行'
            wrapper = OxmlElement('w:customXml')
            wrapper.append(table.rows[1]._tr)
            table._tbl.append(wrapper)
        elif case == 'header':
            doc.sections[0].header.paragraphs[0].text = '不能丢失的页眉'
        elif case == 'tracked':
            element = OxmlElement('w:ins')
            p._p.append(element)
        else:
            element = OxmlElement('w:fldChar' if case == 'field' else 'w:footnoteReference')
            p.add_run()._r.append(element)
        data = word_bytes(doc)
    with pytest.raises(ConversionError, match='原件已保留') as failure:
        convert_file(store, name, data, template_id)
    _, _, original = source_files(store, failure.value.source_id)
    assert original.read_bytes() == data
    assert store.rows('SELECT * FROM briefs') == []
    assert store.rows('SELECT * FROM runs') == []
    assert store.rows('SELECT * FROM jobs') == []
    assert_no_research(store)


def test_unready_template_rejected_before_storing_upload(workspace):
    store, template_id = workspace
    with store.tx() as connection:
        connection.execute("UPDATE templates SET status='preparing' WHERE id=?", (template_id,))
    with pytest.raises(ValueError, match='尚未准备完成'):
        convert_file(store, 'original.txt', b'Preserve this.', template_id)
    assert store.rows('SELECT * FROM sources') == []
    assert list((store.root / 'sources').glob('*.original.*')) == []
