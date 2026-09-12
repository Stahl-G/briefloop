"""The bundled built-in template registers, prepares and exports like an uploaded one."""
from io import BytesIO
from zipfile import ZipFile
import threading
from docx import Document
from briefloop.store import Store
from briefloop.templates import import_builtin, template
from briefloop.export_jobs import enqueue_export, generate_word, output_path


def test_builtin_ships_prepares_is_idempotent_and_exports(tmp_path):
    store = Store(tmp_path)
    import_builtin(store)
    rows = store.rows('SELECT id,name,status,origin FROM templates')
    assert [(r['name'], r['status'], r['origin']) for r in rows] == [('研报版式', 'ready', 'builtin')]
    record = template(store, rows[0]['id'])
    assert [s['section_id'] for s in record['spec']['sections']] == ['summary', 'dynamics', 'data', 'analysis', 'risks', 'appendix']
    import_builtin(store)
    assert len(store.rows('SELECT id FROM templates')) == 1

    source = store.add_source('Synthetic material', 'Synthetic evidence for the report.')
    run = store.create_run({'title': 'AI 行业周报', 'objective': 'Explain', 'period': '2026 年第 37 周',
                            'organization': '示例机构', 'template_id': record['id']}, [source['id']])
    document = {'type': 'doc', 'content': [
        {'type': 'heading', 'attrs': {'level': 2}, 'content': [{'type': 'text', 'text': '一、摘要'}]},
        {'type': 'paragraph', 'content': [{'type': 'text', 'text': '本期核心判断保持不变。'}]},
        {'type': 'heading', 'attrs': {'level': 2}, 'content': [{'type': 'text', 'text': '1.1 市场规模'}]},
        {'type': 'paragraph', 'content': [{'type': 'text', 'text': '规模数据保持增长。'}]},
        {'type': 'table', 'content': [
            {'type': 'tableRow', 'content': [
                {'type': 'tableHeader', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': '指标'}]}]},
                {'type': 'tableHeader', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': '本期'}]}]}]},
            {'type': 'tableRow', 'content': [
                {'type': 'tableCell', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': '装机容量'}]}]},
                {'type': 'tableCell', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': '12 GW'}]}]}]}]}]}
    brief = store.publish(run['id'], {'title': 'AI 行业周报', 'editor_document': document})
    job = enqueue_export(store, brief['id'])
    generate_word(store, job, threading.Event())
    with ZipFile(output_path(store, job)) as archive:
        names = archive.namelist()
        document_xml = archive.read('word/document.xml').decode()
        styles_xml = archive.read('word/styles.xml').decode()
        footers = ''.join(archive.read(n).decode() for n in names if n.startswith('word/footer'))
        headers = ''.join(archive.read(n).decode() for n in names if n.startswith('word/header'))
        assert '{{' not in document_xml and '{{' not in headers
        assert 'AI 行业周报' in document_xml and '示例机构' in document_xml and '2026 年第 37 周' in document_xml
        assert 'w:instr=" PAGE "' in footers or ' PAGE ' in footers
        assert 'w:eastAsia="黑体"' in styles_xml and 'w:eastAsia="宋体"' in styles_xml
        assert 'w:tblBorders' in document_xml and 'w:insideV' in document_xml
        from lxml import etree
        tree = etree.fromstring(archive.read('word/document.xml'))
        text_of = lambda p: ''.join(t.text or '' for t in p.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
        style_of = lambda p: p.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle')
        paragraphs = tree.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p')
        subsection = next(p for p in paragraphs if text_of(p) == '1.1 市场规模')
        assert style_of(subsection).get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') == 'Heading2'
        chapter = next(p for p in tree.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p') if text_of(p) == '一、摘要')
        chapter_style = style_of(chapter).get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
        assert chapter_style.replace(' ', '').startswith('BLHeadingsummary')


def test_uploaded_templates_keep_upload_origin(tmp_path):
    from docx import Document as NewDocument
    store = Store(tmp_path)
    original = NewDocument()
    original.add_paragraph('Monthly Report')
    original.add_paragraph('1. Summary')
    stream = BytesIO()
    original.save(stream)
    from briefloop.templates import import_template
    record = import_template(store, 'Monthly.docx', stream.getvalue(), prepare_job=False)
    assert record['origin'] == 'upload'
