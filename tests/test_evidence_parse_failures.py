"""Actual malformed documents must only invalidate their own saved bindings."""
import json
from io import BytesIO
from zipfile import ZipFile
import pytest
from openpyxl import Workbook
from pypdf import PdfWriter
from briefloop.store import Store
from briefloop.evidence import create_span,create_claim,bind_claim,inspect_bindings
from briefloop.sources import upload
from briefloop.media import source_files


def attached(store,kind):
    stream=BytesIO()
    if kind=='pdf':
        writer=PdfWriter();writer.add_blank_page(width=100,height=100);writer.write(stream)
        locator={'kind':'pdf','page':1}
    else:
        book=Workbook();book.active.title='Facts';book.active['A1']='Synthetic';book.save(stream)
        locator={'kind':'xlsx','sheet':'Facts','cells':'A1'}
    source=upload(store,'synthetic.'+kind,stream.getvalue())
    return source,locator


def corrupt(store,source,kind,variant):
    _,_,original=source_files(store,source['id'])
    if variant in ('xml','missing_member'):
        buffer=BytesIO()
        with ZipFile(original) as archive,ZipFile(buffer,'w') as damaged:
            for name in archive.namelist():
                if variant=='missing_member' and name=='xl/workbook.xml':continue
                damaged.writestr(name,b'<broken' if variant=='xml' and name=='xl/worksheets/sheet1.xml' else archive.read(name))
        original.write_bytes(buffer.getvalue())
    else:original.write_bytes(b'%PDF-1.7\ntruncated' if kind=='pdf' else b'PK\x03\x04truncated')
    # Legacy imports may have no raw checksum. Exercise parsing rather than
    # stopping earlier at the separate, already supported checksum rejection.
    provenance=store.root/'sources'/(source['id']+'.provenance.json')
    data=json.loads(provenance.read_text());data.pop('raw_sha256',None);provenance.write_text(json.dumps(data))


@pytest.mark.parametrize('kind,variant',[('pdf','truncated'),('xlsx','truncated'),('xlsx','xml'),('xlsx','missing_member')])
def test_bad_original_is_local_and_new_span_rejected(tmp_path,kind,variant):
    store=Store(tmp_path);bad,locator=attached(store,kind);good=store.add_source('Healthy','Healthy evidence.')
    run=store.create_run({'title':'Synthetic','objective':'Inspect bindings'},[bad['id'],good['id']])
    brief=store.publish(run['id'],{'title':'Synthetic','editor_document':{'type':'doc','content':[{'type':'paragraph','content':[{'type':'text','text':'Synthetic assertion. Healthy assertion.'}]}]}})
    block=json.loads(brief['editor_document'])['content'][0]['attrs']['blockId']
    for source,loc,quote in [(bad,locator,'Synthetic assertion.'),(good,{'kind':'text','start_line':1,'end_line':1},'Healthy assertion.')]:
        span=create_span(store,{'source_id':source['id'],'locator':loc})
        claim=create_claim(store,run['id'],{'statement':quote,'kind':'fact','supports':[{'span_id':span['id'],'supports_quote':quote}]})
        bind_claim(store,brief['id'],claim['id'],block,quote)
    corrupt(store,bad,kind,variant)
    result=inspect_bindings(store,brief['id']);assert len(result['bindings'])==2
    by_source={row['evidence'][0]['source_id']:row for row in result['bindings']}
    assert by_source[good['id']]['status']=='unreviewed' and by_source[good['id']]['evidence'][0]['intact']
    damaged=by_source[bad['id']];assert damaged['status']=='source_changed'
    assert damaged['evidence'][0]['intact'] is False
    assert damaged['evidence'][0]['location_error']['code']==kind+'_parse_error'
    assert damaged['evidence'][0]['location_error']['source_id']==bad['id']
    with pytest.raises(ValueError,match='原件无法解析'):create_span(store,{'source_id':bad['id'],'locator':locator})


def test_unexpected_parser_bug_is_not_hidden(tmp_path,monkeypatch):
    store=Store(tmp_path);source,locator=attached(store,'pdf')
    def bug(*args,**kwargs):raise RuntimeError('unexpected implementation bug')
    monkeypatch.setattr('pypdf.PdfReader',bug)
    with pytest.raises(RuntimeError,match='implementation bug'):create_span(store,{'source_id':source['id'],'locator':locator})


def test_unrelated_workbook_keyerror_remains_visible(tmp_path,monkeypatch):
    store=Store(tmp_path);source,locator=attached(store,'xlsx')
    def bug(*args,**kwargs):raise KeyError('unexpected internal state')
    monkeypatch.setattr('openpyxl.load_workbook',bug)
    with pytest.raises(KeyError,match='internal state'):create_span(store,{'source_id':source['id'],'locator':locator})


def test_partial_workbook_load_closes_both_owned_streams(tmp_path,monkeypatch):
    import openpyxl
    from zipfile import BadZipFile
    store=Store(tmp_path);source,locator=attached(store,'xlsx');load=openpyxl.load_workbook
    streams=[];closed=[]
    def tracked(stream,**kwargs):
        streams.append(stream)
        if len(streams)==2:raise BadZipFile('Synthetic second-load failure')
        book=load(stream,**kwargs);close=book.close
        def finish():closed.append(True);close()
        book.close=finish
        return book
    monkeypatch.setattr(openpyxl,'load_workbook',tracked)
    with pytest.raises(ValueError,match='原件无法解析'):create_span(store,{'source_id':source['id'],'locator':locator})
    assert len(streams)==2 and all(s.closed for s in streams) and closed==[True]
