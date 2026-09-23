"""A native host failure cannot discard an already admitted Word template."""
from hashlib import sha256
from io import BytesIO
import json
import threading
from types import SimpleNamespace

from docx import Document
import pytest

from briefloop.export_jobs import enqueue_export, generate_word, output_path
from briefloop.native_orchestrator import template_submit
from briefloop.runtime import Worker
from briefloop.store import Store, dump
from briefloop.templates import import_template, prepare, template


def template_job(tmp_path):
    store = Store(tmp_path)
    original = Document()
    original.add_paragraph('Monthly Report')
    original.add_heading('Summary', level=1)
    original.add_paragraph('Old body text.')
    buffer = BytesIO()
    original.save(buffer)
    record = import_template(store, 'Monthly.docx', buffer.getvalue(), prepare_job=False)
    job = store.enqueue('prepare_template', {'template_id': record['id']})
    store.update_job(job['id'], 'running')
    return store, record, store.one('jobs', job['id'])


def template_spec():
    return {
        'keep_blocks': [0], 'paragraph_index': 2,
        'sections': [{'section_id': 'summary', 'title': 'Summary', 'index': 1}],
        'fields': [{'old': 'Monthly Report', 'field': 'title'}],
    }


def save_template(store, record):
    return prepare(store, record['id'], template_spec())


class Host:
    def __init__(self, action):
        self.action = action
        self.cancelled = threading.Event()

    def execute(self, *args, **kwargs):
        self.action()
        raise RuntimeError('host disconnected after the template turn')

    def cancel(self):
        self.cancelled.set()


def test_saved_template_survives_late_host_failure_and_queued_word_export(tmp_path):
    store, record, job = template_job(tmp_path)

    def submit_then_fail():
        template_submit(store, {'template_id': record['id'], 'job_id': job['id'],
            'packet_root': str(store.root / 'jobs' / job['id'] / 'packet'),
            'session_id': 'native-fixture',
            '_harness': SimpleNamespace(cancel_requested=lambda session_id: False)}, template_spec())

    worker = Worker(store, Host(submit_then_fail))
    worker._execute_main_job(job)

    prepared = template(store, record['id'])
    finished = store.one('jobs', job['id'])
    assert prepared['status'] == 'ready'
    assert finished['status'] == 'complete' and finished['error'] is None
    assert json.loads(finished['result']) == {
        'template_id': record['id'], 'revision': record['revision'], 'status': 'ready'}
    assert prepared['spec']['prepared_hash'] == sha256(
        (store.root / 'templates' / record['id'] / 'prepared.docx').read_bytes()).hexdigest()
    assert (store.root / 'jobs' / job['id'] / 'template.json').exists()

    source = store.add_source('Synthetic material', 'Current evidence.')
    run = store.create_run({'title': 'Current Report', 'objective': 'Explain',
                            'template_id': record['id']}, [source['id']])
    document = {'type': 'doc', 'content': [{'type': 'paragraph',
        'content': [{'type': 'text', 'text': 'Current body text.'}]}]}
    brief = store.publish(run['id'], {'title': 'Current Report', 'editor_document': document})
    export = enqueue_export(store, brief['id'])
    result = generate_word(store, export, threading.Event())
    assert result['version_id'] == brief['id']
    assert 'Current body text.' in '\n'.join(p.text for p in Document(output_path(store, export)).paragraphs)


def test_failure_before_save_keeps_template_and_job_failed(tmp_path):
    store, record, job = template_job(tmp_path)
    Worker(store, Host(lambda: None))._execute_main_job(job)
    assert template(store, record['id'])['status'] == 'failed'
    assert store.one('jobs', job['id'])['status'] == 'failed'


@pytest.mark.parametrize('damage', ['missing', 'hash_mismatch', 'invalid_archive'])
def test_bad_saved_file_does_not_turn_host_failure_into_success(tmp_path, damage):
    store, record, job = template_job(tmp_path)

    def save_then_corrupt():
        save_template(store, record)
        path = store.root / 'templates' / record['id'] / 'prepared.docx'
        if damage == 'missing':
            path.unlink()
            return
        damaged = b'not a DOCX archive'
        path.write_bytes(damaged)
        if damage == 'invalid_archive':
            spec = template(store, record['id'])['spec']
            spec['prepared_hash'] = sha256(damaged).hexdigest()
            with store.tx() as connection:
                connection.execute('UPDATE templates SET spec=? WHERE id=?', (dump(spec), record['id']))

    Worker(store, Host(save_then_corrupt))._execute_main_job(job)
    assert store.one('jobs', job['id'])['status'] == 'failed'
    assert template(store, record['id'])['status'] == 'failed'


def test_cancel_after_save_keeps_cancelled_job_and_ready_template(tmp_path):
    store, record, job = template_job(tmp_path)
    worker = None

    def save_then_cancel():
        save_template(store, record)
        worker.stop_job(job['id'])

    worker = Worker(store, Host(save_then_cancel))
    worker.current = job['id']
    worker._execute_main_job(job)
    assert store.one('jobs', job['id'])['status'] == 'cancelled'
    assert template(store, record['id'])['status'] == 'ready'


def test_concurrent_ready_admission_is_not_overwritten_by_stale_failure(tmp_path):
    store, record, job = template_job(tmp_path)
    worker = Worker(store, Host(lambda: None))
    read_saved = worker._saved_template_result
    inspected = False

    def save_after_stale_read(template_id):
        nonlocal inspected
        if inspected:
            return read_saved(template_id)
        inspected = True
        observed = store.rows('SELECT id,revision,status,spec FROM templates WHERE id=?', (template_id,))[0]
        save_template(store, record)
        return None, observed

    worker._saved_template_result = save_after_stale_read
    worker._execute_main_job(job)
    assert store.one('jobs', job['id'])['status'] == 'complete'
    assert template(store, record['id'])['status'] == 'ready'
