import threading
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import pytest
from docx import Document

from briefloop import export_jobs
from briefloop.store import Store, dump


def _indexed_document(source_ids):
    def paragraph(text):
        return {'type': 'paragraph', 'content': [{'type': 'text', 'text': text}]}
    def row(values, kind):
        return {'type': 'tableRow', 'content': [
            {'type': kind, 'content': [paragraph(value)]} for value in values]}
    return {'type': 'doc', 'content': [
        {'type': 'paragraph', 'content': [{'type': 'citation', 'attrs': {'sourceId': source_ids[0]}}]},
        {'type': 'heading', 'attrs': {'level': 2}, 'content': [{'type': 'text', 'text': '来源'}]},
        {'type': 'table', 'content': [row(['编号', '主体/来源', 'source_id'], 'tableHeader'),
                                    *[row([str(i), 'Source', sid], 'tableCell')
                                      for i, sid in enumerate(source_ids, 1)]]}]}


def _complete_word(store, job):
    result = export_jobs.generate_word(store, job, threading.Event())
    with store.tx() as c:
        c.execute("UPDATE jobs SET status='complete',result=? WHERE id=?", (dump(result), job['id']))
    return Document(BytesIO(export_jobs.output_path(store, job).read_bytes()))


def _has_source_index(document, source_id):
    return any(len(table.columns) == 3
               and [cell.text.strip() for cell in table.rows[0].cells] == ['编号', '主体/来源', 'source_id']
               and any(row.cells[2].text.strip() == source_id for row in table.rows[1:])
               for table in document.tables if table.rows)


def test_source_projection_invalidates_word_cache_only_when_rendered_source_changes(tmp_path):
    from briefloop.templates import import_builtin
    store = Store(tmp_path)
    import_builtin(store)
    template_id = store.rows('SELECT id FROM templates LIMIT 1')[0]['id']
    first = store.add_source('First source', 'First evidence', url='https://example.org/first')
    later = store.add_source('Later source', 'Later evidence', url='https://example.org/later')
    unrelated = store.add_source('Unrelated', 'Other evidence')
    run = store.create_run({'title': 'Source cache', 'objective': 'Check Word'}, [first['id']])
    brief = store.publish(run['id'], {'title': 'Source cache',
                                      'editor_document': _indexed_document([first['id'], later['id']])})
    original = store.one('briefs', brief['id'])
    before_jobs = {}
    for layout in (None, template_id):
        before = export_jobs.enqueue_export(store, brief['id'], layout)
        before_jobs[layout] = before
        first_doc = _complete_word(store, before)
        assert _has_source_index(first_doc, later['id'])
        assert export_jobs.enqueue_export(store, brief['id'], layout)['id'] == before['id']
    store.attach_source(run['id'], later['id'])
    assert store.one('briefs', brief['id']) == original
    changed_jobs = {}
    for layout in (None, template_id):
        changed = export_jobs.enqueue_export(store, brief['id'], layout)
        changed_jobs[layout] = changed
        assert changed['id'] != before_jobs[layout]['id']
        with pytest.raises(ValueError, match='导出输入已变化'):
            export_jobs.generate_word(store, before_jobs[layout], threading.Event())
        new_doc = _complete_word(store, changed)
        assert not _has_source_index(new_doc, later['id'])
        assert any('Later source' in p.text for p in new_doc.paragraphs)
        assert export_jobs.enqueue_export(store, brief['id'], layout)['id'] == changed['id']
    store.attach_source(run['id'], unrelated['id'])
    for layout in (None, template_id):
        assert export_jobs.enqueue_export(store, brief['id'], layout)['id'] == changed_jobs[layout]['id']
    with store.tx() as c:
        c.execute('UPDATE sources SET name=? WHERE id=?', ('Corrected source', later['id']))
    renamed_jobs = {}
    for layout in (None, template_id):
        renamed = export_jobs.enqueue_export(store, brief['id'], layout)
        renamed_jobs[layout] = renamed
        assert renamed['id'] != changed_jobs[layout]['id']
        renamed_doc = _complete_word(store, renamed)
        assert any('Corrected source' in p.text for p in renamed_doc.paragraphs)
        assert export_jobs.enqueue_export(store, brief['id'], layout)['id'] == renamed['id']
    with store.tx() as c:
        c.execute('UPDATE sources SET url=? WHERE id=?', ('https://example.org/corrected', later['id']))
    for layout in (None, template_id):
        relinked = export_jobs.enqueue_export(store, brief['id'], layout)
        assert relinked['id'] != renamed_jobs[layout]['id']
        relinked_doc = _complete_word(store, relinked)
        assert any(rel.target_ref == 'https://example.org/corrected'
                   for rel in relinked_doc.part.rels.values() if rel.is_external)
        assert export_jobs.enqueue_export(store, brief['id'], layout)['id'] == relinked['id']


@pytest.mark.parametrize('use_template', [False, True])
def test_word_renders_fingerprinted_sources_when_the_run_changes_at_stage_two(tmp_path, monkeypatch, use_template):
    from briefloop.templates import import_builtin
    store = Store(tmp_path)
    layout = None
    if use_template:
        import_builtin(store)
        layout = store.rows('SELECT id FROM templates LIMIT 1')[0]['id']
    first = store.add_source('Original source', 'Initial evidence', url='https://example.org/original')
    later = store.add_source('Later source', 'Additional evidence', url='https://example.org/later')
    run = store.create_run({'title': 'Source snapshot', 'objective': 'Check Word'}, [first['id']])
    brief = store.publish(run['id'], {'title': 'Source snapshot',
                                      'editor_document': _indexed_document([first['id'], later['id']])})
    job = export_jobs.enqueue_export(store, brief['id'], layout)
    original_event = store.event
    changed = []

    def change_during_stage_two(job_id, kind, data):
        original_event(job_id, kind, data)
        if kind == 'export_progress' and data['step'] == 2 and not changed:
            changed.append(True)
            store.attach_source(run['id'], later['id'])
            with store.tx() as connection:
                connection.execute('UPDATE sources SET name=?,url=? WHERE id=?',
                                   ('Updated source', 'https://example.org/updated', first['id']))

    monkeypatch.setattr(store, 'event', change_during_stage_two)
    old_doc = _complete_word(store, job)
    assert changed
    assert _has_source_index(old_doc, later['id'])
    assert any('Original source' in p.text for p in old_doc.paragraphs)
    assert not any('Updated source' in p.text or 'Later source' in p.text for p in old_doc.paragraphs)
    old_urls = {rel.target_ref for rel in old_doc.part.rels.values() if rel.is_external}
    assert 'https://example.org/original' in old_urls
    assert 'https://example.org/updated' not in old_urls

    new_job = export_jobs.enqueue_export(store, brief['id'], layout)
    assert new_job['id'] != job['id']
    new_doc = _complete_word(store, new_job)
    assert not _has_source_index(new_doc, later['id'])
    assert any('Updated source' in p.text for p in new_doc.paragraphs)
    assert any('Later source' in p.text for p in new_doc.paragraphs)
    new_urls = {rel.target_ref for rel in new_doc.part.rels.values() if rel.is_external}
    assert 'https://example.org/updated' in new_urls


def test_concurrent_exports_share_job_and_rebuild_missing_or_damaged_file(tmp_path, monkeypatch):
    store = Store(tmp_path)
    source = store.add_source('Synthetic', '本周两项交付。')
    run = store.create_run({'title': '验收', 'objective': '合成测试'}, [source['id']])
    brief = store.publish(run['id'], {'title': '验收', 'markdown': '# 验收\n\n本周两项交付。'})
    # Independent Store instances model separate request handlers/processes.
    stores = [Store(tmp_path) for _ in range(4)]
    original = export_jobs.export_input
    ready = threading.Barrier(len(stores))
    def concurrent_input(*args, **kwargs):
        result = original(*args, **kwargs)
        ready.wait(timeout=10)
        return result
    monkeypatch.setattr(export_jobs, 'export_input', concurrent_input)
    with ThreadPoolExecutor(max_workers=len(stores)) as pool:
        jobs = list(pool.map(lambda s: export_jobs.enqueue_export(s, brief['id']), stores))
    assert len({job['id'] for job in jobs}) == 1
    monkeypatch.setattr(export_jobs, 'export_input', original)
    job = jobs[0]
    result = export_jobs.generate_word(store, job, threading.Event())
    with store.tx() as c:
        c.execute("UPDATE jobs SET status='complete',result=? WHERE id=?", (dump(result), job['id']))
    assert export_jobs.enqueue_export(store, brief['id'])['id'] == job['id']
    export_jobs.output_path(store, job).write_bytes(b'corrupt')
    rebuilt = export_jobs.enqueue_export(store, brief['id'])
    assert rebuilt['id'] != job['id']
    result = export_jobs.generate_word(store, rebuilt, threading.Event())
    with store.tx() as c:
        c.execute("UPDATE jobs SET status='complete',result=? WHERE id=?", (dump(result), rebuilt['id']))
    export_jobs.output_path(store, rebuilt).unlink()
    assert export_jobs.enqueue_export(store, brief['id'])['id'] != rebuilt['id']
    revised = store.revise(brief['id'], '# 验收\n\n本周三项交付。')
    changed = export_jobs.enqueue_export(store, revised['id'])
    assert changed['id'] not in {job['id'] for job in store.rows("SELECT * FROM jobs WHERE kind='export_docx'") if job['payload'] == rebuilt['payload']}


def test_reader_source_upgrade_invalidates_both_renderer_caches(tmp_path, monkeypatch):
    from briefloop.templates import import_builtin
    store = Store(tmp_path)
    import_builtin(store)
    template_id = store.rows('SELECT id FROM templates LIMIT 1')[0]['id']
    source = store.add_source('Synthetic', '本周两项交付。')
    run = store.create_run({'title': '验收', 'objective': '合成测试'}, [source['id']])
    brief = store.publish(run['id'], {'title': '验收', 'markdown': '# 验收\n\n本周两项交付。'})
    current_input = export_jobs.export_input
    def old_input(*args, **kwargs):
        identity, figures = current_input(*args, **kwargs)
        identity['renderer'] = 24 if identity['requirements'].get('template_id') else 25
        return identity, figures
    for layout in (None, template_id):
        monkeypatch.setattr(export_jobs, 'export_input', old_input)
        old = export_jobs.enqueue_export(store, brief['id'], layout)
        result = export_jobs.generate_word(store, old, threading.Event())
        with store.tx() as c:
            c.execute("UPDATE jobs SET status='complete',result=? WHERE id=?", (dump(result), old['id']))
        assert export_jobs.enqueue_export(store, brief['id'], layout)['id'] == old['id']
        monkeypatch.setattr(export_jobs, 'export_input', current_input)
        upgraded = export_jobs.enqueue_export(store, brief['id'], layout)
        assert upgraded['id'] != old['id']
        assert export_jobs.enqueue_export(store, brief['id'], layout)['id'] == upgraded['id']
        assert export_jobs.output_path(store, old).is_file()
