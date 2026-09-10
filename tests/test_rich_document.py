from io import BytesIO
from zipfile import ZipFile
import copy
import json
import subprocess
from pathlib import Path
import pytest
from PIL import Image
from briefloop.store import Store, Conflict
from briefloop.document_model import normalize_document, document_markdown, markdown_document
from briefloop.exports import docx_bytes
from briefloop.figures import register_figure
from briefloop.figure_support import export_figures


def para(text, **attrs):
    return {'type': 'paragraph', 'attrs': attrs, 'content': [{'type': 'text', 'text': text}]}


def test_default_fonts_keep_bilingual_content_format_and_unicode_bullets():
    from lxml import etree
    from docx import Document
    W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    A='http://schemas.openxmlformats.org/drawingml/2006/main'
    ns={'w':W,'a':A,'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    markdown='# 中文与English\n\n保留正文提到的 Calibri 名称和 **重点数字12.5%**。\n\n- 项目甲\n- Bullet B\n\n| 指标 | Value |\n|---|---|\n|收入|12.5|\n\n```text\n数据 = 12.5\n```'
    rich=markdown_document(markdown)
    rich['content'][1]['content'][0].setdefault('marks',[]).append({'type':'textStyle','attrs':{'color':'#C00000'}})
    # Exercise rich generic and retained Markdown industry entry points, which
    # both start from a stock default DOCX and have different typography setup.
    outputs=[docx_bytes(document=rich),docx_bytes(markdown,report_profile='industry_periodic',title='中文与English',language='中文')]
    for payload in outputs:
        doc=Document(BytesIO(payload))
        assert any('保留正文提到的 Calibri 名称' in p.text for p in doc.paragraphs)
        assert any(run.bold for p in doc.paragraphs for run in p.runs if '重点数字' in run.text)
        assert doc.tables[0].cell(1,0).text=='收入'
        assert any('数据 = 12.5' in p.text for p in doc.paragraphs)
        with ZipFile(BytesIO(payload)) as archive:
            styles=etree.fromstring(archive.read('word/styles.xml'))
            assert styles.find('w:docDefaults/w:rPrDefault/w:rPr/w:lang',ns).get('{'+W+'}eastAsia')=='zh-CN'
            for name in archive.namelist():
                if not name.startswith('word/') or not name.endswith('.xml'):continue
                xml=etree.fromstring(archive.read(name))
                for fonts in xml.findall('.//w:rFonts',ns):
                    values={etree.QName(key).localname:value for key,value in fonts.attrib.items()}
                    assert values.get('ascii')=='Arial' and values.get('hAnsi')=='Arial'
                    assert 'eastAsia' not in values and not any('theme' in key.lower() for key in values)
                assert all(node.get('{'+W+'}name')=='Arial' for node in xml.findall('.//w:font',ns))
                assert all(node.get('typeface')=='Arial' for node in xml.findall('.//a:fontScheme//a:latin',ns))
                assert not xml.findall('.//a:fontScheme//a:font',ns)
                assert not xml.findall('.//m:mathFont',ns)
                for level in xml.findall('.//w:lvl',ns):
                    kind=level.find('w:numFmt',ns);text=level.find('w:lvlText',ns)
                    if kind is not None and kind.get('{'+W+'}val')=='bullet':assert text.get('{'+W+'}val')=='•'
    with ZipFile(BytesIO(outputs[0])) as archive:assert b'C00000' in archive.read('word/document.xml')
    # Explicit non-Simplified-Chinese requests do not inherit the Chinese
    # heuristic merely because their text also contains Han characters.
    explicit=docx_bytes('日本語の報告',language='ja-JP')
    with ZipFile(BytesIO(explicit)) as archive:
        styles=etree.fromstring(archive.read('word/styles.xml'))
        assert all(node.get('{'+W+'}eastAsia')!='zh-CN' for node in styles.findall('.//w:lang',ns))


def test_rich_edit_survives_reload_and_word_without_markdown_loss(tmp_path):
    store = Store(tmp_path)
    source = store.add_source('Evidence', 'Revenue 12, prior 13.')
    run = store.create_run({'title': 'Report', 'objective': 'Explain'}, [source['id']])
    path = store.root / 'chart.png'; Image.new('RGB', (200, 100), 'blue').save(path)
    figure = register_figure(store, run['id'], path, 'Trend', 'USD; observed period', [source['id']])
    document = {'type': 'doc', 'content': [
        {'type': 'heading', 'attrs': {'level': 2, 'blockId': 'summary'}, 'content': [{'type': 'text', 'text': 'Results'}]},
        para('Revenue 12'),
        {'type': 'table', 'content': [
            {'type': 'tableRow', 'content': [{'type': 'tableHeader', 'attrs': {'colspan': 2}, 'content': [para('Values')]}]},
            {'type': 'tableRow', 'content': [{'type': 'tableCell', 'attrs': {'rowspan': 2}, 'content': [para('Revenue')]}, {'type': 'tableCell', 'content': [para('12')]}]},
            {'type': 'tableRow', 'content': [{'type': 'tableCell', 'content': [para('13')]}]},
        ]},
        {'type': 'image', 'attrs': {'src': 'briefloop-figure:' + figure['figure_id'], 'width': 150, 'alt': 'Trend', 'caption': 'Current caption'}},
        {'type': 'paragraph', 'content': [{'type': 'citation', 'attrs': {'sourceId': source['id']}}]},
    ]}
    before = store.publish(run['id'], {'title': 'Report', 'editor_document': document})
    edited = copy.deepcopy(document)
    edited['content'][1]['content'][0]['marks'] = [{'type': 'textStyle', 'attrs': {'color': '#C00000'}}]
    edited['content'][2]['content'][0]['content'][0]['attrs']['backgroundColor'] = '#EAF0FF'
    assert document_markdown(document) == document_markdown(edited)
    after = store.revise(before['id'], editor_document=edited)
    assert after['hash'] != before['hash'] and after['parent_id'] == before['id']
    assert not store.rows('SELECT * FROM feedback')  # Color changes are saved, not prose corrections.
    assert store.revise(after['id'], editor_document=edited)['id'] == after['id']
    restored = Store(tmp_path).one('briefs', after['id'])
    assert json.loads(restored['editor_document']) == normalize_document(edited)
    blob = docx_bytes(document=json.loads(restored['editor_document']), figures=export_figures(store, restored), source_records={source['id']: source})
    with ZipFile(BytesIO(blob)) as archive:
        xml = archive.read('word/document.xml').decode().lower()
        assert 'c00000' in xml and 'eaf0ff' in xml
        assert 'w:gridspan w:val="2"' in xml and 'w:vmerge' in xml
        assert 'current caption' in xml and 'evidence' in xml and 'w:drawing' in xml
        assert 'briefloop-figure:' not in xml and '***' not in xml
    with pytest.raises(Conflict): store.revise(before['id'], editor_document=edited)


def test_legacy_import_and_complete_publication_identity(tmp_path):
    md = '# Heading\n\nText **bold** and *italic*.\n\n| A | B |\n|---|---|\n| 1 | 2 |'
    document = markdown_document(md)
    assert any(node['type'] == 'table' for node in document['content'])
    store = Store(tmp_path); source = store.add_source('source', 'text')
    run = store.create_run({'title': 'Test', 'objective': 'Test'}, [source['id']])
    old = store.publish(run['id'], {'title': 'Test', 'markdown': md})
    new = store.revise(old['id'], editor_document=document)
    assert store.one('briefs', old['id'])['editor_document'] is None
    assert new['editor_document']
    with pytest.raises(Conflict):
        store.publish(run['id'], {'title': 'Changed metadata', 'markdown': md}, version_id=old['id'])
    with pytest.raises(ValueError):
        normalize_document({'type': 'doc', 'content': [{'type': 'image', 'attrs': {'src': 'file:///private/data.png'}}]})


def test_real_editor_schema_accepts_color_spans_and_citations():
    repo = Path(__file__).parents[1]
    script = '''
import {getSchema} from '@tiptap/core';
import StarterKit from '@tiptap/starter-kit';
import {TableKit} from '@tiptap/extension-table';
import {MarkdownManager,Markdown} from '@tiptap/markdown';
import {TextStyle,Layout,ReportImage,Citation} from './frontend/rich-document.js';
const extensions=[StarterKit,TableKit,ReportImage,TextStyle,Layout,Citation,Markdown];
const d={type:'doc',content:[{type:'paragraph',attrs:{textAlign:'center'},content:[{type:'text',text:'Color',marks:[{type:'textStyle',attrs:{color:'#c00000'}},{type:'bold'}]},{type:'citation',attrs:{sourceId:'src_test'}}]}]};
getSchema(extensions).nodeFromJSON(d).check();
d.content.push({type:'table',content:[{type:'tableRow',content:[{type:'tableHeader',content:[{type:'paragraph',content:[{type:'text',text:'Value'}]}]}]}]});
console.log(JSON.stringify(getSchema(extensions).nodeFromJSON(d).toJSON()));
const markdown=new MarkdownManager({extensions}).serialize(d);
if(!markdown.includes('Color')||!markdown.includes('[@src_test]'))throw Error(markdown);
'''
    result = subprocess.run(['node', '--input-type=module', '-e', script], cwd=repo, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    normalized=normalize_document(json.loads(result.stdout))
    assert normalized['content'][-1]['type']=='table'


def test_word_job_keeps_clicked_version_while_user_edits(tmp_path):
    import threading
    from briefloop.export_jobs import enqueue_export, generate_word, output_path, export_input
    store=Store(tmp_path);source=store.add_source('Evidence','Revenue 12')
    run=store.create_run({'title':'Report','objective':'Explain'},[source['id']])
    before=store.publish(run['id'],{'title':'Report','editor_document':{'type':'doc','content':[para('Revenue 12')]}})
    assert export_input(store,before)[0]['renderer']==4
    job=enqueue_export(store,before['id'])
    assert enqueue_export(store,before['id'])['id']==job['id']
    after=store.revise(before['id'],editor_document={'type':'doc','content':[para('Revenue 14')]})
    result=generate_word(store,job,threading.Event())
    store.update_job(job['id'],'complete',result=result)
    with ZipFile(output_path(store,job)) as archive:
        xml=archive.read('word/document.xml').decode()
        assert 'Revenue 12' in xml and 'Revenue 14' not in xml
    assert result['version_id']==before['id']
    assert enqueue_export(store,after['id'])['id']!=job['id']
