from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
from briefloop.document_export import render_document, reader_source_blocks


def fixture():
    def p(text):
        return {'type': 'paragraph', 'content': [{'type': 'text', 'text': text}]}
    def row(values, kind):
        return {'type': 'tableRow', 'content': [{'type': kind, 'content': [p(v)]} for v in values]}
    return {'type': 'doc', 'content': [
        {'type': 'paragraph', 'content': [{'type': 'citation', 'attrs': {'sourceId': 'src_one'}}]},
        {'type': 'heading', 'attrs': {'level': 2}, 'content': [{'type': 'text', 'text': '来源'}]},
        {'type': 'table', 'content': [row(['编号', '主体/来源', 'source_id'], 'tableHeader'),
                                    row(['1', 'Source', 'src_one'], 'tableCell')]}]}


def test_reader_export_replaces_internal_index_with_linked_title_without_mutating_draft():
    draft = fixture(); original = deepcopy(draft); doc = Document()
    render_document(doc, draft, sources={'src_one': {'name': 'Official source', 'url': 'https://example.org/report'}})
    assert draft == original
    assert len(doc.tables) == 0
    assert sum(p.text == '来源' for p in doc.paragraphs) == 1
    assert 'src_one' not in doc.element.xml
    links = doc.element.xpath('//w:hyperlink')
    assert len(links) == 2
    assert links[0].xpath('.//w:t')[0].text == '[1]'
    assert doc.part.rels[links[0].get(qn('r:id'))].target_ref == 'https://example.org/report'
    links = links[1:]
    assert links[0].xpath('.//w:t')[0].text == 'Official source'
    relation = doc.part.rels[links[0].get(qn('r:id'))]
    assert relation.is_external and relation.target_ref == 'https://example.org/report'
    assert links[0].xpath('.//w:color')[0].get(qn('w:val')) == '0563C1'


def test_unrecognized_or_unresolved_tables_are_not_discarded():
    draft = fixture()
    assert reader_source_blocks(draft, {})[0] == draft['content']
    draft['content'][-1]['content'][0]['content'][2]['content'][0]['content'][0]['text'] = '业务指标'
    assert reader_source_blocks(draft, {'src_one': {'name': 'Source'}})[0] == draft['content']


def test_export_without_source_append_keeps_authored_index():
    doc = Document()
    render_document(doc, fixture(), sources={'src_one': {'name': 'Source'}}, append_sources=False)
    assert len(doc.tables) == 1


def test_local_source_citation_has_a_resolving_document_anchor():
    doc = Document()
    render_document(doc, fixture(), sources={'src_one': {'name': 'Local source'}})
    citation = doc.element.xpath('//w:hyperlink')[0]
    anchor = citation.get(qn('w:anchor'))
    assert anchor and anchor in [n.get(qn('w:name')) for n in doc.element.xpath('//w:bookmarkStart')]


def test_legacy_citation_metadata_does_not_break_word_export():
    from io import BytesIO

    citations = [None, {}, {'source_id': ['src_one'], 'locator': 'Invalid source'},
                 {'source_id': 'src_one', 'locator': None},
                 {'source_id': 'src_one', 'locator': 17},
                 {'source_id': 'src_one', 'locator': ' Page 17 '}]
    original = deepcopy(citations)
    doc = Document()
    render_document(doc, fixture(), sources={'src_one': {'name': 'Local source'}}, citations=citations)
    blob = BytesIO(); doc.save(blob); blob.seek(0)
    restored = Document(blob)
    paragraphs = [''.join(p._p.xpath('.//w:t/text()')) for p in restored.paragraphs]
    assert '[1]' in paragraphs
    assert '1. Local source · Page 17' in paragraphs
    assert 'Invalid source' not in '\n'.join(paragraphs)
    assert citations == original


def test_saved_citation_locators_survive_plain_template_and_http_word_exports(tmp_path):
    import http.client
    from io import BytesIO
    import threading
    from briefloop.export_jobs import enqueue_export, generate_word, output_path
    from briefloop.server import make_server, _close_service
    from briefloop.templates import import_builtin

    server=make_server(tmp_path/'workspace',port=0,paused=True)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    try:
        store=server.store
        import_builtin(store)
        template_id=store.rows("SELECT id FROM templates WHERE name='通用报告·品牌绿'")[0]['id']
        source=store.add_source('Official synthetic source','Revenue 12',url='https://example.org/report')
        unused=store.add_source('Uncited material','Not used in the saved body')
        run=store.create_run({'title':'Revenue report','objective':'Preserve saved locating text'},[source['id']])
        locators=['Page 17, table Revenue, lines 2-3','附表 A 第 4 行']
        document={'type':'doc','content':[{'type':'paragraph','content':[
            {'type':'text','text':'Revenue 12 '},{'type':'citation','attrs':{'sourceId':source['id']}}]}]}
        brief=store.publish(run['id'],{'title':'Revenue report','editor_document':document,'citations':[
            {'source_id':source['id'],'locator':value} for value in [locators[0],' '+locators[0]+' ',locators[1],'',' \t ']
        ]+[{'source_id':unused['id'],'locator':'Do not append uncited locating text'}]})

        def check(blob):
            doc=Document(BytesIO(blob))
            paragraphs=[''.join(p._p.xpath('.//w:t/text()')) for p in doc.paragraphs]
            assert 'Revenue 12 [1]' in paragraphs
            assert '1. Official synthetic source · '+'；'.join(locators) in paragraphs
            assert all('\n'.join(paragraphs).count(locator)==1 for locator in locators)
            assert 'Do not append uncited locating text' not in '\n'.join(paragraphs)
            assert 'https://example.org/report' in [rel.target_ref for rel in doc.part.rels.values() if rel.is_external]

        for layout in (None,template_id):
            job=enqueue_export(store,brief['id'],layout)
            result=generate_word(store,job,threading.Event())
            assert result['version_id']==brief['id']
            check(output_path(store,job).read_bytes())
        connection=http.client.HTTPConnection('127.0.0.1',server.server_port,timeout=5)
        try:
            connection.request('GET','/api/download?format=docx&version='+brief['id'])
            response=connection.getresponse()
            assert response.status==200
            check(response.read())
        finally:connection.close()
        assert store.one('briefs',brief['id'])==brief
    finally:
        server.shutdown();thread.join(timeout=5);_close_service(server)
