from io import BytesIO
from zipfile import ZipFile
import json
import threading
from docx import Document
from docx.shared import Pt, RGBColor
from briefloop.store import Store
from briefloop.templates import import_template, prepare, export_template
from briefloop.export_jobs import enqueue_export, generate_word, output_path, export_input


def test_rebuild_keeps_local_run_overrides_out_of_body_defaults(tmp_path):
    from docx.enum.style import WD_STYLE_TYPE
    from briefloop.templates import rebuild_template_version
    original=Document()
    base=original.styles.add_style('Report Body',WD_STYLE_TYPE.PARAGRAPH)
    base.base_style=original.styles['Normal']
    base.font.color.rgb=RGBColor.from_string('222222');base.font.size=Pt(11);base.font.name='Calibri'
    original.add_paragraph('Monthly Report');original.add_paragraph('1. Summary')
    sample=original.add_paragraph(style=base)
    lead=sample.add_run('Key point: ');lead.bold=True;lead.font.color.rgb=RGBColor.from_string('CC0000')
    lead.font.size=Pt(16);lead.font.name='Cambria';lead.font.all_caps=True
    sample.add_run('This ordinary paragraph inherits the dark body style; only the lead is red and enlarged.')
    stream=BytesIO();original.save(stream);store=Store(tmp_path)
    record=import_template(store,'Monthly.docx',stream.getvalue(),prepare_job=False)
    prepared=prepare(store,record['id'],{'keep_blocks':[0],'paragraph_index':2,
        'sections':[{'section_id':'summary','title':'1. Summary','index':1}],
        'fields':[{'old':'Monthly Report','field':'title'}]})
    previous_path=store.root/'templates'/prepared['id']/'prepared.docx';previous_bytes=previous_path.read_bytes()
    rebuilt=rebuild_template_version(store,prepared['id'])
    source=store.add_source('Current material','The current operating result.')
    run=store.create_run({'title':'Current Report','objective':'Explain','template_id':rebuilt['id']},[source['id']])
    document={'type':'doc','content':[{'type':'paragraph','content':[
        {'type':'text','text':'Ordinary current text. '},
        {'type':'text','text':'Explicit blue text.','marks':[{'type':'textStyle','attrs':{'color':'#0000FF'}}]}]}]}
    brief=store.publish(run['id'],{'title':'Current Report','editor_document':document})
    rendered=Document(BytesIO(export_template(store,brief,document,{})))
    body=next(p for p in rendered.paragraphs if p.text.startswith('Ordinary current text.'))
    style=body.style
    assert style.font.color.rgb is None and style.font.size is None and style.font.name is None
    assert style.font.all_caps is None and style.font.bold is None
    assert style.base_style.name=='Report Body'
    assert str(style.base_style.font.color.rgb)=='222222' and style.base_style.font.size==Pt(11)
    assert style.base_style.font.name=='Calibri'
    assert str(body.runs[-1].font.color.rgb)=='0000FF'
    assert previous_path.read_bytes()==previous_bytes


def test_rebuild_is_a_new_layout_version_preserving_legacy_and_preamble(tmp_path):
    from copy import deepcopy
    import hashlib
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from briefloop.templates import rebuild_template_version,template
    original=Document()
    original.add_paragraph('Monthly Report')
    original.add_paragraph('August 2025')
    original.add_paragraph('1. Summary')
    callout=original.add_paragraph('An emphasized historical callout must not become the ordinary body style.')
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E3F2FD');callout._p.get_or_add_pPr().append(shade)
    border=OxmlElement('w:pBdr');left=OxmlElement('w:left');left.set(qn('w:val'),'single');border.append(left)
    callout._p.get_or_add_pPr().append(border)
    callout.runs[0].bold=True;callout.runs[0].font.color.rgb=RGBColor.from_string('1A568E')
    ordinary=original.add_paragraph()
    ordinary.add_run('Overview: ').bold=True
    ordinary.add_run('The ordinary body uses an unshaded paragraph and keeps emphasis limited to its lead phrase.')
    for run in ordinary.runs:run.font.size=Pt(11);run.font.color.rgb=RGBColor.from_string('222222')
    original.add_paragraph('2. Operating implications')
    original.add_paragraph('This is another ordinary paragraph with enough text to identify its role.')
    source=BytesIO();original.save(source)
    store=Store(tmp_path)
    old=import_template(store,'Monthly.docx',source.getvalue())
    old=prepare(store,old['id'],{'keep_blocks':[0,1],'paragraph_index':3,
        'sections':[{'section_id':'summary','title':'1. Summary','index':2,'purpose':'Changes'},
                    {'section_id':'impact','title':'2. Operating implications','index':5,'purpose':'Impact'}],
        'fields':[{'old':'Monthly Report','field':'title'},{'old':'August 2025','field':'period'}]})
    assert old['spec']['body_sample_index']==4
    # A synthetic already-prepared v1 record models the shipped callout bug.
    # Only this fixture is changed; rebuilding must leave it byte-for-byte alone.
    legacy_path=store.root/'templates'/old['id']/'prepared.docx'
    legacy_doc=Document(legacy_path);body_style=legacy_doc.styles[old['spec']['styles']['paragraph']]
    if body_style.element.pPr is not None:body_style.element.remove(body_style.element.pPr)
    body_style.element.append(deepcopy(callout._p.pPr));body_style.font.bold=True;body_style.font.color.rgb=RGBColor.from_string('1A568E')
    for section in old['spec']['sections']:
        style=legacy_doc.styles[old['spec']['styles']['heading:'+section['section_id']]]
        style.paragraph_format.keep_with_next=None;style.paragraph_format.keep_together=None
    legacy_doc.save(legacy_path)
    legacy_spec={key:value for key,value in old['spec'].items() if key not in ('layout_version','cover_title_present','body_sample_index','preparation_spec')}
    legacy_spec['prepared_hash']=hashlib.sha256(legacy_path.read_bytes()).hexdigest()
    with store.tx() as c:c.execute('UPDATE templates SET spec=? WHERE id=?',(json.dumps(legacy_spec),old['id']))
    before_bytes=legacy_path.read_bytes();before_spec=template(store,old['id'])['spec']
    title='Current monthly report'
    paragraph=lambda text:{'type':'paragraph','content':[{'type':'text','text':text}]}
    document={'type':'doc','content':[paragraph('Preface before the first heading must survive.'),
        {'type':'table','content':[{'type':'tableRow','content':[{'type':'tableCell','content':[paragraph('Preamble table must survive.')]}]}]},
        {'type':'heading','attrs':{'level':1},'content':[{'type':'text','text':title}]},
        paragraph('Reporting period note must survive.'),
        {'type':'heading','attrs':{'level':2,'blockId':'summary'},'content':[{'type':'text','text':'1. Summary'}]},
        *[paragraph('This paragraph presents the current business change and its operating consequences. '*18) for _ in range(3)],
        {'type':'heading','attrs':{'level':2,'blockId':'impact'},'content':[{'type':'text','text':'2. Operating implications'}]},
        paragraph('The next section remains with its explanation, while the earlier preface and table stay in place.')]}
    current_source=store.add_source('Current synthetic material','Current business change and operating consequences.')
    def export_for(record):
        run=store.create_run({'title':title,'objective':'Explain','period':'August 2026','template_id':record['id']},[current_source['id']])
        brief=store.publish(run['id'],{'title':title,'editor_document':document})
        return export_template(store,brief,document,{})
    legacy_output=export_for(old);(tmp_path/'legacy-output.docx').write_bytes(legacy_output)
    jobs_before=len(store.rows('SELECT id FROM jobs'))
    rebuilt=rebuild_template_version(store,old['id'])
    assert rebuilt['status']=='ready' and rebuilt['parent_id']==old['id'] and rebuilt['revision']==old['revision']+1
    assert len(store.rows('SELECT id FROM jobs'))==jobs_before
    assert legacy_path.read_bytes()==before_bytes and template(store,old['id'])['spec']==before_spec
    rendered=export_for(rebuilt);(tmp_path/'rebuilt-output.docx').write_bytes(rendered)
    current=Document(BytesIO(rendered));legacy=Document(BytesIO(legacy_output))
    assert sum(p.text==title for p in current.paragraphs)==1 and sum(p.text==title for p in legacy.paragraphs)==2
    assert any(p.text=='Preface before the first heading must survive.' for p in current.paragraphs)
    assert any(p.text=='Reporting period note must survive.' for p in current.paragraphs)
    assert current.tables[0].cell(0,0).text=='Preamble table must survive.'
    body_style=current.styles[rebuilt['spec']['styles']['paragraph']]
    assert not body_style.element.xpath('./w:pPr/w:shd|./w:pPr/w:pBdr') and body_style.font.bold is not True
    assert body_style.font.size==Pt(11) and str(body_style.font.color.rgb)=='222222'
    for section in rebuilt['spec']['sections']:
        style=current.styles[rebuilt['spec']['styles']['heading:'+section['section_id']]]
        assert style.paragraph_format.keep_with_next is True and style.paragraph_format.keep_together is True
    assert (tmp_path/'legacy-output.docx').read_bytes()==legacy_output


def test_template_reuses_direct_styles_replaces_old_facts_and_locks_version(tmp_path):
    original=Document();original.sections[0].header.paragraphs[0].text='Example Corp · August 2025'
    original.add_paragraph('Monthly Report')
    original.add_paragraph('August 2025')
    heading=original.add_paragraph('一、核心摘要');heading.runs[0].bold=True;heading.runs[0].font.size=Pt(17);heading.runs[0].font.color.rgb=RGBColor.from_string('17466B')
    body=original.add_paragraph('OLD FACT REVENUE 987654');body.runs[0].font.size=Pt(11)
    original.add_paragraph('二、经营影响')
    original.add_paragraph('OLD FACT 2025')
    stream=BytesIO();original.save(stream)
    store=Store(tmp_path);record=import_template(store,'Monthly.docx',stream.getvalue())
    prepared=prepare(store,record['id'],{'keep_blocks':[0,1],'paragraph_index':3,
        'sections':[{'section_id':'summary','title':'一、核心摘要','index':2,'purpose':'变化'},
                    {'section_id':'impact','title':'二、经营影响','index':4,'purpose':'影响'}],
        'fields':[{'old':'Monthly Report','field':'title'},{'old':'August 2025','field':'period'}]})
    assert prepared['status']=='ready'
    source=store.add_source('Current','Current revenue 12')
    run=store.create_run({'title':'August Report','objective':'Explain','period':'August 2026','template_id':record['id']},[source['id']])
    assert json.loads(run['requirements'])['sections'][0]['section_id']=='summary'
    doc={'type':'doc','content':[{'type':'heading','attrs':{'level':2,'blockId':'summary'},'content':[{'type':'text','text':'一、核心摘要'}]},
        {'type':'paragraph','content':[{'type':'text','text':'CURRENT FACT 12'}]},
        {'type':'heading','attrs':{'level':2,'blockId':'impact'},'content':[{'type':'text','text':'二、经营影响'}]},
        {'type':'paragraph','content':[{'type':'text','text':'Impact on customers'}]}]}
    brief=store.publish(run['id'],{'title':'August Report','editor_document':doc})
    job=enqueue_export(store,brief['id']);result=generate_word(store,job,threading.Event())
    with ZipFile(output_path(store,job)) as archive:
        xml=archive.read('word/document.xml').decode();styles=archive.read('word/styles.xml').decode()
        header=archive.read('word/header1.xml').decode()
        assert 'CURRENT FACT 12' in xml and 'OLD FACT' not in xml and 'August 2025' not in xml
        assert 'August 2026' in xml and 'August 2026' in header
        assert '17466B' in styles and 'w:sz w:val="34"' in styles
    store.set_meta('settings',{**store.settings(),'default_template_id':'new-default'})
    assert json.loads(store.one('runs',run['id'])['requirements'])['template_id']==record['id']
    assert result['version_id']==brief['id']
