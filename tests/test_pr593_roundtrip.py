from io import BytesIO
from zipfile import ZipFile
import json
from docx import Document
from docx.shared import Pt, RGBColor
from PIL import Image
from briefloop.store import Store
from briefloop.exports import docx_bytes
from briefloop.word_import import import_revision
from briefloop.templates import import_template, prepare, export_template
from briefloop.figures import register_figure
from briefloop.figure_support import markdown_bundle


def paragraph(text):return {'type':'paragraph','content':[{'type':'text','text':text}]}
def heading(text):return {'type':'heading','attrs':{'level':2,'blockId':'summary'},'content':[{'type':'text','text':text}]}
def setup(tmp_path):
    store=Store(tmp_path);source=store.add_source('Evidence','Revenue data')
    run=store.create_run({'title':'Report','objective':'Read'},[source['id']])
    return store,source,run


def test_word_roundtrip_preserves_all_preheading_blocks(tmp_path):
    store,source,run=setup(tmp_path)
    table={'type':'table','content':[{'type':'tableRow','content':[{'type':'tableCell','content':[paragraph('Intro table 12')]}]}]}
    document={'type':'doc','content':[paragraph('Important introduction 12'),table,heading('Summary'),paragraph('Body')]}
    brief=store.publish(run['id'],{'title':'Report','editor_document':document})
    result=import_revision(store,brief['id'],'unchanged.docx',docx_bytes(document=document))
    assert result['status']=='imported'
    nodes=json.loads(result['version']['editor_document'])['content']
    assert [n['type'] for n in nodes]==['paragraph','table','heading','paragraph']
    assert 'Important introduction 12' in result['version']['markdown']
    assert 'Intro table 12' in result['version']['markdown']


def test_markdown_bundle_uses_each_node_caption_without_registered_stale_caption(tmp_path):
    store,source,run=setup(tmp_path);path=store.root/'image.png';Image.new('RGB',(20,20),'red').save(path)
    figure=register_figure(store,run['id'],path,'Revenue','USD millions',[source['id']])
    nodes=[{'type':'image','attrs':{'src':'briefloop-figure:'+figure['figure_id'],'caption':caption}} for caption in ['USD thousands','']]
    brief=store.publish(run['id'],{'title':'Report','editor_document':{'type':'doc','content':nodes}})
    with ZipFile(BytesIO(markdown_bundle(store,brief))) as archive:
        text=archive.read('report.md').decode()
        assert 'USD millions' not in text and text.count('USD thousands')==1
        assert text.count('来源：Evidence')==2


def test_template_keeps_style_chain_and_updates_even_page_fields(tmp_path):
    store,source,run=setup(tmp_path);doc=Document()
    doc.settings.odd_and_even_pages_header_footer=True
    doc.sections[0].even_page_header.paragraphs[0].text='Header 2025'
    doc.sections[0].even_page_footer.paragraphs[0].text='Footer 2025'
    doc.styles['Heading 1'].font.size=Pt(28);doc.styles['Heading 1'].font.color.rgb=RGBColor.from_string('FF0000')
    doc.add_paragraph('Cover 2025');doc.add_paragraph('Summary',style='Heading 1');doc.add_paragraph('Old body')
    stream=BytesIO();doc.save(stream)
    template=import_template(store,'template.docx',stream.getvalue())
    prepare(store,template['id'],{'sections':[{'section_id':'summary','title':'Summary','index':1,'purpose':'Summary'}],
        'paragraph_index':2,'keep_blocks':[0],'fields':[{'old':'2025','field':'period'}]})
    newrun=store.create_run({'title':'Report','objective':'Read','period':'2026','template_id':template['id']},[source['id']])
    document={'type':'doc','content':[heading('Summary'),paragraph('New body')]}
    brief=store.publish(newrun['id'],{'title':'Report','editor_document':document})
    final=Document(BytesIO(export_template(store,brief,document,{})))
    style=next(p.style for p in final.paragraphs if p.text=='Summary')
    assert style.base_style.name=='Heading 1'
    assert style.base_style.font.size==Pt(28) and str(style.base_style.font.color.rgb)=='FF0000'
    assert final.sections[0].even_page_header.paragraphs[0].text=='Header 2026'
    assert final.sections[0].even_page_footer.paragraphs[0].text=='Footer 2026'
