from io import BytesIO
import json
from docx import Document
from briefloop.store import Store
from briefloop.exports import docx_bytes
from briefloop.word_import import import_revision


def test_word_revision_binds_selected_base_preserves_rich_text_and_stops_bad_alignment(tmp_path):
    store=Store(tmp_path);source=store.add_source('Evidence','Revenue 14')
    run=store.create_run({'title':'Report','objective':'Explain'},[source['id']])
    document={'type':'doc','content':[{'type':'heading','attrs':{'level':2,'blockId':'summary'},'content':[{'type':'text','text':'Summary'}]},
        {'type':'paragraph','content':[{'type':'text','text':'Revenue 12','marks':[{'type':'textStyle','attrs':{'color':'#c00000'}}]},{'type':'citation','attrs':{'sourceId':source['id']}}]}]}
    base=store.publish(run['id'],{'title':'Report','editor_document':document})
    word=Document(BytesIO(docx_bytes(document=document,source_records={source['id']:source})))
    word.paragraphs[1].runs[0].text='Revenue 14'
    data=BytesIO();word.save(data)
    result=import_revision(store,base['id'],'revision.docx',data.getvalue())
    assert result['status']=='imported'
    revised=result['version'];assert revised['parent_id']==base['id'] and 'Revenue 14' in revised['markdown']
    parsed=json.loads(revised['editor_document'])
    assert parsed['content'][0]['attrs']['blockId']=='summary'
    assert parsed['content'][1]['content'][0]['marks'][0]['attrs']['color']=='#c00000'
    assert len(store.rows('SELECT * FROM feedback'))==1
    word.paragraphs[0].text='Unrelated section';data=BytesIO();word.save(data)
    pending=import_revision(store,revised['id'],'different.docx',data.getvalue())
    assert pending['status']=='needs_alignment'
    assert len(store.rows('SELECT * FROM feedback'))==1
    assert store.one('sources',pending['source_id'])['status']=='ready'
